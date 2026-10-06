"""Build the supplied #126 timetable at source-cell granularity, without an API.

This reproducible local artifact includes both timetable grids and the detailed
main-session table. It does not replace the normal multi-source Gemini pipeline.
"""
from __future__ import annotations

import argparse
from dataclasses import asdict
from datetime import date, datetime, timedelta
import hashlib
import json
from pathlib import Path
import re
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from bs4 import BeautifulSoup
from build import render_site
from shared.schedule import Schedule, Session, DaySchedule, save_schedule, time_to_minutes
from working_groups.ran1.models import DAY_ORDER, ran1_timeline
from working_groups.ran1.parser import parse_docx, build_room_list
from working_groups.ran1.agreements import parse_agreements

SCHEDULE_FILE = Path('downloads/ran1/Chair_notes/RAN1#126 online and offline schedules - v00.docx')
NOTE_FILE = Path('downloads/ran1/Chair_notes/Chair notes RAN1#126_eom.docx')
# The manual interpretation below is deliberately locked to the reviewed input.
SCHEDULE_SHA = 'd52af6b894d4ccde2b1d1180bb1565b3f75128d5993f48fc60dadba58eb7496f'
TOPICS = {'AI/ML': '9.1', 'MIMO': '9.2', 'A-IoT': '9.3', 'NTN-NR': '9.4', 'TEI': '11'}


def agendas(text, sections):
    """Exact IDs and explicit .x wildcards only; never infer a numeric prefix."""
    found = []
    for line in text.splitlines():
        explicit = re.match(r'^\s*\.?((?:\d+\.)+\d+(?:\.x)?|\d+\.x)\b', line, re.I)
        if explicit:
            found.append(explicit[1])
        match = re.search(r'\b(?:AI|Agenda items?)\s+([\d.,/ ]+)', line, re.I)
        if match:
            found.extend(re.findall(r'\d+(?:\.\d+)*', match[1]))
        for topic, ai in TOPICS.items():
            if re.search(r'(?<![\w-])' + re.escape(topic) + r'(?![\w-])', line):
                found.append(ai)
    expanded, mappings = [], []
    for token in found:
        if token.lower().endswith('.x'):
            prefix = token[:-1]
            targets = [ai for ai in sections if ai.startswith(prefix)]
            mappings.append({'source': token, 'targets': targets})
            expanded.extend(targets)
        else:
            expanded.append(token)
    return list(dict.fromkeys(expanded)), mappings


def build(output):
    if hashlib.sha256(SCHEDULE_FILE.read_bytes()).hexdigest() != SCHEDULE_SHA:
        raise ValueError('Timetable bytes changed; review source-cell interpretation before rebuilding')
    data = parse_agreements(NOTE_FILE, 'ran1#126')
    sections = data['sections']
    cells, metadata = parse_docx(SCHEDULE_FILE)
    assert len(cells) == 86 and len(metadata) == 2
    day_rooms = build_room_list(metadata)
    doc = Document(SCHEDULE_FILE)
    assert len(doc.tables) == 3
    detailed = {}
    for row, start in [(1, '08:30'), (3, '11:00'), (5, '14:30'), (7, '17:00')]:
        for col, day in enumerate(DAY_ORDER, 1):
            value = doc.tables[2].rows[row].cells[col].text.strip()
            if value:
                detailed[(day, start)] = value
    days = {day: [] for day in DAY_ORDER}
    audit = []
    missing = set()
    for index, cell in enumerate(cells):
        rooms = day_rooms[cell.day]
        columns = [i for i, room in enumerate(rooms)
                   if room.table_index == cell.table_index
                   and room.room_index_in_table in cell.room_indices]
        assert columns and columns == list(range(min(columns), max(columns) + 1))
        start = cell.fallback_start_time or cell.time_block_start
        end = cell.time_block_end
        # Opening and dinner times are explicit in the source, not inferred
        # subtopic allocations. Keep the source's approximate Friday closing.
        notes = [line.strip() for line in cell.text.splitlines() if line.strip()]
        opening = re.search(r'commences at (\d{2}:\d{2})', cell.text)
        if opening:
            start = opening[1]
        closing = re.search(r'expected to close at (\d{2}:\d{2})', cell.text)
        if closing:
            end = closing[1]
        if not cell.fallback_start_time:
            following = [x.fallback_start_time for x in cells
                         if x.day == cell.day and x.table_index == cell.table_index
                         and x.time_block_index == cell.time_block_index and x.fallback_start_time
                         and set(x.room_indices).intersection(cell.room_indices)]
            if following:
                end = min(following)
        extra = detailed.get((cell.day, cell.time_block_start)) if cell.table_index == 0 and 0 in cell.room_indices and not cell.fallback_start_time else None
        source_text = cell.text
        if extra:
            notes += ['Detailed main-session schedule:'] + extra.splitlines()
            source_text += '\n' + extra
        ids, mappings = agendas(source_text, sections)
        missing.update(ai for ai in ids if ai not in sections)
        chairs = list(dict.fromkeys(re.findall(r'\b(?:Xiaodong|Hiroki|Sorouri?)\b', cell.text)))
        label = ' · '.join(line.strip() for line in cell.text.splitlines() if line.strip())
        descriptions = [{'agenda_item': ai, 'description': sections[ai]['title']}
                        for ai in ids if ai in sections]
        session = Session(
            name=label, duration_minutes=time_to_minutes(end) - time_to_minutes(start),
            start_time=start, end_time=end, day=cell.day,
            room_col_start=min(columns) + 2, room_col_end=max(columns) + 3,
            chair=', '.join(chairs) or None, agenda_item=', '.join(ids) or None,
            group_header='Online' if cell.table_index == 0 else 'Offline',
            agenda_descriptions=descriptions, notes=notes,
        )
        assert session.duration_minutes > 0
        days[cell.day].append(session)
        audit.append({'source_cell': index, **asdict(cell), 'detailed_main_text': extra,
                      'start': start, 'end': end, 'room_columns': columns,
                      'agenda_items': ids, 'wildcards': mappings})
    assert not missing, f'Agenda IDs absent in EOM: {missing}'
    schedule = Schedule(
        meeting_name='RAN1#126 · Online and offline schedules',
        days=[DaySchedule(day, day_rooms[day], days[day],
                          date=(date(2026, 8, 24) + timedelta(days=i)).isoformat(),
                          timeline=ran1_timeline()) for i, day in enumerate(DAY_ORDER)],
        source_file=SCHEDULE_FILE.name, source_files=[SCHEDULE_FILE.name],
        generated_at=datetime.now().astimezone().isoformat(timespec='seconds'),
        timezone='Europe/Amsterdam', wg_id='ran1', meeting_id='ran1#126',
        starts_on='2026-08-24', ends_on='2026-08-28',
        starts_at='2026-08-24T09:00:00+02:00', ends_at='2026-08-28T17:00:00+02:00',
        chairman_agreements=data,
    )
    save_schedule(schedule, output / 'ran1/schedule.json')
    render_site({'default_wg': 'ran1', 'working_groups': [{'id': 'ran1'}],
                 'presentation': {'notice': 'Schedule v00 · Chairman notes EOM · 24–28 August 2026. Times show timetable blocks; topic durations are preserved in cell details.'}},
                output, {'ran1': schedule})
    page = BeautifulSoup((output / 'ran1/index.html').read_text(), 'html.parser')
    manifest = json.loads(page.select_one('#agreement-data').string)
    assert len(page.select('.session-block')) == 86
    assert manifest['meeting_id'] == 'ran1#126' and manifest['status'] == 'ready'
    for ai, entry in manifest['sections'].items():
        if entry.get('url'):
            fragment = output / 'ran1' / entry['url']
            assert fragment.is_file() and sections[ai]['html'] in fragment.read_text()
    # Every source cell and supplemental main-session paragraph has an audit
    # entry. The popup includes all source lines without assigning extra times.
    assert len(audit) == len(cells)
    assert sum(bool(x['detailed_main_text']) for x in audit) == len(detailed)
    linked = {ai for x in audit for ai in x['agenda_items']}
    fallbacks = {ai: {'equations': value['html'].count('math-fallback'),
                      'objects': value['html'].count('[Object:'),
                      'images': value['html'].count('[Image/shape:')}
                 for ai, value in sections.items()
                 if any(token in value['html'] for token in ['math-fallback', '[Object:', '[Image/shape:'])}
    report = {'meeting_id': 'ran1#126', 'schedule_file': SCHEDULE_FILE.name,
              'schedule_sha256': SCHEDULE_SHA, 'chairman_file': NOTE_FILE.name,
              'chairman_sha256': data['sha256'], 'source_tables': 3,
              'schedule_cells': len(cells), 'detailed_main_cells': len(detailed),
              'days': {day: len(sessions) for day, sessions in days.items()},
              'agenda_sections': len(sections), 'nonempty_sections': sum(bool(s['html']) for s in sections.values()),
              'excluded_tdoc_rows': sum(s['excluded_tdoc_rows'] for s in sections.values()),
              'linked_agendas': sorted(linked), 'missing_agendas': sorted(missing),
              'rendering_fallbacks': fallbacks, 'warnings': data['warnings'],
              'interpretation': 'Source-cell timetable; no Gemini call or cached merged sessions used.',
              'cells': audit}
    (output / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps({k: v for k, v in report.items() if k not in {'cells', 'linked_agendas'}}, ensure_ascii=False, indent=2))
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=Path('test_runs/ran1-126/site'))
    build(parser.parse_args().output_dir)
