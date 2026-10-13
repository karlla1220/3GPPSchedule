"""Vice-chair session notes read with the chairman note: one merged state per item."""

from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from bs4 import BeautifulSoup
from docx import Document
import pytest

from shared.agreement_assets import agreement_assets, package_agreements
from working_groups.ran1 import agreements as a

MEETING = "ran1#126bis"
BOILERPLATE = "Note 1: Including proposals for downlink control channel design."
CHAIR = "Chair notes RAN1#126bis_v{:02d}.docx"
SESSION = "draft_R1-2607288_Session Notes of AI 10.5.4_v{:02d}.docx"


def chair_note(path, extra=()):
    """A chairman note whose vice-chair items hold the agenda text alone, or more."""
    d = Document()
    d.add_heading("10.5 Multi-antenna system", level=1)
    d.add_paragraph("Note 1: Including proposals for both Sing-TRP and Multi-TRP")
    d.add_heading("10.5.4 Downlink control channel, scheduling and HARQ operation", level=2)
    d.add_heading("10.5.4.1 Downlink control channel, scheduling for downlink and uplink", level=3)
    d.add_paragraph(BOILERPLATE)
    for text in extra:
        d.add_paragraph(text)
    d.add_heading("10.5.4.2 HARQ related Aspects", level=3)
    d.add_paragraph("Note 1: Including proposals for timing line.")
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path


def session_note(path, agreements, *, meeting="#126bis", numbers=("10.5.0", "10.5.0.1")):
    """A session note: TDoc header, then the agenda headings numbered from scratch."""
    d = Document()
    d.add_paragraph(f"3GPP TSG RAN WG1 {meeting}\t\t\tR1-2607288")
    d.add_paragraph("Agenda Item:\t10.5.4")
    d.add_paragraph("Title:\tSession Notes of AI 10.5.4")
    d.add_heading(f"{numbers[0]} Downlink control channel, scheduling and HARQ operation", level=2)
    d.add_heading(f"{numbers[1]} Downlink control channel, scheduling for downlink and uplink", level=3)
    d.add_paragraph(BOILERPLATE)
    for text in agreements:
        d.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path


def text(section):
    return BeautifulSoup(section["html"], "html.parser").get_text("\n")


def added(section):
    soup = BeautifulSoup(section["html"], "html.parser")
    return [b.get_text() for b in soup.select(".agreement-added")]


def offline_build(root, chair, notes, previous, now):
    """Build from ref_in_manual and the session note download folder only."""
    manual = root / "ref_in_manual/ran1"
    sessions = root / a.SESSION_NOTES_DIR / "Hiroki_notes"
    for folder in (manual, sessions):
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.glob("*"):
            old.unlink()
    if chair is not None:
        (manual / chair.name).write_bytes(chair.read_bytes())
    for note in notes:
        (sessions / note.name).write_bytes(note.read_bytes())
    result, identity = a.build_agreements(
        cfg={}, meeting_id=MEETING, schedule_path=root / "s.docx",
        offline=True, previous=previous, now=now,
    )
    return result, identity


@pytest.mark.parametrize(("lines", "expected"), [
    (["3GPP TSG RAN WG1 #126bis\t\tR1-2607288", "Agenda Item:\t10.5.4"], ("ran1#126bis", "10.5.4")),
    (["3GPP TSG-RAN WG1 Meeting #126-bis", "Agenda item: 9.2."], ("ran1#126bis", "9.2")),
    (["3GPP TSG RAN WG1 #126b", "Agenda Item:\t7"], ("ran1#126bis", "7")),
    (["3GPP TSG RAN WG1 #127", "Source: Ad-Hoc Chair"], ("ran1#127", None)),
    (["Jeju, Korea, Oct. 12th-16th, 2026"], (None, None)),
])
def test_header_states_the_meeting_and_the_agenda_item(lines, expected):
    fields = a._header_fields(lines)
    assert (fields["meeting_id"], fields["agenda_item"]) == expected


@pytest.mark.parametrize(("name", "ai", "tdoc"), [
    ("draft_R1-2607285_Session Notes of AI 8.2_v01.docx", "8.2", "R1-2607285"),
    ("Draft R1-2607314 Session Notes of AI 10.5.2 (MIMO Operation) - v02.docx", "10.5.2", "R1-2607314"),
    ("R1-2606764_Session Notes of AI 10.5.4_final.docx", "10.5.4", "R1-2606764"),
])
def test_session_note_names(name, ai, tdoc):
    assert a.session_note_agenda(name) == ai
    assert a.session_note_id(name) == tdoc


def test_restarted_numbering_is_anchored_to_the_stated_item(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    chair = a.parse_agreements(chair_note(tmp_path / CHAIR.format(5)), MEETING)
    titles = {ai: s["title"] for ai, s in chair["sections"].items()}
    note = a.parse_session_note(session_note(tmp_path / SESSION.format(1), ["Agreement:"]),
                                MEETING, titles=titles)
    assert list(note["sections"]) == ["10.5.4", "10.5.4.1"]
    # Without a chairman note the first heading is the stated item itself.
    alone = a.parse_session_note(tmp_path / SESSION.format(1), MEETING)
    assert list(alone["sections"]) == ["10.5.4", "10.5.4.1"]


def test_a_note_opening_with_the_parent_item_is_anchored_by_title(tmp_path, monkeypatch):
    """A note for 10.5.4.1 that opens with the 10.5.4 heading."""
    monkeypatch.chdir(tmp_path)
    chair = a.parse_agreements(chair_note(tmp_path / CHAIR.format(5)), MEETING)
    titles = {ai: s["title"] for ai, s in chair["sections"].items()}
    d = Document()
    d.add_paragraph("3GPP TSG RAN WG1 #126bis")
    d.add_paragraph("Agenda Item:\t10.5.4.1")
    d.add_heading("8.1 Downlink control channel, scheduling and HARQ operation", level=2)
    d.add_paragraph("Email discussion")
    d.add_heading("8.1.1 Downlink control channel, scheduling for downlink and uplink", level=3)
    d.add_paragraph("Agreement: one")
    d.add_heading("8.1.2 HARQ related Aspects", level=3)
    d.add_paragraph("Agreement: two")
    d.save(tmp_path / "R1-2600001 Session Notes of AI 10.5.4.1 - v00.docx")
    note = a.parse_session_note(tmp_path / "R1-2600001 Session Notes of AI 10.5.4.1 - v00.docx",
                                MEETING, titles=titles)
    assert list(note["sections"]) == ["10.5.4", "10.5.4.1", "10.5.4.2"]


def test_a_note_of_another_meeting_is_rejected_by_its_header(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = session_note(tmp_path / SESSION.format(0), ["Agreement:"], meeting="#126")
    with pytest.raises(ValueError, match="ran1#126"):
        a.parse_session_note(path, MEETING)


def test_shared_blocks_appear_once_and_session_blocks_stay_in_place(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    result, identity = offline_build(
        tmp_path, chair_note(tmp_path / "in" / CHAIR.format(5)),
        [session_note(tmp_path / "in" / SESSION.format(1), ["Agreement:", "Support DCI format 1_X."])],
        None, "2026-10-13T08:00:00+00:00",
    )
    section = result["sections"]["10.5.4.1"]
    shown = text(section)
    assert shown.count(BOILERPLATE) == 1
    assert shown.index(BOILERPLATE) < shown.index("Agreement:") < shown.index("Support DCI format 1_X.")
    assert set(section["sources"]) == {"chair", "R1-2607288"}
    assert "parts" not in section
    # The chairman note's other items are untouched.
    assert set(result["sections"]["10.5.4.2"]["sources"]) == {"chair"}
    assert [n["name"] for n in identity["session_notes"]] == [SESSION.format(1)]
    assert result["session_notes"]["R1-2607288"]["agenda_item"] == "10.5.4"


def test_gathering_into_the_chairman_note_is_no_change_and_session_additions_are(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    src = tmp_path / "in"
    s1 = session_note(src / SESSION.format(1), ["Agreement:", "Support DCI format 1_X."])
    first, _ = offline_build(tmp_path, chair_note(src / CHAIR.format(5)), [s1],
                             None, "2026-10-13T08:00:00+00:00")
    baseline, _ = package_agreements(first)

    # The chair gathers the session note whenever it does: the merged text is the same.
    gathered = chair_note(src / CHAIR.format(6), ["Agreement:", "Support DCI format 1_X."])
    second, _ = offline_build(tmp_path, gathered, [s1], baseline, "2026-10-14T08:00:00+00:00")
    section = second["sections"]["10.5.4.1"]
    assert (section["change"], section["added"]) == ("initial", [])
    assert text(section).count("Support DCI format 1_X.") == 1
    stored, _ = package_agreements(second, baseline)

    # The session note runs ahead again: only its new block is highlighted.
    s2 = session_note(src / SESSION.format(2),
                      ["Agreement:", "Support DCI format 1_X.", "Agreement:", "Support DCI format 1_Y."])
    third, _ = offline_build(tmp_path, gathered, [s2], stored, "2026-10-15T08:00:00+00:00")
    section = third["sections"]["10.5.4.1"]
    assert section["change"] == "updated"
    assert section["changed_in"] == SESSION.format(2)
    assert section["changed_at"] == "2026-10-15T08:00:00+00:00"
    assert added(section) == ["Agreement:", "Support DCI format 1_Y."]
    assert third["sections"]["10.5.4.2"]["change"] == "initial"


def test_a_reworded_block_is_shown_once_in_the_later_wording():
    def part(words, key):
        return {"html": f"<p>{words}</p>", "key": key}

    chair = [part("Note 1", "n"), part("Adopt the TP to TS 38.214 in principle. Final CR in R1-26xxxxx is agreed.", "c")]
    session = [part("Note 1", "n"), part("Adopt the TP to TS 38.214 in principle. Final CR in R1-2607777 is agreed.", "s"),
               part("Agreement: something else entirely", "x")]
    for newer, kept in ((False, "c"), (True, "s")):
        merged = a._merge_parts(chair, session, extra_newer=newer)
        assert [p["key"] for p in merged] == ["n", kept, "x"]


def test_two_tdoc_lists_fold_into_the_longer_one():
    fold = lambda rows: {"html": f'<details class="tdoc-list"><summary>{rows} TDocs</summary>{"r" * rows}</details>', "key": None}
    chair = [{"html": "<p>Note</p>", "key": "n"}, fold(9), {"html": "<p>Agreement</p>", "key": "a"}]
    session = [{"html": "<p>Note</p>", "key": "n"}, fold(7), {"html": "<p>Agreement</p>", "key": "a"}]
    merged = a._merge_parts(chair, session)
    assert [p["html"] for p in merged].count(fold(9)["html"]) == 1 and fold(7) not in merged
    assert a._merge_parts(session, chair)[1] == fold(9)


def test_blank_lines_alone_are_not_merged_in():
    blank = {"html": "<p></p>", "key": None}
    chair = [{"html": "<p>Note</p>", "key": "n"}]
    assert a._merge_parts(chair, chair + [blank, blank]) == chair


def test_first_build_with_session_notes_is_a_baseline(tmp_path, monkeypatch):
    """What earlier builds did not read is no edit, but it is dated by its note."""
    monkeypatch.chdir(tmp_path)
    src = tmp_path / "in"
    chair = chair_note(src / CHAIR.format(5))
    before, _ = offline_build(tmp_path, chair, [], None, "2026-10-12T08:00:00+00:00")
    before, _ = package_agreements(before)
    before.pop("session_notes")  # Published before session notes were read.
    before["parser_version"] = a.PARSER_VERSION  # Not a parser change alone.
    for section in before["sections"].values():
        section.pop("sources")
    s1 = session_note(src / SESSION.format(1), ["Agreement:", "Support DCI format 1_X."])
    after, _ = offline_build(tmp_path, chair, [s1], before, "2026-10-13T08:00:00+00:00")
    section = after["sections"]["10.5.4.1"]
    assert (section["change"], section["added"]) == ("initial", [])
    assert section["changed_at"] == "2026-10-13T08:00:00+00:00"
    assert after["sections"]["10.5.4.2"]["changed_at"] == "2026-10-12T08:00:00+00:00"


def test_an_unreadable_session_note_keeps_its_items_as_published(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    src = tmp_path / "in"
    chair = chair_note(src / CHAIR.format(5))
    s1 = session_note(src / SESSION.format(1), ["Agreement:", "Support DCI format 1_X."])
    first, _ = offline_build(tmp_path, chair, [s1], None, "2026-10-13T08:00:00+00:00")
    stored, fragments = package_agreements(first)
    broken = src / SESSION.format(2)
    broken.write_bytes(b"not a Word document")
    second, identity = offline_build(tmp_path, chair, [broken], stored, "2026-10-14T08:00:00+00:00")
    assert second["sections"]["10.5.4.1"] == stored["sections"]["10.5.4.1"]
    assert any(SESSION.format(2) in w for w in second["warnings"])
    assert identity["session_notes"][0]["name"] == SESSION.format(2)
    # The kept section's page is still published.
    again, _ = package_agreements(second, stored)
    assert again["sections"]["10.5.4.1"]["file"] in fragments


def test_session_notes_alone_before_the_first_chairman_note(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    s1 = session_note(tmp_path / "in" / SESSION.format(1), ["Agreement:"])
    result, identity = offline_build(tmp_path, None, [s1], None, "2026-10-13T08:00:00+00:00")
    assert result["status"] == "ready"
    assert result["sections"]["10.5.4.1"]["changed_in"] == SESSION.format(1)
    assert "name" not in identity and identity["session_notes"]


def test_the_latest_version_of_each_downloaded_note_is_used(tmp_path):
    folder = tmp_path / "Hiroki_notes"
    folder.mkdir()
    for name in (SESSION.format(0), SESSION.format(2), SESSION.format(1),
                 "draft_R1-2607285_Session Notes of AI 8.2_v00.docx", "~$aft_R1-2607285_Session Notes of AI 8.2_v01.docx",
                 "RAN1#126bis schedule for Hiroki Adhoc2 sessions_v05.docx"):
        (folder / name).write_bytes(b"x")
    assert [p.name for p in a.local_session_notes(tmp_path)] == [
        "draft_R1-2607285_Session Notes of AI 8.2_v00.docx", SESSION.format(2),
    ]


def test_listing_takes_vice_chair_folders_and_the_latest_version_per_tdoc():
    listings = {
        "https://x/Inbox/": [{"name": "Chair_notes", "url": "https://x/Inbox/Chair_notes/"},
                             {"name": "Hiroki_notes", "url": "https://x/Inbox/Hiroki_notes/"},
                             {"name": "drafts", "url": "https://x/Inbox/drafts/"}],
    }
    files = {
        "https://x/Inbox/Hiroki_notes": [
            {"name": SESSION.format(1), "url": "u1", "uploaded_at": datetime(2026, 10, 12, 1)},
            {"name": SESSION.format(2), "url": "u2", "uploaded_at": datetime(2026, 10, 13, 1)},
            {"name": "R1-2607288_Session Notes of AI 10.5.4_final.docx", "url": "u3", "uploaded_at": datetime(2026, 10, 16, 1)},
            {"name": "RAN1#126bis schedule for Hiroki Adhoc2 sessions_v05.docx", "url": "u4", "uploaded_at": None},
        ],
    }
    with (patch.object(a, "list_inbox_subfolders", side_effect=listings.__getitem__),
          patch.object(a, "list_remote_files", side_effect=files.__getitem__) as listed):
        notes = a.list_session_notes({"inbox_urls": ["https://x/Inbox/"]})
    assert [n["name"] for n in notes] == ["R1-2607288_Session Notes of AI 10.5.4_final.docx"]
    assert [call.args[0] for call in listed.call_args_list] == ["https://x/Inbox/Hiroki_notes"]
    assert a.session_note_reference(notes[0])["uploaded_at"] == "2026-10-16T01:00:00"


def test_check_identity_adds_session_notes_only_when_there_are_some(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    chair = {"name": CHAIR.format(5), "url": "c", "sha256": "abc"}
    note = {"name": SESSION.format(1), "url": "u1", "uploaded_at": datetime(2026, 10, 13, 1)}
    with (patch.object(a, "remote_reference", return_value=chair),
          patch.object(a, "list_session_notes", return_value=[])):
        assert a.agreement_reference({}, MEETING) == (chair, chair)
    with (patch.object(a, "remote_reference", return_value=chair),
          patch.object(a, "list_session_notes", return_value=[note])):
        _, identity = a.agreement_reference({}, MEETING)
    assert identity == {**chair, "session_notes": [a.session_note_reference(note)]}


def test_remote_build_dates_a_session_note_by_its_upload_and_reports_the_listing(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    chair = chair_note(tmp_path / "in" / CHAIR.format(5))
    body = session_note(tmp_path / "in" / SESSION.format(1), ["Agreement:"]).read_bytes()
    info = {"name": SESSION.format(1), "url": "https://x/Inbox/Hiroki_notes/n.docx",
            "uploaded_at": datetime(2026, 10, 13, 6, 48), "folder": "Hiroki_notes"}
    manual = tmp_path / "ref_in_manual/ran1"
    manual.mkdir(parents=True)
    (manual / chair.name).write_bytes(chair.read_bytes())
    meta = {"last_modified": "Tue, 13 Oct 2026 06:48:12 GMT"}
    with (patch.object(a, "list_session_notes", return_value=[info]),
          patch.object(a.remote_files, "fetch_file", return_value=(body, meta))):
        result, identity = a.build_agreements(
            cfg={}, meeting_id=MEETING, schedule_path=None, offline=False,
            previous=None, now="2026-10-13T09:00:00+00:00",
        )
    assert identity["session_notes"] == [a.session_note_reference(info)]
    assert result["session_notes"]["R1-2607288"]["changed_at"] == "2026-10-13T06:48:12+00:00"
    assert (tmp_path / a.SESSION_NOTES_DIR / "Hiroki_notes" / SESSION.format(1)).exists()


def test_panel_manifest_names_the_notes_each_item_holds(tmp_path, monkeypatch):
    from tests.test_agreements import schedule

    monkeypatch.chdir(tmp_path)
    result, _ = offline_build(
        tmp_path, chair_note(tmp_path / "in" / CHAIR.format(5)),
        [session_note(tmp_path / "in" / SESSION.format(1), ["Agreement:"])],
        None, "2026-10-13T08:00:00+00:00",
    )
    stored, fragments = package_agreements(result)
    data = schedule(stored, meeting=MEETING)
    data.agreement_fragments = fragments
    manifest, _ = agreement_assets(data)
    assert manifest["session_notes"]["R1-2607288"]["name"] == SESSION.format(1)
    assert manifest["sections"]["10.5.4.1"]["notes"] == ["R1-2607288"]
    assert "notes" not in manifest["sections"]["10.5.4.2"]
    assert all("parts" not in s and "sources" not in s for s in manifest["sections"].values())


def test_a_session_note_that_appears_later_is_highlighted(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    src = tmp_path / "in"
    chair = chair_note(src / CHAIR.format(5))
    before, _ = offline_build(tmp_path, chair, [], None, "2026-10-12T08:00:00+00:00")
    before, _ = package_agreements(before)
    s1 = session_note(src / SESSION.format(0), ["Agreement:", "Support DCI format 1_X."])
    after, _ = offline_build(tmp_path, chair, [s1], before, "2026-10-13T08:00:00+00:00")
    section = after["sections"]["10.5.4.1"]
    assert section["change"] == "updated" and section["changed_in"] == SESSION.format(0)
    assert added(section) == ["Agreement:", "Support DCI format 1_X."]
