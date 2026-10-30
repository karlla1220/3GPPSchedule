"""SA2: finding the Session Plan, building the page, and the LLM fallback."""
import io
import json
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace
from urllib.parse import unquote
import zipfile

from docx import Document
import httpx
import pytest

from shared import officials
from working_groups.sa2 import document as docmod, lifecycle, llm, pipeline, sessions, sources

FIXTURES = Path(__file__).parent / 'fixtures/sa2'
HISTORY = (Path(__file__).parent / 'fixtures/officials/S2_officialsHistory.htm').read_text(encoding='utf-8')
R20 = 'S2-2608833_SA2-176-Session Plan - r20.docx'
R9 = 'S2-2601721_SA2-174-Session Plan -r9.docx'
META_176 = {'starts_on': '2026-08-24', 'ends_on': '2026-08-28', 'timezone': 'Europe/Prague'}
SYNC = sources.DEFAULTS['sync_url']
ARCHIVE = sources.DEFAULTS['archive_url']
FOLDER_176 = ARCHIVE + 'TSGS2_176_Prague_2026-08/'
FOLDER_177 = ARCHIVE + 'TSGS2_177_Prague_2026-10/'
ROWS = [{'Title': '3GPPSA2#176', 'TBId': 385, 'Id': 1, 'StartDate': '2026-08-24 09:00:00',
         'EndDate': '2026-08-28 16:30:00', 'StartTimeZone': '(GMT+02.00)  Belgrade, Bratislava, Budapest, Ljubljana, Prague',
         'Country': 'CZ', 'Location': 'Prague', 'MtgDocURL': 'https://ftp.3gpp.org//tsg_sa/WG2_Arch/TSGS2_176_Prague_2026-08/'},
        {'Title': '3GPPSA2#177', 'TBId': 385, 'Id': 2, 'StartDate': '2026-10-12 09:00:00',
         'EndDate': '2026-10-16 16:30:00', 'StartTimeZone': '(GMT+02.00)  Belgrade, Bratislava, Budapest, Ljubljana, Prague',
         'Country': 'CZ', 'Location': 'Prague', 'MtgDocURL': 'https://ftp.3gpp.org//tsg_sa/WG2_Arch/TSGS2_177_Prague_2026-10/'},
        {'Title': '3GPPSA2#178-AH-e', 'TBId': 385, 'Id': 3, 'StartDate': '2027-01-25 09:00:00',
         'EndDate': '2027-01-29 16:30:00', 'StartTimeZone': '', 'Country': '', 'Location': 'Online'}]


def listing(entries, base):
    rows = ''.join(f'<tr><td><a href="{base}{n}">{n}</a></td><td>{t}</td></tr>' for n, t in entries)
    return f'<table>{rows}</table>'


def mock_server(files, stamps=None, seen=None):
    """3GPP-like server: listings for folders (with upload times), 403 for missing paths."""
    stamps = stamps or {}

    def handler(request):
        url = unquote(str(request.url))
        if seen is not None:
            seen.append(url)
        if url.endswith('/'):
            names = sorted({u[len(url):].split('/')[0] + ('/' if '/' in u[len(url):] else '')
                            for u in files if u.startswith(url) and u != url})
            if not names:
                return httpx.Response(403)
            return httpx.Response(200, text=listing([(n, stamps.get(url + n, '2026/08/20 8:00')) for n in names], url))
        if url in files:
            return httpx.Response(200, content=files[url], headers={'etag': f'"{len(files[url])}"'})
        return httpx.Response(403)
    return httpx.Client(transport=httpx.MockTransport(handler))


def zipped(name, data):
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w') as archive:
        archive.writestr(name, data)
    return out.getvalue()


def candidate(meeting, version, uploaded, url='https://www.3gpp.org/ftp/tsg_sa/x'):
    return {'meeting_id': meeting, 'version': version, 'url': url, 'uploaded_at': uploaded, 'name': f'{meeting}-{version}'}


# ------------------------------------------------------------------ sources

def test_a_repeated_revision_is_ordered_by_upload_time():
    # "-r14.docx" (Thu 12:34) then "-r14 - Friday plan.docx" (Fri 07:48); #176's r15 was uploaded before r14.
    found = [candidate('sa2#177', 14, '2026-10-15 10:34'), candidate('sa2#177', 14, '2026-10-16 05:48'),
             candidate('sa2#177', 13, '2026-10-15 06:01')]
    assert sources.choose_schedule(found, {}, '2026-10-16')['uploaded_at'] == '2026-10-16 05:48'
    dates = {'sa2#177': ('2026-10-12', '2026-10-16'), 'sa2#178': ('2026-11-16', '2026-11-20')}
    early = found + [candidate('sa2#178', 0, '2026-11-01 09:00')]
    assert sources.choose_schedule(early, dates, '2026-10-14')['meeting_id'] == 'sa2#177'
    assert sources.choose_schedule(early, dates, '2026-10-17')['meeting_id'] == 'sa2#178'
    with pytest.raises(ValueError, match='No SA2 Session Plan'):
        sources.choose_schedule([], {}, '2026-10-17')


def test_a_pre_meeting_zip_holds_the_plan():
    data = (FIXTURES / R20).read_bytes()
    assert sources.unpack(zipped('S2-2608833_SA2-176-Session Plan - r20.docx', data), 'S2-x_SA2-176-Session Plan - r6.zip') \
        == (data, 'S2-x_SA2-176-Session Plan - r6.docx')
    assert sources.unpack(data, R20) == (data, R20)
    with pytest.raises(ValueError, match='0 DOCX'):
        sources.unpack(zipped('notes.txt', b'x'), 'S2-x_SA2-176-Session Plan - r6.zip')


def test_fetch_bundle_reads_the_sync_folder_and_the_archive_agenda(tmp_path, monkeypatch):
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')
    plan = (FIXTURES / R20).read_bytes()
    files = {SYNC + R20: plan,
             SYNC + 'S2-2608833_SA2-176-Session Plan - r19.docx': b'older',
             FOLDER_176 + 'INBOX/DRAFTS/_Session_Plan_Updates/Draft S2-2606837_SA2-176-Session Plan - r1.docx': b'draft',
             FOLDER_176 + 'Agenda/agenda.csv': (FIXTURES / 'agenda_176.csv').read_bytes()}
    seen = []
    with mock_server(files, seen=seen) as http:
        bundle = sources.fetch_bundle(sources.DEFAULTS, {}, http, rows=ROWS[:1], today='2026-08-27',
                                      cache_dir=tmp_path / 'inputs')
    assert (bundle.manifest['meeting_id'], bundle.manifest['version']) == ('sa2#176', 20)
    assert bundle.schedule == plan and bundle.manifest['agenda']['url'] == FOLDER_176 + 'Agenda/agenda.csv'
    assert bundle.manifest['portal_metadata']['timezone'] == 'Europe/Prague'
    # The folder came from Portal's MtgDocURL: the archive root is never listed.
    assert ARCHIVE not in seen


def test_before_the_meeting_the_archive_zip_is_read_and_an_absent_sync_folder_is_normal(tmp_path, monkeypatch):
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')
    monkeypatch.setattr(sources.remote_files.ftp_transport, 'ftp_response',
                        lambda *a, **k: pytest.fail('an absent folder is not read over FTP'))
    early = zipped(R20, (FIXTURES / R20).read_bytes())
    name = 'S2-2610195_SA2-177-Session Plan - r6.zip'
    files = {FOLDER_177 + 'INBOX/DRAFTS/_Session_Plan_Updates/' + name: early,
             FOLDER_177 + 'INBOX/DRAFTS/_Session_Plan_Updates/S2-260nnnn_SA2-177-Session Plan.zip': b'r0',
             FOLDER_177 + 'Agenda/S2-2609298_SA2#177 Meeting Agenda.zip': b'agenda docx',
             FOLDER_177 + 'Agenda/agenda.csv': b'"1.1","Opening of the meeting"\n'}
    with mock_server(files) as http:
        bundle = sources.fetch_bundle(sources.DEFAULTS, {}, http, rows=ROWS[1:], today='2026-10-05',
                                      cache_dir=tmp_path / 'inputs')
    assert (bundle.manifest['meeting_id'], bundle.manifest['version'], bundle.warnings) == ('sa2#177', 6, [])
    assert bundle.document()[1] == 'S2-2610195_SA2-177-Session Plan - r6.docx'
    # Nothing at all: the check says so.
    sources.remote_files.new_run()
    with mock_server({}) as http, pytest.raises(ValueError, match='No SA2 Session Plan'):
        sources.fetch_bundle(sources.DEFAULTS, {}, http, rows=ROWS[1:], today='2026-10-05', cache_dir=tmp_path / 'inputs')


def test_a_vanished_plan_is_kept_instead_of_stepping_back(tmp_path, monkeypatch):
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')
    friday = SYNC + 'S2-2610195_SA2-177-Session Plan -r14 - Friday plan.docx'
    thursday = SYNC + 'S2-2610195_SA2-177-Session Plan -r14.docx'
    stamps = {friday: '2026/10/16 5:48', thursday: '2026/10/15 10:34'}
    with mock_server({friday: b'friday', thursday: b'thursday'}, stamps) as http:
        first = sources.fetch_bundle(sources.DEFAULTS, {}, http, rows=ROWS[1:2], today='2026-10-16',
                                     cache_dir=tmp_path / 'inputs')
    assert first.schedule == b'friday'
    previous = {**first.manifest, 'metadata': {'starts_on': '2026-10-12', 'ends_on': '2026-10-16'}}
    sources.remote_files.new_run()
    with mock_server({thursday: b'thursday'}, stamps) as http:
        kept = sources.fetch_bundle(sources.DEFAULTS, previous, http, rows=ROWS[1:2], today='2026-10-16',
                                    cache_dir=tmp_path / 'inputs')
    assert kept.schedule == b'friday'


def test_config_rejects_unknown_keys_and_bad_folders(tmp_path, monkeypatch):
    path = tmp_path / 'config.json'
    monkeypatch.setattr(sources, 'CONFIG_PATH', path)
    path.write_text(json.dumps({'chair_url': 'x'}))
    with pytest.raises(ValueError, match='Unknown'):
        sources.load_config()
    path.write_text(json.dumps({'meeting_specific_override': [ARCHIVE + 'Inbox/']}))
    with pytest.raises(ValueError, match='meeting_specific'):
        sources.load_config()
    path.write_text(json.dumps({'meeting_specific_override': [FOLDER_177], 'meetings': {'178': {'timezone': 'America/Edmonton'}}}))
    cfg = sources.load_config()
    assert sources.override_for(cfg, 'sa2#178') == {'timezone': 'America/Edmonton'}
    assert sources.config_hash(cfg) == sources.config_hash({**cfg, 'meeting_specific_override': []})


# ----------------------------------------------------------- pipeline and CI

@pytest.fixture
def offline_config(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = {**sources.DEFAULTS, 'local_agenda': str(FIXTURES / 'agenda_176.csv'), 'meetings': {'176': dict(META_176)}}
    monkeypatch.setattr(pipeline, 'load_config', lambda: cfg)
    monkeypatch.setattr(sources, 'load_config', lambda: cfg)
    monkeypatch.setattr(pipeline, 'portal_rows', lambda: pytest.fail('Portal queried offline'))
    monkeypatch.setattr(sources, 'portal_rows', lambda: [])
    monkeypatch.setattr(pipeline, 'has_key', lambda: False)   # never reach Gemini from tests
    monkeypatch.setattr(officials, 'fetch_history', lambda code: pytest.fail('officials fetched offline'))
    return cfg


def test_local_build_needs_no_network_and_records_state(offline_config, tmp_path, monkeypatch):
    monkeypatch.setattr(httpx.Client, 'send', lambda *a, **k: pytest.fail('Offline network request'))
    schedule = pipeline.build_schedule(SimpleNamespace(local=str(FIXTURES / R20), output_dir=tmp_path / 'out'))
    assert (schedule.wg_id, schedule.meeting_name) == ('sa2', 'SA2#176')
    assert schedule.source_labels == {R20: 'Session Plan - r20', 'agenda.csv': 'agenda.csv'}
    state = sources.read_json(tmp_path / 'out/sa2/.schedule_state.json')
    assert state['meeting_id'] == 'sa2#176' and state['llm'] == {'used': False}
    # Offline the convener names stay as written; nothing is guessed.
    assert state['chairs'] == {'Mike': 'Mike', 'Wanqiang': 'Wanqiang'}
    again = pipeline.build_schedule(SimpleNamespace(no_download=True, output_dir=tmp_path / 'out'))
    assert again.days == schedule.days


def test_convener_names_come_from_the_officials_in_office():
    current = officials.in_office(officials.parse_history(HISTORY), officials.date(2026, 8, 24))
    names = {label: officials.resolve(label, current, sources.DEFAULTS['chair_aliases']) for label in ('Mike', 'Wanqiang')}
    assert names == {'Mike': 'Michael STARSINIC', 'Wanqiang': 'Wanqiang ZHANG'}
    # #174's Convener 1 had stepped down as vice chair; his name stays as written.
    april = officials.in_office(officials.parse_history(HISTORY), officials.date(2026, 4, 13))
    assert pipeline.chair_names(['Dario'], sources.DEFAULTS, 'sa2#174', '2026-04-13', {}, True) == {'Dario': 'Dario'}
    assert officials.resolve('Dario', april, sources.DEFAULTS['chair_aliases']) is None


def test_online_build_resolves_conveners(offline_config, tmp_path, monkeypatch):
    monkeypatch.setattr(officials, 'fetch_history', lambda code: officials.parse_history(HISTORY))
    bundle = sources.local_bundle(FIXTURES / R20, offline_config)
    bundle.manifest['schedule']['uploaded_at'] = '2026-08-28 05:48'
    bundle.save(sources.DOWNLOADS / 'prepared')
    monkeypatch.setattr(pipeline, 'portal_rows', lambda: [])
    schedule = pipeline.build_schedule(SimpleNamespace(output_dir=tmp_path / 'out'))
    assert schedule.source_labels[R20] == 'Session Plan - r20'   # the file name, not the upload time
    chairs = {s.chair for d in schedule.days for s in d.sessions if s.room_ids == ['stream2']}
    assert chairs == {'Michael STARSINIC'}


def test_lifecycle_check_stages_without_advancing_state(offline_config, tmp_path, monkeypatch):
    bundle = sources.local_bundle(FIXTURES / R20, offline_config)
    monkeypatch.setattr(sources, 'fetch_bundle', lambda *a, **k: bundle)
    before = {'meeting_id': 'sa2#175'}
    sources.write_json(sources.OUTPUT / '.schedule_state.json', before)
    result = lifecycle.check_updates()
    assert result.changed and 'meeting_id changed' in result.reasons and not result.errors
    assert sources.read_json(sources.OUTPUT / '.schedule_state.json') == before
    lifecycle.prepare_build()
    assert sources.Bundle.load(sources.DOWNLOADS / 'prepared').schedule == bundle.schedule
    sources.write_json(sources.OUTPUT / '.schedule_state.json', {**bundle.manifest, 'metadata': META_176})
    assert not lifecycle.check_updates().changed
    monkeypatch.setattr(sources, 'fetch_bundle', lambda *a, **k: (_ for _ in ()).throw(ValueError('down')))
    assert lifecycle.check_updates().errors == ['SA2 source check failed: down']


def test_reset_keeps_the_last_successful_output(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    sources.write_json(sources.OUTPUT / '.schedule_state.json', {'meeting_id': 'sa2#177'})
    sources.write_json(sources.DOWNLOADS / 'inputs/manifest.json', {})
    lifecycle.reset_cache()
    assert (sources.OUTPUT / '.schedule_state.json').exists() and not sources.DOWNLOADS.exists()


def test_lifecycle_import_is_isolated():
    result = subprocess.run([sys.executable, '-c', """
import sys
from working_groups.registry import get_working_group
get_working_group('sa2')
assert not any(name.startswith(('working_groups.ran', 'google.genai', 'docx')) for name in sys.modules)
"""], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


# --------------------------------------------------------------- LLM fallback

class FakeGemini:
    def __init__(self, *responses):
        self.responses, self.calls = list(responses), []

    def __call__(self, payload, schema, model, thinking=None):
        self.calls.append(payload)
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response


def run_together():
    """#174 r9, Monday Q5 Stream 1, left: "TEI20_DAMP [1]TEI20_UDCE-ARC [2]" runs two topics together."""
    document = docmod.extract_document((FIXTURES / R9).read_bytes(), R9)
    return document, next(c for c in document['cells'] if c['day'] == 'Monday' and c['quarters'] == [5]
                          and c['streams'] == [1] and c['halves'] == [1])


def test_written_times_and_shorthand_agenda_items():
    assert llm.written_times('Plenary session (1330 - 1630)\n14:00 List\nLS’s, 9779, 10337\nDrafting: (1300)') == \
        {'13:30', '16:30', '14:00', '13:00'}
    assert llm.item_errors('c', ['20.6.13', '19.49'], 'KI#12 – KI#16 (20.6.12 - 16)\nRel-19 (19.48/49/50)', {}) == []
    assert llm.item_errors('c', ['20.6.25'], 'KI#12 – KI#16 (20.6.12 - 16)', {}) == [
        "c: agenda item '20.6.25' is not written in its cited text"]


def test_checked_llm_answers_replace_the_rules_for_uncertain_cells(tmp_path):
    document, cell = run_together()
    rules = {s['id']: s for s in sessions.interpret(document, {})[0]}
    doubtful = [s for s in rules.values() if s['reasons']]
    assert [s['cell'] for s in doubtful] == [cell['id']]
    item = doubtful[0]
    title = 'TEI20 _NetShare_Ph2-ARC / TEI20_DAMP / TEI20_UDCE-ARC …'
    gemini = FakeGemini({'cells': [{'id': item['id'], 'start': '18:00', 'end': '19:30', 'title': title,
                                    'agenda_items': ['20.10.1']}]})
    refine = llm.CellRefiner('test-model', gemini, tmp_path)
    found = {s['id']: s for s in sessions.interpret(document, {}, refine)[0]}
    assert (found[item['id']]['name'], found[item['id']]['part']) == (title, [1, 2])
    assert refine.summary == {'model': 'test-model', 'uncertain': 1, 'applied': 1, 'kept_rules': 0}
    assert '"part": [1, 2]' in gemini.calls[0] and 'two topics run together' in gemini.calls[0]


@pytest.mark.parametrize('change, problem', [
    ({'title': 'Two TEI items'}, 'title words'),
    ({'start': '18:45'}, 'not written in the source'),
    ({'agenda_items': ['20.15']}, 'agenda item'),
])
def test_unchecked_answers_are_retried_once_then_the_rules_stand(tmp_path, change, problem):
    document, _ = run_together()
    item = next(s for s in sessions.interpret(document, {})[0] if s['reasons'])
    bad = {'id': item['id'], 'start': '18:00', 'end': '19:30', 'title': 'TEI20_DAMP', 'agenda_items': [], **change}
    gemini = FakeGemini({'cells': [bad]}, {'cells': [bad]})
    refine = llm.CellRefiner('test-model', gemini, tmp_path)
    found = next(s for s in sessions.interpret(document, {}, refine)[0] if s['id'] == item['id'])
    assert found['name'] == item['name'] and refine.summary['kept_rules'] == 1
    assert len(gemini.calls) == 2 and problem in gemini.calls[1]


def test_an_llm_outage_never_breaks_the_build(tmp_path):
    document, _ = run_together()
    refine = llm.CellRefiner('test-model', FakeGemini(RuntimeError('quota')), tmp_path)
    assert [s['name'] for s in sessions.interpret(document, {}, refine)[0]] == \
        [s['name'] for s in sessions.interpret(document, {})[0]]
    assert refine.summary['error'] == 'quota'


def slot_list():
    """#173's slot list kind of document: a table, but not the template."""
    doc = Document()
    doc.add_paragraph('SA2 #179 meeting session plan')
    table = doc.add_table(rows=3, cols=3)
    for row, cells in enumerate([('Monday', 'Main Room', 'Breakout 1'), ('0900 - 1030', 'KI#18 (AI for Network) (20.6.18)',
                                                                         'KI#22 (Computing) (20.6.22)'),
                                 ('1100 - 1230', 'KI#1.1 (NAS) (20.6.1)', '-')]):
        for column, text in enumerate(cells):
            table.cell(row, column).text = text
    data = io.BytesIO()
    doc.save(data)
    return data.getvalue()


def slot_answer(texts):
    ref = {text: key for key, text in texts.items()}
    session = lambda start, end, room, title, item, text: {
        'day': 'Monday', 'start': start, 'end': end, 'part': None, 'room_ids': [room], 'title': title,
        'agenda_items': [item], 'refs': [ref[text]]}
    return {'rooms': [{'id': 'main', 'name': 'Main'}, {'id': 'brk1', 'name': 'Breakout 1'}], 'days': ['Monday'],
            'sessions': [session('09:00', '10:30', 'main', 'KI#18', '20.6.18', 'KI#18 (AI for Network) (20.6.18)'),
                         session('09:00', '10:30', 'brk1', 'KI#22', '20.6.22', 'KI#22 (Computing) (20.6.22)'),
                         session('11:00', '12:30', 'main', 'KI#1.1', '20.6.1', 'KI#1.1 (NAS) (20.6.1)')]}


def test_a_plan_that_is_not_the_template_is_read_by_the_llm(offline_config, tmp_path, monkeypatch):
    data = slot_list()
    name = 'S2-2611111_SA2-179-Session Plan - r1.docx'
    with pytest.raises(docmod.LayoutError):
        docmod.extract_document(data, name)
    blocks = llm.outline(Document(io.BytesIO(data)))
    texts = {c['id']: c['text'] for b in blocks for c in b.get('cells', []) if c['text']}
    local = tmp_path / name
    local.write_bytes(data)
    offline_config['local_agenda'] = None
    offline_config['meetings'] = {'179': {'starts_on': '2027-02-22', 'ends_on': '2027-02-26', 'timezone': 'Asia/Seoul'}}
    monkeypatch.setattr(pipeline, 'CACHE', tmp_path / 'cache')
    monkeypatch.setattr(llm, 'CACHE', tmp_path / 'cache')
    gemini = FakeGemini(slot_answer(texts))
    schedule = pipeline.build_schedule(SimpleNamespace(local=str(local), output_dir=tmp_path / 'out'), request=gemini)
    monday = schedule.days[0]
    assert [(s.start_time, s.room_ids, s.name, s.agenda_item) for s in monday.sessions] == [
        ('09:00', ['main'], 'KI#18', '20.6.18'), ('09:00', ['brk1'], 'KI#22', '20.6.22'),
        ('11:00', ['main'], 'KI#1.1', '20.6.1')]
    assert schedule.supplements[0]['text'].startswith('The Session Plan table was not in its usual layout')
    assert sources.read_json(tmp_path / 'out/sa2/.schedule_state.json')['llm']['document'] is True


def test_an_unchecked_document_answer_fails_the_build(tmp_path):
    data = slot_list()
    blocks = llm.outline(Document(io.BytesIO(data)))
    texts = {c['id']: c['text'] for b in blocks for c in b.get('cells', []) if c['text']}
    invented = slot_answer(texts)
    invented['sessions'][0]['part'] = [1, 2]   # a piece that cites no cell of its own
    gemini = FakeGemini(invented, invented)
    with pytest.raises(docmod.LayoutError, match='could not read the Session Plan either'):
        llm.read_document(data, 'S2-2611111_SA2-179-Session Plan - r1.docx', 'test-model', None, gemini, tmp_path)
    assert len(gemini.calls) == 2 and 'their own table cell' in gemini.calls[1]
