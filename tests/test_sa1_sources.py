"""SA1 sources, pipeline and CI: choosing the chair file version to build from."""
import io
import json
import subprocess
import sys
from types import SimpleNamespace
from urllib.parse import quote, unquote
import zipfile

from docx import Document
import httpx
import pytest

from test_sa1 import EOM, FIXTURES, HISTORY, META_115, M113, TIMETABLE, find
from working_groups.sa1 import lifecycle, pipeline, sources


# ----------------------------------------------------------------- sources

@pytest.mark.parametrize('name, kind, order, label, meeting', [
    (EOM, 'chair', [3, 4, 950], "Chair's Notes EOM", 'sa1#115'),
    ('Agenda for SA1#115_Thursday_1030.docx', 'chair', [2, 3, 630], "Chair's Notes Thu 10:30", 'sa1#115'),
    ('S1-263002_Agenda for SA1#115_Monday_10_30.docx', 'chair', [2, 0, 630], "Chair's Notes Mon 10:30", 'sa1#115'),
    ('S1-263002_Agenda for SA1#115.docx', 'chair', [2, -1, 0], "Chair's Notes", 'sa1#115'),
    (M113, 'chair', [2, 4, 960], "Chair's Notes Fri 16:00", 'sa1#113'),
    ('SA1#116_ChairNotes.docx', 'chair', [2, -1, 0], "Chair's Notes", 'sa1#116'),
    (TIMETABLE, 'timetable', [2, 2, 1050], 'Breakout schedule Wed 17:30', None),
    ('S1-264001.zip', 'chair', [0, 0, 0], "Chair's Notes S1-264001", None),
    ('Agenda_AI_SA1#115_WED_10_30.docx', None, None, None, None),
    ('AgendaDrafting_Massive_SA1#115 Wednesday 1800.docx', None, None, None, None),
    ('doc_list_S1_115_Aug_28_at_12_10.docm', None, None, None, None),
])
def test_file_names_give_kind_order_and_label(name, kind, order, label, meeting):
    found = sources.classify({'name': name, 'url': 'u', 'uploaded_at': ''})
    if kind is None:
        assert found is None
    else:
        assert (found['kind'], found['order'], found['label'], found['meeting_id']) == (kind, order, label, meeting)


SYNC = 'https://www.3gpp.org/ftp/Meetings_3GPP_SYNC/SA1/Inbox/'
ARCHIVE = 'https://www.3gpp.org/ftp/tsg_sa/WG1_Serv/'
ROWS = [{'Title': '3GPPSA1#115', 'TBId': 384, 'Id': 1, 'StartDate': '2026-08-24 09:00:00',
         'EndDate': '2026-08-28 17:30:00', 'Location': 'Prague', 'Country': 'CZ',
         'StartTimeZone': '(GMT+02.00)  Belgrade, Bratislava, Budapest, Ljubljana, Prague'},
        {'Title': '3GPPSA1#116', 'TBId': 384, 'Id': 2, 'StartDate': '2026-11-16 09:00:00',
         'EndDate': '2026-11-20 17:30:00', 'Location': 'Calgary', 'Country': 'CA',
         'StartTimeZone': '(GMT-07:00) Mountain Time (US & Canada)'}]


def mock_server(files):
    """3GPP-like server: listings with upload times, 403 for missing paths."""
    def handler(request):
        url = unquote(str(request.url))
        if url.endswith('/'):
            names = sorted({u[len(url):].split('/')[0] + ('/' if '/' in u[len(url):] else '')
                            for u in files if u.startswith(url) and u != url})
            if not names:
                return httpx.Response(403)
            rows = ''.join(f'<tr><td><a href="{quote(url + n, safe=":/")}">{n}</a></td>'
                           f'<td>{files.get(url + n, (None, "2026/08/26 15:22"))[1]}'
                           '</td></tr>' for n in names)
            return httpx.Response(200, text=f'<table>{rows}</table>')
        if url in files:
            return httpx.Response(200, content=files[url][0], headers={'etag': f'"{len(files[url][0])}"'})
        return httpx.Response(403)
    return httpx.Client(transport=httpx.MockTransport(handler))


def blank_docx(title='3GPP TSG-SA WG1 Meeting #116'):
    buffer = io.BytesIO()
    doc = Document()
    doc.add_paragraph(title)
    doc.save(buffer)
    return buffer.getvalue()


def renumbered(data, meeting='#116'):
    """The fixture as if it were another meeting's chair file (its title paragraph)."""
    doc = Document(io.BytesIO(data))
    title = doc.paragraphs[0]
    text = title.text.replace('#115', meeting)
    for i, run in enumerate(title.runs):
        run.text = text if i == 0 else ''
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def zipped(name, data):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as archive:
        archive.writestr(name, data)
    return buffer.getvalue()


@pytest.fixture
def server_cache(tmp_path, monkeypatch):
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')


def fetch(files, today, previous=None, tmp_path=None):
    with mock_server(files) as http:
        return sources.fetch_bundle(sources.DEFAULTS, previous or {}, http, rows=ROWS, today=today,
                                    cache_dir=tmp_path / 'inputs')


def test_a_newer_timetable_only_file_wins_and_the_chair_file_gives_the_agenda(server_cache, tmp_path):
    eom = (FIXTURES / EOM).read_bytes()
    files = {SYNC + 'agenda/previous_ones/S1-263002_Agenda for SA1#115_Tuesday_09_00.docx': (eom, '2026/08/28 10:18'),
             SYNC + 'agenda/previous_ones/' + TIMETABLE: ((FIXTURES / TIMETABLE).read_bytes(), '2026/08/26 15:22')}
    bundle = fetch(files, '2026-08-26', tmp_path=tmp_path)
    assert (bundle.manifest['meeting_id'], bundle.manifest['version']) == ('sa1#115', 'Breakout schedule Wed 17:30')
    assert bundle.manifest['chair']['name'] == 'S1-263002_Agenda for SA1#115_Tuesday_09_00.docx'
    assert bundle.manifest['portal_metadata']['timezone'] == 'Europe/Prague'
    # A later chair file wins again, and needs no second file.
    sources.remote_files.new_run()
    files[SYNC + 'agenda/Agenda for SA1#115_Thursday_0800.docx'] = (eom, '2026/08/27 05:50')
    bundle = fetch(files, '2026-08-27', tmp_path=tmp_path)
    assert (bundle.manifest['version'], bundle.manifest['chair']) == ("Chair's Notes Thu 08:00", None)


def test_a_version_without_a_schedule_table_is_skipped(server_cache, tmp_path, capsys):
    eom = (FIXTURES / EOM).read_bytes()
    files = {SYNC + 'agenda/Agenda for SA1#115_Thursday_0800.docx': (eom, '2026/08/27 05:50'),
             ARCHIVE + 'TSGS1_115_Prague/temp_reports/Agenda for SA1#115_END_OF_MEETING.docx':
                 (blank_docx('3GPP TSG-SA WG1 Meeting #115'), '2026/08/31 13:53')}
    previous = {'meeting_id': 'sa1#115', 'metadata': META_115}
    bundle = fetch(files, '2026-09-01', previous, tmp_path)
    assert bundle.manifest['version'] == "Chair's Notes Thu 08:00"
    assert 'has no schedule table' in capsys.readouterr().out


def test_agenda_tdoc_zips_are_opened_for_the_next_meeting_only(server_cache, tmp_path):
    eom = (FIXTURES / EOM).read_bytes()
    files = {SYNC + 'agenda/' + EOM: (eom, '2026/08/28 13:52'),
             SYNC + 'S1-263001.zip': (zipped('S1-263001_2nd Draft Agenda for SA1#115_v11.docx', eom), '2026/08/23 15:15')}
    # Before #116 has any file, the zips tell which meeting they are; #115 keeps its closing file.
    assert fetch(files, '2026-11-01', tmp_path=tmp_path).manifest['version'] == "Chair's Notes EOM"
    sources.remote_files.new_run()
    files[SYNC + 'S1-264001.zip'] = (zipped('S1-264001_2nd Draft Agenda for SA1#116.docx', renumbered(eom)),
                                     '2026/11/15 15:15')
    bundle = fetch(files, '2026-11-15', tmp_path=tmp_path)
    assert (bundle.manifest['meeting_id'], bundle.manifest['version']) == ('sa1#116', "Chair's Notes S1-264001")
    assert bundle.documents()[0].startswith(b'PK')   # the DOCX inside the zip


def test_bundle_cache_rejects_corruption(tmp_path):
    bundle = sources.local_bundle(FIXTURES / EOM, sources.DEFAULTS)
    bundle.save(tmp_path)
    (tmp_path / 'schedule.bin').write_bytes(b'changed')
    with pytest.raises(ValueError, match='Corrupt'):
        sources.Bundle.load(tmp_path)


def test_config_rejects_unknown_keys_and_bad_folders(tmp_path, monkeypatch):
    path = tmp_path / 'config.json'
    monkeypatch.setattr(sources, 'CONFIG_PATH', path)
    for bad in ({'extra': 1}, {'meeting_specific_override': ['https://www.3gpp.org/ftp/tsg_ran/WG2_RL2/TSGR2_135/']},
                {'meetings': {'x': {}}}, {'officials_code': ''}):
        path.write_text(json.dumps(bad))
        with pytest.raises(ValueError):
            sources.load_config()
    path.write_text(json.dumps({'meeting_specific_override': [ARCHIVE + 'TSGS1_116_Calgary/']}))
    assert sources.load_config()['meeting_specific_override'] == [ARCHIVE + 'TSGS1_116_Calgary/']


# ----------------------------------------------------------- pipeline and CI

@pytest.fixture
def offline_config(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = {**sources.DEFAULTS, 'meetings': {'115': dict(META_115)}}
    monkeypatch.setattr(pipeline, 'load_config', lambda: cfg)
    monkeypatch.setattr(sources, 'load_config', lambda: cfg)
    monkeypatch.setattr(pipeline, 'portal_rows', lambda: pytest.fail('Portal queried offline'))
    monkeypatch.setattr(sources, 'portal_rows', lambda: [])
    monkeypatch.setattr(pipeline.llm_cells, 'has_key', lambda: False)   # never reach Gemini from tests
    return cfg


def test_local_build_needs_no_network_and_records_state(offline_config, tmp_path, monkeypatch):
    monkeypatch.setattr(httpx.Client, 'send', lambda *a, **k: pytest.fail('Offline network request'))
    options = SimpleNamespace(local=str(FIXTURES / EOM), no_download=True, output_dir=tmp_path / 'out')
    schedule = pipeline.build_schedule(options)
    assert schedule.wg_id == 'sa1' and schedule.meeting_name == 'SA1#115'
    assert schedule.source_labels == {EOM: "Chair's Notes EOM"}
    # Offline, the NOTE names are shown as written.
    assert find(schedule, 'Tuesday', '09:00', 'breakout-2').chair == 'Qun Wei'
    state = sources.read_json(tmp_path / 'out/sa1/.schedule_state.json')
    assert state['meeting_id'] == 'sa1#115' and state['llm'] == {'used': False}
    again = pipeline.build_schedule(SimpleNamespace(no_download=True, output_dir=tmp_path / 'out'))
    assert again.days == schedule.days


def test_officials_spell_the_note_names_and_their_failure_is_harmless(offline_config, tmp_path, monkeypatch):
    options = SimpleNamespace(local=str(FIXTURES / EOM), output_dir=tmp_path / 'out')
    monkeypatch.setattr(pipeline, 'portal_rows', lambda: [])
    schedule = pipeline.build_schedule(options, history=HISTORY)
    assert find(schedule, 'Tuesday', '09:00', 'breakout-2').chair == 'QUN WEI'
    monkeypatch.setattr(pipeline.officials, 'fetch_history', lambda code: (_ for _ in ()).throw(ValueError('down')))
    assert pipeline.chair_resolver(offline_config, '2026-08-24') is None


def test_a_new_meeting_without_dates_fails_clearly(offline_config, tmp_path, monkeypatch):
    offline_config['meetings'] = {}
    monkeypatch.setattr(pipeline, 'portal_rows', lambda: None)
    with pytest.raises(ValueError, match='Add them under "meetings"'):
        pipeline.build_schedule(SimpleNamespace(local=str(FIXTURES / EOM), output_dir=tmp_path / 'out'))
    assert not (tmp_path / 'out/sa1/.schedule_state.json').exists()


def test_lifecycle_check_stages_without_advancing_state(offline_config, tmp_path, monkeypatch):
    bundle = sources.local_bundle(FIXTURES / EOM, offline_config)
    monkeypatch.setattr(sources, 'fetch_bundle', lambda *a, **k: bundle)
    before = {'meeting_id': 'sa1#114'}
    sources.write_json(sources.OUTPUT / '.schedule_state.json', before)
    result = lifecycle.check_updates()
    assert result.changed and 'meeting_id changed' in result.reasons and not result.errors
    assert sources.read_json(sources.OUTPUT / '.schedule_state.json') == before
    lifecycle.prepare_build()
    assert sources.Bundle.load(sources.DOWNLOADS / 'prepared').schedule == bundle.schedule
    sources.write_json(sources.OUTPUT / '.schedule_state.json', {**bundle.manifest, 'metadata': META_115})
    assert not lifecycle.check_updates().changed
    monkeypatch.setattr(sources, 'fetch_bundle', lambda *a, **k: (_ for _ in ()).throw(ValueError('down')))
    assert lifecycle.check_updates().errors == ['SA1 source check failed: down']


def test_reset_keeps_the_last_successful_output(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    sources.write_json(sources.OUTPUT / '.schedule_state.json', {'meeting_id': 'sa1#115'})
    sources.write_json(sources.DOWNLOADS / 'inputs/manifest.json', {})
    lifecycle.reset_cache()
    assert (sources.OUTPUT / '.schedule_state.json').exists() and not sources.DOWNLOADS.exists()


def test_lifecycle_import_is_isolated():
    result = subprocess.run([sys.executable, '-c', """
import sys
from working_groups.registry import get_working_group
get_working_group('sa1')
assert not any(name.startswith(('working_groups.ran', 'google.genai', 'docx')) for name in sys.modules)
"""], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
