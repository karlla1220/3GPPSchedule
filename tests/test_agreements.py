"""Real OOXML extraction, isolation, update detection and static artifact contracts."""

import base64
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor, Inches
import pytest

from shared.agreement_assets import agreement_assets, write_agreement_assets, agenda_section_ids
from shared.renderer import generate_html
from shared.schedule import (
    Schedule,
    DaySchedule,
    RoomInfo,
    Session,
    save_schedule,
    load_schedule,
)
from working_groups.ran1 import agreements as a

REAL_NOTE = Path("tests/fixtures/ran1/Chair notes RAN1#124 - v09.docx")


@pytest.mark.parametrize(('value', 'expected'), [
    ('10.4.x', ['10.4.1', '10.4.2', '10.4.10']),
    # A parent brings its subsections; 10.40.1 shares only a numeric prefix.
    ('10.4', ['10.4', '10.4.1', '10.4.2', '10.4.10']),
    ('10.4.1, 10.4', ['10.4.1', '10.4', '10.4.2', '10.4.10']),
    ('10.8.1., 8.1 AIML, 8.1 NES', ['10.8.1', '8.1']),
    ('10.5.3.1/4', ['10.5.3.1', '10.5.3.4']),
    ('6/8.1', ['6', '8.1']),
    ('9.3.2, 99.x', ['9.3.2', '99.x']),
    ('10.4.1', ['10.4.1']),
])
def test_explicit_agenda_notation(value, expected):
    sections = dict.fromkeys(['10.4', '10.4.1', '10.4.2', '10.40.1', '10.4.10',
                             '8.1', '10.5.3.1', '10.5.3.4'])
    assert agenda_section_ids(value, sections) == expected


def test_heading_style_depth_does_not_override_numbering_depth(tmp_path):
    document = Document()
    document.add_heading('9.3 Ambient IoT', level=2)
    numbering = document.part.numbering_part.element
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), '100')
    for index in range(3):
        level = OxmlElement('w:lvl')
        level.set(qn('w:ilvl'), str(index))
        for tag, value in [('start', '1'), ('numFmt', 'decimal'),
                           ('lvlText', '.'.join('%' + str(i + 1) for i in range(index + 1)))]:
            child = OxmlElement('w:' + tag)
            child.set(qn('w:val'), value)
            level.append(child)
        abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), '100')
    child = OxmlElement('w:abstractNumId')
    child.set(qn('w:val'), '100')
    num.append(child)
    numbering.append(num)
    for title in ['R2D', 'D2R', 'Other procedures']:
        paragraph = document.add_heading(title, level=4)
        props = paragraph._p.get_or_add_pPr()
        num_props = OxmlElement('w:numPr')
        for tag, value in [('ilvl', '2'), ('numId', '100')]:
            child = OxmlElement('w:' + tag)
            child.set(qn('w:val'), value)
            num_props.append(child)
        props.append(num_props)
        document.add_paragraph('Original content for ' + title)
    path = tmp_path / 'Chair notes RAN1#126_v15.docx'
    document.save(path)
    sections = a.parse_agreements(path, 'ran1#126')['sections']
    for ai, title in [('9.3.1', 'R2D'), ('9.3.2', 'D2R'), ('9.3.3', 'Other procedures')]:
        assert sections[ai]['title'] == title
        assert 'Original content for ' + title in sections[ai]['html']
    assert 'Original content for D2R' not in sections['9.3.1']['html']


@pytest.fixture(autouse=True)
def isolated_cache(tmp_path, monkeypatch):
    monkeypatch.setattr(a, "CACHE_DIR", tmp_path / "cache")


@pytest.fixture
def note(tmp_path):
    d = Document()
    d.add_heading("10.1 Evaluation", level=1)
    d.add_paragraph("Agreement")
    p = d.add_paragraph("Exact original <script>alert(1)</script> ")
    r = p.add_run("formatted")
    r.bold = True
    r.italic = True
    r.underline = WD_UNDERLINE.SINGLE
    r.font.color.rgb = RGBColor.from_string("FF0000")
    r.font.strike = True
    p.add_run("sub").font.subscript = True
    p.add_run("sup").font.superscript = True
    p.paragraph_format.left_indent = Inches(0.5)
    d.add_paragraph("First bullet", style="List Bullet")
    d.add_paragraph("First numbered", style="List Number")
    p = d.add_paragraph()
    for url, label in [
        ("javascript:alert(1)", "unsafe link"),
        ("https://example.org/source", "safe link"),
    ]:
        rel = d.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = label
        r.append(t)
        link.append(r)
        p._p.append(link)
    image = tmp_path / "pixel.png"
    image.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
        )
    )
    d.add_picture(str(image), width=Inches(0.3))
    table = d.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Horizontal merge"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Vertical merge"
    table.cell(1, 1).text = "Cell content"
    d.add_paragraph("")
    d.add_paragraph("R1-2600001\tContribution title\tExample company")
    d.add_paragraph("Discussion after TDoc")
    d.add_heading("10.10 Different agenda", level=1)
    d.add_paragraph("Agreement:")
    d.add_paragraph("Only agenda 10.10")
    d.add_heading("10.2 No agreement", level=1)
    d.add_paragraph("Mention of an agreement is not a marker.")
    d.add_heading("Unknown heading", level=1)
    d.add_paragraph("Agreement")
    d.add_paragraph("Unnumbered subsection content")
    d.add_heading("10.3 TDoc entries only", level=1)
    d.add_paragraph("R1-2600999")
    path = tmp_path / "Chair notes RAN1#124 - v01.docx"
    d.save(path)
    return path


def schedule(data=None, meeting="ran1#124", wg="ran1"):
    return Schedule(
        "Test",
        [
            DaySchedule(
                "Monday",
                [RoomInfo("Main")],
                [
                    Session(
                        "Multi",
                        60,
                        "09:00",
                        "10:00",
                        "Monday",
                        agenda_item="10.1, 10.10",
                    )
                ],
            )
        ],
        "schedule.docx",
        "2026-09-29",
        wg_id=wg,
        meeting_id=meeting,
        chairman_agreements=data or {},
    )


def test_exact_agenda_boundaries_and_original_text(note):
    result = a.parse_agreements(note, "ran1#124")
    s = result["sections"]
    html = s["10.1"]["html"]
    assert "Exact original" in html and "&lt;script&gt;" in html
    assert "Discussion after TDoc" in html and "Only agenda 10.10" not in html
    assert "Only agenda 10.10" in s["10.10"]["html"]
    assert "Mention of an agreement is not a marker." in s["10.2"]["html"]
    assert "Unnumbered subsection content" in s["10.2"]["html"]
    assert s["10.3"]["html"] == ""
    assert s["10.1"]["excluded_tdoc_rows"] == 1
    assert "R1-2600001" not in html


def test_run_list_table_image_and_link_fidelity(note):
    result = a.parse_agreements(note, "ran1#124")
    html = result["sections"]["10.1"]["html"]
    dom = BeautifulSoup(html, "html.parser")
    for css in (
        "font-weight:bold",
        "font-style:italic",
        "text-decoration:underline line-through",
        "color:#FF0000",
        "margin-left:36pt",
    ):
        assert css in html
    assert dom.find("sub").text == "sub" and dom.find("sup").text == "sup"
    assert dom.select('[colspan="2"]') and dom.select('[rowspan="2"]')
    assert dom.find("img")["src"].startswith("data:image/png;base64,")
    assert dom.find("a")["href"] == "https://example.org/source"
    assert "javascript:" not in html and not dom.find("script")
    assert dom.select("ul > li") and dom.select('ol > li[value="1"]')
    assert dom.get_text().count("Vertical merge") == 1


@pytest.mark.parametrize("meeting", ["ran1#124bis", "ran1#125", "", "ran1#12"])
def test_reject_other_or_unknown_meeting(note, meeting):
    with pytest.raises(ValueError, match="exact"):
        a.parse_agreements(note, meeting)


def test_real_chairman_note_numbering_and_counts():
    result = a.parse_agreements(REAL_NOTE, "ran1#124")
    sections = result["sections"]
    assert len(sections) == 90
    assert sections["9.2.1"]["title"] == "Improvement of SRS capacity and coverage"
    assert "9.10.1" not in sections
    assert sum(s["excluded_tdoc_rows"] for s in sections.values()) == 1539
    assert sections["10.5.0"]["title"] == "General aspects and frameworks"
    assert result["content_scope"] == "agenda-section-excluding-tdocs"
    html = sections["10.1"]["html"]
    assert "<table>" in html and "<math " in html
    assert "R1-2601415" not in BeautifulSoup(html, "html.parser").get_text()
    assert "data:image/svg+xml;base64," in html
    assert not any("could not be converted" in w for w in result["warnings"])


def test_same_filename_new_bytes_invalidates_cache(note):
    first = a.parse_agreements(note, "ran1#124")
    d = Document(note)
    d.paragraphs[2].add_run(" CONTENT UPDATE")
    d.save(note)
    second = a.parse_agreements(note, "ran1#124")
    assert first["sha256"] != second["sha256"]
    assert "CONTENT UPDATE" in second["sections"]["10.1"]["html"]


def test_docm_is_read_without_loading_macros(note, tmp_path):
    docm = tmp_path / "Chair notes RAN1#124 - v01.docm"
    docm.write_bytes(note.read_bytes())
    assert a.parse_agreements(docm, "ran1#124")["sections"]["10.1"]["html"]


def test_offline_missing_and_exact_local_note(note, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = {"inbox_urls": [], "extra_folders": []}
    with patch.object(a, "remote_reference", side_effect=AssertionError("network")):
        found, ref = a.build_agreements(
            cfg=cfg,
            meeting_id="ran1#124",
            schedule_path=tmp_path / "schedule.docx",
            offline=True,
        )
        missing, _ = a.build_agreements(
            cfg=cfg,
            meeting_id="ran1#124bis",
            schedule_path=tmp_path / "schedule.docx",
            offline=True,
        )
    assert found["status"] == "ready" and ref["origin"] == "local"
    assert missing["status"] == "unavailable"


def test_remote_revalidation_and_download_failure_do_not_succeed(tmp_path):
    cfg = {"inbox_urls": [], "extra_folders": []}
    info = {
        "name": "Chair notes RAN1#124 - v02.docx",
        "url": "https://example.org/note.docx",
    }
    with (
        patch.object(a, "get_latest_chair_notes_info", return_value=info) as listing,
        patch.object(
            a.remote_files, "fetch_file", return_value=(b"new", {"sha256": "NEW"})
        ) as fetch,
    ):
        assert a.remote_reference(cfg, "ran1#124")["sha256"] == "NEW"
        listing.assert_called_once_with(
            urls=[], extra_folders=[], preferred_meeting_id="ran1#124", strict=True
        )
        fetch.assert_called_once()
    with (
        patch.object(a, "remote_reference", return_value=info),
        patch.object(a, "download_latest_chair_notes", return_value=None),
    ):
        with pytest.raises(RuntimeError):
            a.build_agreements(
                cfg=cfg,
                meeting_id="ran1#124",
                schedule_path=tmp_path / "schedule.docx",
                offline=False,
            )


def test_asset_manifest_is_lazy_content_addressed_and_roundtrips(note, tmp_path):
    s = schedule(a.parse_agreements(note, "ran1#124"))
    save_schedule(s, tmp_path / "schedule.json")
    s = load_schedule(tmp_path / "schedule.json")
    page = generate_html(s)
    manifest, assets = agreement_assets(s)
    assert "Exact original" not in page
    assert "./agreements/" in page and "solid-js@1.9.9" in page
    assert 'aria-controls="agreement-panel"' in page
    assert "iframe" not in page
    assert "Agreements for" in page
    assert "dompurify@3.4.16" in page
    write_agreement_assets(s, tmp_path)
    for name, content in assets.items():
        assert (tmp_path / name).exists()
    assert "source_download" not in manifest
    assert not any(name.endswith((".docx", ".docm")) for name in assets)
    html = (tmp_path / manifest["sections"]["10.1"]["url"]).read_text()
    assert html.startswith('<div class="docx-document">')
    assert "<html" not in html and "<style" not in html
    assert "<article>" in html
    assert ":host" in manifest["document_css"]
    assert manifest["sections"]["10.2"].get("url")
    assert not manifest["sections"]["10.3"].get("url")


def test_wrong_meeting_payload_and_other_wg_never_show_agreements(note):
    data = a.parse_agreements(note, "ran1#124")
    manifest, assets = agreement_assets(schedule(data, meeting="ran1#124bis"))
    assert manifest["status"] == "unavailable" and not assets
    assert (
        BeautifulSoup(
            generate_html(schedule(data, wg="ran-plenary")), "html.parser"
        ).find(id="agreement-panel")
        is None
    )


def test_portal_timezone_does_not_suppress_agreement_check():
    from test_check_update import _run_check

    state = {
        "files": [],
        "meeting_id": "ran1#124",
        "timezone": "Europe/Malta",
        "timezone_ref": {"type": "portal"},
        "agreements_ref": {"name": "same", "url": "https://x/note", "sha256": "OLD"},
    }
    current = {"name": "same", "url": "https://x/note", "sha256": "NEW"}
    with (
        patch(
            "working_groups.ran1.check_update.lookup_timezone_reference",
            return_value=None,
        ),
        patch.object(a, "remote_reference", return_value=current) as probe,
    ):
        outputs, _ = _run_check(state=state, local_refs={}, remote=[])
    assert ("changed", "true") in outputs
    probe.assert_called_once()
    state["agreements_ref"] = current
    with (
        patch(
            "working_groups.ran1.check_update.lookup_timezone_reference",
            return_value=None,
        ),
        patch.object(a, "remote_reference", return_value=current),
    ):
        outputs, _ = _run_check(state=state, local_refs={}, remote=[])
    assert ("changed", "false") in outputs


def test_listing_error_is_not_document_absence():
    from working_groups.ran1.downloader import get_latest_chair_notes_info

    with patch(
        "working_groups.ran1.downloader.list_remote_files",
        side_effect=RuntimeError("offline"),
    ):
        with pytest.raises(RuntimeError):
            get_latest_chair_notes_info(
                urls=["https://example.org"],
                strict=True,
                preferred_meeting_id="ran1#124",
            )


def test_real_drawing_relationships_have_no_raster_fallback():
    from lxml import etree as ET
    from shared.docx_html import NS

    with ZipFile(REAL_NOTE) as z:
        root = ET.fromstring(z.read("word/document.xml"))
        rels = {
            r.get("Id"): r.get("Target")
            for r in ET.fromstring(z.read("word/_rels/document.xml.rels"))
        }
        refs = [e.get("{%s}embed" % NS["r"]) for e in root.findall(".//a:blip", NS)]
        assert [rels[r] for r in refs] == [
            "media/image1.wmf",
            "media/image2.wmf",
            "media/image3.wmf",
        ]
        assert not [
            n for n in z.namelist() if n.lower().endswith((".png", ".jpg", ".jpeg"))
        ]
        assert not [n for n in z.namelist() if "/embeddings/" in n]
        assert len(root.findall(".//m:oMath", NS)) == 32
        assert not root.xpath(
            '//*[local-name()="AlternateContent" or local-name()="Fallback"]'
        )


def test_all_real_omml_equations_convert_independently():
    from shared.docx_html import DocxHTML, NS
    from shared.docx_math import to_mathml

    doc = DocxHTML(REAL_NOTE)
    converted = [to_mathml(m) for m in doc.body.findall(".//m:oMath", NS)]
    assert len(converted) == 32
    assert all("<math" in m for m in converted)
    assert any("<mfrac>" in m for m in converted)
    assert any("<msubsup>" in m for m in converted)


def test_metafile_symbol_mapping_and_active_svg_rejection(monkeypatch):
    from shared import wmf2svg

    monkeypatch.setattr(wmf2svg, "available", lambda: False)
    from shared.metafile_images import render_metafile_svg, sanitize_svg
    from lxml import etree as ET

    with ZipFile(REAL_NOTE) as z:
        for i in range(1, 4):
            svg, warnings = render_metafile_svg(z.read(f"word/media/image{i}.wmf"))
            root = ET.fromstring(svg)
            texts = "".join(root.itertext())
            assert any(char in texts for char in "θφϕ")
            assert "metadata" not in svg.decode()
            assert 'font-family="Symbol"' not in svg.decode()
            assert not any("partially" in w for w in warnings)
    with pytest.raises(ValueError):
        sanitize_svg(
            b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
        )
    result = sanitize_svg(
        b'<svg xmlns="http://www.w3.org/2000/svg"><image href="https://tracker/x" onload="evil()"/></svg>'
    )
    assert b"https://" not in result and b"onload" not in result


def test_manual_note_removal_returns_to_remote_detection(note, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    manual = tmp_path / "ref_in_manual/ran1"
    manual.mkdir(parents=True)
    target = manual / note.name
    target.write_bytes(note.read_bytes())
    cfg = {"extra_files": []}
    assert a.local_note_reference(cfg, "ran1#124")["origin"] == "local"
    target.unlink()
    assert a.local_note_reference(cfg, "ran1#124") is None
    assert a.local_note_reference(cfg, "ran1#124bis") is None


def test_wmf2svg_real_headless_conversion():
    from shared import wmf2svg
    from shared.metafile_images import render_metafile_svg

    if not wmf2svg.available():
        pytest.skip(
            "Optional Java renderer: run scripts/setup_wmf2svg.py with Java installed"
        )
    with ZipFile(REAL_NOTE) as z:
        for i in range(1, 4):
            svg, notices = render_metafile_svg(z.read(f"word/media/image{i}.wmf"))
            assert any("wmf2svg 0.10.6" in n for n in notices)
            assert not any("failed" in n for n in notices)
            assert b"<text" in svg and b"<style" not in svg and b"<!DOCTYPE" not in svg


def test_java_zero_exit_missing_output_and_hash_are_failures(tmp_path, monkeypatch):
    import hashlib
    import subprocess
    from shared import wmf2svg

    jar = tmp_path / "renderer.jar"
    jar.write_bytes(b"test jar")
    monkeypatch.setattr(wmf2svg, "JAR_PATH", jar)
    with pytest.raises(ValueError, match="checksum"):
        wmf2svg.convert(b"test")
    monkeypatch.setattr(wmf2svg, "SHA256", hashlib.sha256(jar.read_bytes()).hexdigest())
    with patch.object(
        subprocess, "run", return_value=subprocess.CompletedProcess([], 0, b"", b"")
    ) as run:
        with pytest.raises(ValueError, match="clean SVG"):
            wmf2svg.convert(b"test")
        assert run.call_args.kwargs["timeout"] == 30
        assert "-Djava.awt.headless=true" in run.call_args.args[0]


def test_java_failure_falls_back_and_backend_changes_invalidate_cache(
    note, monkeypatch
):
    from shared import wmf2svg
    from shared.metafile_images import render_metafile_svg

    monkeypatch.setattr(wmf2svg, "available", lambda: True)
    with (
        ZipFile(REAL_NOTE) as z,
        patch.object(wmf2svg, "convert", side_effect=TimeoutError),
    ):
        svg, notices = render_metafile_svg(z.read("word/media/image1.wmf"))
        assert b"<svg" in svg and any("wmf2svg failed" in n for n in notices)
    first = a.parse_agreements(note, "ran1#124")
    monkeypatch.setattr(wmf2svg, "available", lambda: False)
    second = a.parse_agreements(note, "ran1#124")
    assert first["sha256"] == second["sha256"]
    assert len(list(a.CACHE_DIR.glob("*.json"))) == 2


def test_bullet_lists_use_depth_and_not_word_symbol_glyphs(tmp_path):
    from shared.docx_html import DocxHTML

    d = Document()
    numbering = d.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "500")
    for depth, glyph in enumerate(["\uf09e", "o", "\uf09f", "-"]):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(depth))
        for tag, value in [("start", "1"), ("numFmt", "bullet"), ("lvlText", glyph)]:
            node = OxmlElement("w:" + tag)
            node.set(qn("w:val"), value)
            lvl.append(node)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "500")
    aid = OxmlElement("w:abstractNumId")
    aid.set(qn("w:val"), "500")
    num.append(aid)
    numbering.append(num)
    for depth, text in [
        (0, "Parent"),
        (1, "Child"),
        (2, "Grandchild"),
        (3, "Deep"),
        (0, "Sibling"),
    ]:
        p = d.add_paragraph(text)
        pr = p._p.get_or_add_pPr()
        np = OxmlElement("w:numPr")
        for tag, value in [("numId", "500"), ("ilvl", str(depth))]:
            node = OxmlElement("w:" + tag)
            node.set(qn("w:val"), value)
            np.append(node)
        pr.append(np)
        p.paragraph_format.left_indent = Inches(0.5 * (depth + 1))
    path = tmp_path / "lists.docx"
    d.save(path)
    doc = DocxHTML(path)
    # Exercise the real agreement assembly seam, including list grouping.
    from shared.docx_html import render_blocks

    html = render_blocks(doc.block(p) for p in doc.body)
    dom = BeautifulSoup(html, "html.parser")
    assert len(dom.select("ul > li")) == 5
    assert dom.select_one("ul > li > ul > li > ul > li > ul > li").get_text() == "Deep"
    assert [
        li.find("p", recursive=False).text
        for li in dom.ul.find_all("li", recursive=False)
    ] == ["Parent", "Sibling"]
    assert not any(c in dom.get_text() for c in ["\uf09e", "\uf09f"])
    assert "margin-left:216pt" not in html and "text-indent" not in html


def test_real_agreements_have_semantic_lists_without_private_use_markers():
    result = a.parse_agreements(REAL_NOTE, "ran1#124")
    html = "".join(
        s["html"] for s in result["sections"].values()
    )
    dom = BeautifulSoup(html, "html.parser")
    assert len(dom.select("ul li")) > 100
    assert dom.select("ul ul li")
    assert not dom.select(".list-marker")
    assert not any(0xE000 <= ord(c) <= 0xF8FF for c in dom.get_text())


def test_direct_first_line_reset_overrides_inherited_hanging_indent():
    from shared.docx_html import DocxHTML, NS, text_of

    doc = DocxHTML(REAL_NOTE)
    for text in [
        "Regarding the gNB transmission power",
        "Note: The values defined in option1",
    ]:
        paragraph = next(
            p for p in doc.body.findall(".//w:p", NS) if text_of(p).startswith(text)
        )
        html = doc.paragraph(paragraph)
        assert "text-indent:-85.05pt" not in html
        assert html.count("text-indent:") == 1
        assert "text-indent:0pt" in html


def test_numbered_lists_keep_start_and_continuation_across_normal_paragraphs(tmp_path):
    from shared.docx_html import DocxHTML, render_blocks

    d = Document()
    nid = d.styles["List Number"].element.find(".//" + qn("w:numId")).get(qn("w:val"))
    num = next(
        n
        for n in d.part.numbering_part.element
        if n.tag == qn("w:num") and n.get(qn("w:numId")) == nid
    )
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "3")
    override.append(start)
    num.append(override)
    d.add_paragraph("Third", style="List Number")
    d.add_paragraph("Explanation")
    d.add_paragraph("Fourth", style="List Number")
    path = tmp_path / "numbered.docx"
    d.save(path)
    doc = DocxHTML(path)
    dom = BeautifulSoup(render_blocks(doc.block(p) for p in doc.body), "html.parser")
    assert [ol["start"] for ol in dom.select("ol")] == ["3", "4"]
    assert [li["value"] for li in dom.select("ol > li")] == ["3", "4"]
    assert dom.select_one("ol + p").text == "Explanation"


def test_real_pdcch_list_uses_paragraph_indent_instead_of_skipped_ilvl():
    result = a.parse_agreements(REAL_NOTE, 'ran1#124')
    dom = BeautifulSoup(result['sections']['10.5.2.1']['html'], 'html.parser')
    core = next(p for p in dom.select('li > p') if p.text.strip() == 'CORESET')
    assert len(core.find_parents('ul')) == 2
    assert not dom.select('.list-gap')
    assert len(core.find_parents('ul')[-1].find_all('li', recursive=False)) == 3


def test_indent_noise_and_sparse_levels_do_not_create_phantom_nesting():
    from shared.docx_html import ListItem, render_blocks
    rows = [ListItem(0, 'ul', 1, '<p>Parent</p>', 400),
            ListItem(3, 'ul', 1, '<p>Child</p>', 709),
            ListItem(5, 'ul', 1, '<p>Same indent</p>', 713),
            ListItem(0, 'ul', 1, '<p>Sibling</p>', 400)]
    dom = BeautifulSoup(render_blocks(rows), 'html.parser')
    assert [p.text for p in dom.select('ul > li > ul > li > p')] == ['Child', 'Same indent']
    assert not dom.select('ul ul ul')
    assert len(dom.ul.find_all('li', recursive=False)) == 2


def test_unlabelled_idle_mobility_continues_agreement_but_tdocs_are_excluded():
    result = a.parse_agreements(REAL_NOTE, 'ran1#124')
    last = BeautifulSoup(result['sections']['10.5.1.1']['html'], 'html.parser').get_text()
    assert 'Study 6GR signals, channels and procedures for initial access and idle mobility' in last
    assert 'Study 6GR signals, channels and procedures for idle mobility' in last
    assert last.count('Agreement') == 3  # Never invent a missing source label.
    assert 'R1-2601577' not in last and 'FL summary' not in last
    assert 'On synchronization acquisition and beam measurement' not in last


def test_full_section_keeps_nonstandard_labels_and_text_after_tdoc(tmp_path):
    d = Document()
    d.add_heading("10.1 Section", level=1)
    expected = ["Introduction without a label", "Potential agreement", "Tentative wording",
                "Proposal 1:", "Proposed text", "Observation", "Observation text",
                "Conclusion", "Conclusion text", "Remaining issues", "Unresolved text"]
    for text in expected:
        d.add_paragraph(text)
    d.add_paragraph("R1-2600001\tSummary title\tCompany")
    d.add_paragraph("R1-2600002")
    d.add_paragraph("R1-2600003 proposes a different approach.")
    d.add_paragraph("See R1-2600004 for the assumptions.")
    d.add_paragraph("Agreement")
    d.add_paragraph("Content after the TDoc list")
    d.add_heading("Discussion", level=2)
    d.add_paragraph("Unnumbered heading and discussion remain")
    d.add_heading("10.10 Next section", level=1)
    d.add_paragraph("Other agenda content")
    path = tmp_path / "Chair notes RAN1#124 - v02.docx"
    d.save(path)
    section = a.parse_agreements(path, "ran1#124")["sections"]["10.1"]
    text = BeautifulSoup(section["html"], "html.parser").get_text("\n")
    assert all(value in text for value in expected)
    assert [text.index(value) for value in expected] == sorted(text.index(value) for value in expected)
    assert "R1-2600001" not in text and "R1-2600002" not in text
    assert "R1-2600003 proposes" in text and "See R1-2600004" in text
    assert "Content after the TDoc list" in text
    assert "Unnumbered heading and discussion remain" in text
    assert "Other agenda content" not in text
    assert section["excluded_tdoc_rows"] == 2


def test_tdoc_tables_remove_only_metadata_rows_and_preserve_mixed_content(tmp_path):
    d = Document()
    d.add_heading("10.1 Section", level=1)
    table = d.add_table(rows=3, cols=3)
    for row, values in zip(table.rows, [
        ["TDoc", "Title", "Source"],
        ["R1-2600001", "Contribution", "Company"],
        ["R1-2600002", "Another contribution", "Company"],
    ]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    mixed = d.add_table(rows=3, cols=2)
    for row, values in zip(mixed.rows, [
        ["Parameter", "Value"], ["R1-2600003", "Summary title"],
        ["Reference", "See R1-2600004"],
    ]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    d.add_paragraph("Text after tables")
    path = tmp_path / "Chair notes RAN1#124 - v03.docx"
    d.save(path)
    section = a.parse_agreements(path, "ran1#124")["sections"]["10.1"]
    dom = BeautifulSoup(section["html"], "html.parser")
    assert len(dom.select("table")) == 1 and len(dom.select("tr")) == 2
    assert "Parameter" in dom.text and "See R1-2600004" in dom.text
    assert "Text after tables" in dom.text and "Contribution" not in dom.text
    assert "R1-2600003" not in dom.text
    assert section["excluded_tdoc_rows"] == 3


def test_tdoc_id_paragraph_with_image_is_not_metadata(note):
    d = Document(note)
    picture = next(p for p in d.paragraphs if p._p.xpath('.//w:drawing'))
    picture.add_run("R1-2600123")
    d.save(note)
    html = a.parse_agreements(note, "ran1#124")["sections"]["10.1"]["html"]
    assert "R1-2600123" in html and "data:image/png" in html


def test_tdoc_id_list_excludes_only_complete_metadata_paragraphs():
    assert a._tdoc_entry("R1-2600001\nR1-2600002")
    assert a._tdoc_entry("R1-2600001 rev 2\tTitle\tSource")
    assert not a._tdoc_entry("R1-2600001\nSubstantive explanation")
    assert not a._tdoc_entry("Refer to R1-2600001\tfor details")


def test_section_extraction_invalidates_previous_marker_cache(note, monkeypatch):
    old = a.parse_agreements(note, "ran1#124")
    cache_file = next(a.CACHE_DIR.glob("*.json"))
    import json
    cache_file.rename(cache_file.with_name("11-" + cache_file.name.split("-", 1)[1]))
    old["sections"] = {}
    next(a.CACHE_DIR.glob("*.json")).write_text(json.dumps(old))
    assert a.parse_agreements(note, "ran1#124")["sections"]["10.2"]["html"]


def test_embedded_ole_equation_uses_static_preview_with_vml_size(note):
    from lxml import etree as ET
    from shared.docx_html import DocxHTML, NS
    doc = DocxHTML(note)
    rid = next(e.get('{%s}embed' % NS['r']) for e in doc.body.findall('.//a:blip', NS))
    paragraph = ET.fromstring(f'''<w:p xmlns:w="{NS['w']}" xmlns:r="{NS['r']}"
      xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">
      <w:r><w:object><v:shape style="width:31.3pt;height:15.65pt">
      <v:imagedata r:id="{rid}"/></v:shape>
      <o:OLEObject ProgID="Equation.3" r:id="ignored-active-object"/>
      </w:object></w:r></w:p>''')
    rendered = doc.paragraph(paragraph)
    assert 'data:image/png;base64,' in rendered
    assert 'width:41.7333px' in rendered
    assert 'ignored-active-object' not in rendered and '[Object:' not in rendered
