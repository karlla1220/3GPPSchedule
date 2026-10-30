"""RAN3 time plans: the chair's living agenda, the offline room sheets and the rules they rely on."""
import io
import subprocess
import sys
from types import SimpleNamespace

from bs4 import BeautifulSoup
from docx import Document
import httpx
import pytest

from pathlib import Path
from urllib.parse import unquote

from shared.agenda import agenda_map
from shared.renderer import generate_html
from shared.schedule import load_schedule, save_schedule
from working_groups.ran3 import document as docmod, lifecycle, llm, pipeline, sessions, sources

FIXTURES = Path(__file__).parent / 'fixtures/ran3'
EOM = 'RAN3_133bis_agenda_20261016_EOM.docx'
MID = 'RAN3_133bis_agenda_20261014_1825.docx'
PRE = 'RAN3_134_agenda_20261023.docx'
OLD = 'RAN3_129_agenda_20250829_1620_EOM.docx'
META_133BIS = {'starts_on': '2026-10-12', 'ends_on': '2026-10-16', 'timezone': 'Asia/Seoul'}
META_134 = {'starts_on': '2026-11-16', 'ends_on': '2026-11-20', 'timezone': 'America/Edmonton'}
META_129 = {'starts_on': '2025-08-25', 'ends_on': '2025-08-29', 'timezone': 'Asia/Kolkata'}
NAMES = {'L': 'Lixiang XU', 'X': 'Xudong YANG', 'G': 'Gen CAO'}


def extract(name):
    return docmod.extract_document((FIXTURES / name).read_bytes(), name)


def agenda(name):
    return agenda_map((FIXTURES / name).read_bytes())


def slots(name):
    return docmod.offline_sheets((FIXTURES / name).read_bytes())[0]


def build(name, metadata, agenda_name=None, offline=None, cbs=None, warnings=None, refine=None):
    return sessions.make_schedule(extract(name), agenda(agenda_name) if agenda_name else None, metadata, [name], {},
                                  'now', slots=slots(offline) if offline else [], cbs=cbs,
                                  chair=lambda label: NAMES.get(label, label), refine=refine, warnings=warnings)


@pytest.fixture(scope='module')
def eom():
    """#133bis at its close: the time plan, the offline sheet, and the CB blocks of an earlier version."""
    return build(EOM, META_133BIS, 'agenda_133bis.csv', 'offline_133bis_ver_03.xlsx', cbs=extract(MID)['cbs'])


def cells(schedule, day, start=None, room='main'):
    day = next(d for d in schedule.days if d.day_name == day)
    return [(s.start_time, s.end_time, s.part, s.name, s.agenda_item, s.chair, s.notes) for s in day.sessions
            if s.room_ids == [room] and (start is None or s.start_time == start)]


# ------------------------------------------------------------------ documents

@pytest.mark.parametrize('name, expected', [
    ('RAN3_133bis_agenda_20261014_1825.docx', ('ran3#133bis', '20261014.0.1825')),
    ('RAN3_133bis_agenda_20261016_EOM.docx', ('ran3#133bis', '20261016.1.0000')),
    ('RAN3_129_agenda_20250829_1620_EOM.docx', ('ran3#129', '20250829.1.1620')),
    ('RAN3_134_agenda_20261023.docx', ('ran3#134', '20261023.0.0000')),
    ('RAN3_133bis_agenda_20261014_0900_LX.docx', None),     # a vice chair's copy
    ('RAN3_133_agenda_20260826_1800-yxd.docx', None),
    ('RAN3 offline discussion_ver_03.xlsx', None),
])
def test_living_agenda_filenames(name, expected):
    assert docmod.file_info(name) == expected


def test_versions_sort_with_the_closing_one_last_and_labels_name_them():
    names = ['RAN3_133bis_agenda_20261016_EOM.docx', 'RAN3_133bis_agenda_20261016_1300.docx',
             'RAN3_133bis_agenda_20261009.docx', 'RAN3_133bis_agenda_20261012_0900.docx']
    assert sorted(names, key=lambda n: docmod.file_info(n)[1])[-2:] == names[1::-1]
    assert [docmod.source_label(n) for n in names] == ['agenda_20261016', 'agenda_20261016_1300',
                                                       'agenda_20261009', 'agenda_20261012_0900']
    manifest = {'schedule': {'name': MID}, 'offline': {'name': 'RAN3 offline discussion_ver_03.xlsx'},
                'agenda': {'name': 'agenda.csv'}}
    assert pipeline.source_labels(manifest) == {MID: 'agenda_20261014_1825',
                                                'RAN3 offline discussion_ver_03.xlsx': 'offline discussion_ver_03',
                                                'agenda.csv': 'agenda.csv'}


def test_block_times_come_from_the_time_column():
    document = extract(EOM)
    assert document['meeting_id'] == 'ran3#133bis'
    wednesday = next(d for d in document['days'] if d['day'] == 'Wednesday')
    assert [(b['name'], b['start'], b['end']) for b in wednesday['breaks']] == [
        ('Coffee Break', '10:30', '11:00'), ('Lunch Break', '13:00', '14:30'), ('Coffee Break', '16:30', '17:00')]
    spans = [(c['start'], c['end'], c['part'], c['in_break'], c['lines']) for c in wednesday['cells']]
    assert spans == [
        ('08:00', '09:00', None, False, ['6G Offline*']),
        ('09:00', '10:30', None, False, ['AIML RAN: AI 12', 'L']),
        ('10:30', '11:00', None, True, ['6G Offline (CB#8)']),          # in the coffee cell, under "Coffee Break"
        ('11:00', '13:00', None, False, ['ISAC SI: AI 13']),
        ('14:30', '16:30', None, False, ['6G SI: AI 10.5', 'L']),
        ('16:30', '17:00', None, True, ['6G Offline (CB#9)']),
        ('17:00', '19:00', [1, 2], False, ['MobEnh: AI 15']),            # two untimed rows of one block
        ('17:00', '19:00', [2, 2], False, ['XR: AI 20', '6G Offline (CB#12)']),
        ('19:00', '20:00', None, False, ['RAN3 Social Event', '(self-funded dinner)'])]   # two rows, one text
    tuesday = next(d for d in document['days'] if d['day'] == 'Tuesday')
    assert [(c['start'], c['end']) for c in tuesday['cells'][:2]] == [('08:00', '08:30'), ('08:30', '09:00')]
    assert document['legend'] == [['blue', '6G'], ['L, X', 'chaired by Vice-Chairs'], ['highlighted', 'changed'],
                                  ['strikethrough', 'not treated'], ['*', 'if needed']]


def test_the_oldest_form_reads_too_and_struck_text_is_not_treated():
    document = extract(OLD)                     # #129: "0800", "1030~1100", "Lunch Time", "1930: Social Event"
    monday = next(d for d in document['days'] if d['day'] == 'Monday')
    assert monday['breaks'][1] == {'name': 'Lunch Time', 'start': '13:00', 'end': '14:30'}
    struck = next(c for c in monday['cells'] if c['start'] == '17:00')
    assert struck['lines'] == ['Corrections: AI 9', '', 'Duplex WI: AI 19'] and struck['struck'] == ['/31']
    schedule = build(OLD, META_129)
    assert cells(schedule, 'Monday', '17:00')[0][6] == ['Corrections: AI 9', 'Not treated: /31']
    assert cells(schedule, 'Wednesday', '19:30') == [('19:30', '20:00', None, 'Social Event', None, None,
                                                      ['1930: Social Event', 'Sheraton Hotel'])]


def test_cb_blocks_name_the_agenda_item_and_moderator():
    cbs = extract(MID)['cbs']
    assert cbs['8'] == {'agenda_item': '10.3.1.3', 'name': 'CB: # 8_6GSBI', 'moderator': 'CATT'}
    assert cbs['1']['agenda_item'] == '8.3.1'           # not the form "CB: # 1_Name" in the reminders
    assert cbs['AIML1']['agenda_item'] == '12.2.1' and len(cbs) == 22
    assert extract(EOM)['cbs'] == {}                     # the closing version drops them


def schedule_docx(rows):
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=7)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_a_table_out_of_the_ran3_form_is_a_layout_error():
    head = ['', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', '']
    data = schedule_docx([head] + [['09:00', 'Org: AI 1-8', '', '', '', '', '']] * 5)
    with pytest.raises(docmod.LayoutError, match='6x7'):
        docmod.extract_document(data, 'RAN3_134_agenda_20261023.docx')
    with pytest.raises(docmod.LayoutError, match='Monday to Friday'):
        docmod.extract_document(schedule_docx([['', 'Mon', '', '', '', '', '']]), 'RAN3_134_agenda_20261023.docx')


def test_agenda_csv_and_offline_sheets():
    assert agenda('agenda_133bis.csv')['10.3.1.3'].startswith('Service Based Interface')
    found, unread = docmod.offline_sheets((FIXTURES / 'offline_133bis_ver_03.xlsx').read_bytes())
    assert unread == []
    assert found[0]['day'] == 'Monday' and (found[0]['month'], found[0]['date']) == (10, 12)
    son = [s for s in found if s['topic'].startswith('R20 SON/MDT')]
    assert [(s['start'], s['end']) for s in son][:2] == [('11:00', '11:30'), ('11:30', '12:00')]
    # Merged rows (#131) are filled from their first cell.
    merged = docmod.offline_sheets((FIXTURES / 'offline_131_ver_06.xlsx').read_bytes())[0]
    assert sum(s['topic'] == 'Rel-20 Ambient IoT' for s in merged) > 1


# ------------------------------------------------------------------ cells

@pytest.mark.parametrize('line, kind, items', [
    ('Org: AI 1-8', 'label', ['1', '2', '3', '4', '5', '6', '7', '8']),
    ('AI RAN WI: AI 11', 'label', ['11']),
    ('IoT NTN WI AI 15', 'label', ['15']),
    ('LTE MBS WI: AI23', 'label', ['23']),
    ('Corrections: AI 9/31', 'label', ['9', '31']),
    ('6G SI: AI 10 (incl. 10.5)', 'label', ['10']),
    ('AI32, AI6', 'number', ['32', '6']),
    ('10.3.1.1 (cont.)', 'number', ['10.3.1.1']),
    ('L', 'initial', []),
    ('6G Offline (CB#8)', 'offline', []),
    ('Early CBs', 'progress', []),
    ('(1hr until 19:30 at latest)', 'modifier', []),
    ('1930: Social Event', 'clock', []),
    ('0900 START OF MEETING', 'marker', []),
    ('All CBs in AI8, AI9, AI31 will be treated', 'unknown', []),
])
def test_line_kinds(line, kind, items):
    assert sessions.kind(line) == kind and sessions.line_items(line, None) == items


def test_the_eom_grid_matches_the_report(eom):
    # research/ran3/report/section.json, "예시: 수요일 09:00–13:00".
    assert cells(eom, 'Wednesday', '09:00') == [('09:00', '10:30', None, 'AIML RAN', '12', 'Lixiang XU',
                                                 ['AIML RAN: AI 12', 'L'])]
    assert cells(eom, 'Wednesday', '10:30') == [('10:30', '11:00', None, '6G Offline', '10.3.1.3', None,
                                                 ['6G Offline (CB#8)', 'CB: # 8_6GSBI (moderator - CATT)'])]
    assert cells(eom, 'Wednesday', '11:00')[0][3:6] == ('ISAC SI', '13', None)
    assert cells(eom, 'Wednesday', '09:00', 'offline') == [('09:00', '10:30', None, 'Rel-20 Ambient IoT offline', None,
                                                            'Yan Wang', ['Rel-20 Ambient IoT offline', 'Yan Wang (Huawei)'])]
    assert cells(eom, 'Wednesday', '11:00', 'offline')[0][:6] == ('11:00', '13:00', None, 'R20 SON/MDT', '11.2', 'Xufei Dong')
    # "예시: 목요일 14:30–19:00": label lines and table rows are pieces, in order.
    thursday = [c[:5] for c in cells(eom, 'Thursday') if c[0] >= '14:30']
    assert thursday == [('14:30', '16:30', [1, 2], '6G SI', '10.2.2'), ('14:30', '16:30', [2, 2], 'TEI20', '31'),
                        ('16:30', '17:00', None, '6G Offline*', None), ('17:00', '19:00', [1, 2], 'Early CBs', None),
                        ('17:00', '19:00', [2, 2], '6G Offline', '10.8.1')]
    # "예시: 수요일 16:30–19:00": the offline in XR's row is a note, not a piece.
    assert cells(eom, 'Wednesday', '17:00')[1][2:6] == ([2, 2], 'XR', '20', None)
    assert cells(eom, 'Wednesday', '17:00')[1][6] == ['XR: AI 20', '6G Offline (CB#12)']
    assert cells(eom, 'Wednesday', '17:00', 'offline')[0][3:6] == ('CB: # 1_R20N3pathFail', '8.3.1', 'Nianshan Shi')


def test_days_rooms_breaks_and_colours(eom):
    assert [(d.day_name, d.date, [r.name for r in d.rooms]) for d in eom.days] == [
        ('Monday', '2026-10-12', ['Main session']), ('Tuesday', '2026-10-13', ['Main session']),
        ('Wednesday', '2026-10-14', ['Main session', 'Offline room']),
        ('Thursday', '2026-10-15', ['Main session', 'Offline room']), ('Friday', '2026-10-16', ['Main session'])]
    monday = eom.days[0]
    assert monday.timeline.start == '09:00'            # "0900 START OF MEETING" is no session
    assert all(b['label_position'] == 'time-axis' for d in eom.days for b in d.timeline.breaks)
    groups = {s.name: s.group_header for d in eom.days for s in d.sessions}
    assert groups['6G SI'] == groups['6G Offline'] == 'Study on 6G Radio'
    assert groups['Early CBs'] == '' and groups['TEI20'] == 'Corrections and Enhancements to Rel-20'
    friday = cells(eom, 'Friday', '14:30')[0]
    assert friday[3] == 'CBs*' and friday[6] == ['CBs*', 'END OF MEETING']
    assert eom.supplements[0] == {'type': 'heading', 'text': 'Key to the time plan'}


def test_blank_lines_set_pieces_apart_and_an_initial_chairs_its_group():
    schedule = build(OLD, META_129)
    assert [c[2:6] for c in cells(schedule, 'Thursday', '09:00')] == [
        ([1, 2], 'NR NTN WI / IoT NTN WI', '14, 15', None), ([2, 2], 'Network ES WI', '17', None)]
    assert [c[3:6] for c in cells(schedule, 'Thursday', '14:30')] == [('SON/MDT WI', '10', 'Lixiang XU'),
                                                                       ('Early CBs', None, None)]
    # Without blank lines the initial is written for the whole cell.
    assert [c[3:6] for c in cells(schedule, 'Wednesday', '17:00')] == [('LP WUS WI', '18', 'Gen CAO'),
                                                                        ('SL Relay WI', '22', 'Gen CAO')]
    assert cells(schedule, 'Friday', '14:30')[0][3:5] == ('CBs', '32, 6')    # "END OF MEETING" is no piece


def synthetic(lines, start='14:30', end='16:30', part=None):
    cell = {'id': 'c0', 'day': 'Monday', 'start': start, 'end': end, 'part': part, 'lines': lines, 'struck': [],
            'in_break': False, 'evening': False}
    return {'meeting_id': 'ran3#134', 'days': [{'day': 'Monday', 'cells': [cell], 'breaks': []}], 'legend': [],
            'cbs': {}}


def test_pieces_shorter_than_30_minutes_stay_one_cell():
    labels = [f'Topic {n}: AI {n}' for n in range(1, 6)]
    found = sessions.interpret(synthetic(labels))
    assert len(found) == 1 and found[0]['name'] == 'Topic 1 / Topic 2 / Topic 3 …' and found[0]['part'] is None
    assert found[0]['agenda_items'] == ['1', '2', '3', '4', '5'] and found[0]['notes'] == labels
    # Two labels in the second of two untimed rows: quarters of the block.
    nested = sessions.interpret(synthetic(['XR: AI 20', 'AIML Air: AI 21'], '17:00', '19:00', [2, 2]))
    assert [s['part'] for s in nested] == [[3, 4], [4, 4]]


def test_unknown_lines_flag_the_cell_and_items_not_in_the_agenda_are_reported():
    old = extract(OLD)
    flagged = [c['id'] for d in old['days'] for c in d['cells'] if sessions.uncertainty(c)]
    assert flagged == ['thu.r7']
    pre = extract(PRE)
    assert [c['id'] for d in pre['days'] for c in d['cells'] if sessions.uncertainty(c)] == ['fri.r7']
    warnings = []
    schedule = sessions.make_schedule(synthetic(['6G SI: AI 10.99, 10']), agenda('agenda_134.csv'), META_134,
                                      [PRE], {}, 'now', warnings=warnings)
    assert schedule.days[0].sessions[0].agenda_item == '10'
    assert warnings == ['agenda items not in agenda.csv: 10.99']


def test_offline_room_sessions_join_rows_and_drop_what_is_not_a_session():
    merged = build(EOM, META_133BIS, 'agenda_133bis.csv', 'offline_131_ver_06.xlsx')
    # #131's sheet is another meeting's: its dates are not #133bis's.
    assert all(r.id == 'main' for d in merged.days for r in d.rooms)
    found, warnings = sessions.offline_sessions(slots('offline_131_ver_06.xlsx'), {}, None, {
        'Monday': '2026-02-09', 'Tuesday': '2026-02-10', 'Wednesday': '2026-02-11', 'Thursday': '2026-02-12',
        'Friday': '2026-02-13'})
    assert warnings == []
    names = [(s['day'][:3], s['start'], s['end'], s['name'], s['chair']) for s in found]
    assert ('Wed', '09:00', '10:30', 'Rel-20 Ambient IoT', 'Yan Wang') in names
    assert not any(n[3].startswith(('Reserved', 'No parallel', 'Topic')) for n in names)


# ------------------------------------------------------------------ LLM fallback

class FakeGemini:
    def __init__(self, *responses):
        self.responses, self.calls = list(responses), []

    def __call__(self, payload, schema, model, thinking=None):
        self.calls.append(payload)
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response


def answer(*pieces):
    return {'cells': [{'id': 'thu.r7', 'pieces': [dict(zip(('title', 'chair', 'agenda_items', 'lines'), p))
                                                  for p in pieces]}]}


def thursday_afternoon(refine):
    found = sessions.interpret(extract(OLD), None, None, refine)
    return [(s['part'], s['name'], s['chair_label'], s['agenda_items']) for s in found
            if s['day'] == 'Thursday' and s['start'] == '14:30']


def test_a_checked_answer_replaces_the_rules_for_an_uncertain_cell(tmp_path):
    gemini = FakeGemini(answer(('SON/MDT WI', 'L', ['10'], [0, 1]), ('Early CBs', None, ['8', '9', '31'], [3, 4])))
    refine = llm.CellRefiner('test-model', gemini, cache_dir=tmp_path)
    assert thursday_afternoon(refine) == [([1, 2], 'SON/MDT WI', 'L', ['10']),
                                          ([2, 2], 'Early CBs', None, ['8', '9', '31'])]
    assert refine.summary == {'model': 'test-model', 'uncertain': 1, 'applied': 1, 'kept_rules': 0}
    assert len(gemini.calls) == 1 and '"thu.r7"' in gemini.calls[0]
    again = llm.CellRefiner('test-model', FakeGemini(), cache_dir=tmp_path)    # answered from the cache
    assert thursday_afternoon(again)[1][3] == ['8', '9', '31']


@pytest.mark.parametrize('bad, problem', [
    (answer(('SON/MDT', 'L', ['10'], [0, 1]), ('Early treated batch', None, [], [3, 4])), 'title words'),
    (answer(('SON/MDT WI', 'Lixiang', ['10'], [0, 1]), ('Early CBs', None, [], [3, 4])), 'initial'),
    (answer(('SON/MDT WI', 'L', ['10'], [0, 1]), ('Early CBs', None, ['12'], [3, 4])), 'agenda item'),
    (answer(('SON/MDT WI', 'L', ['10'], [0, 1]), ('Early CBs', None, [], [3])), 'not cited'),
    (answer(('SON/MDT WI', 'L', ['10'], [0]), ('L', None, [], [1]), ('Early CBs', None, [], [3]),
            ('All CBs', None, [], [4]), ('treated', None, [], [])), 'under 30 min'),
])
def test_unchecked_answers_are_retried_once_then_the_rules_stand(tmp_path, bad, problem):
    rules = thursday_afternoon(None)
    gemini = FakeGemini(bad, bad)
    refine = llm.CellRefiner('test-model', gemini, cache_dir=tmp_path)
    assert thursday_afternoon(refine) == rules
    assert len(gemini.calls) == 2 and problem in gemini.calls[1]


def test_an_llm_outage_never_breaks_the_build(tmp_path):
    refine = llm.CellRefiner('test-model', FakeGemini(RuntimeError('quota')), cache_dir=tmp_path)
    assert thursday_afternoon(refine) == thursday_afternoon(None)
    assert refine.summary['kept_rules'] == 1 and 'quota' in refine.summary['error']


# ------------------------------------------------------------------ render

def test_render_roundtrip_and_split_pieces(eom, tmp_path):
    save_schedule(eom, tmp_path / 'schedule.json')
    assert load_schedule(tmp_path / 'schedule.json') == eom
    html = BeautifulSoup(generate_html(eom), 'html.parser')
    assert len(html.select('.session-block')) == sum(len(d.sessions) for d in eom.days)
    piece = html.select_one('#thursday [data-name="TEI20"]')
    assert '2/2 of 120 min' in piece['data-popup']
    assert 'Key to the time plan' in html.get_text()


# ------------------------------------------------------------------ sources

def listing(names, base):
    rows = ''.join(f'<tr><td><a href="{base}{n}">{n}</a></td><td>2026/10/14 9:25</td></tr>' for n in names)
    return f'<table>{rows}</table>'


def mock_server(files, seen=None):
    """3GPP-like server: listings for folders, 403 for missing paths."""
    def handler(request):
        url = unquote(str(request.url))
        if seen is not None:
            seen.append(url)
        if url.endswith('/'):
            names = sorted({u[len(url):].split('/')[0] + ('/' if '/' in u[len(url):] else '')
                            for u in files if u.startswith(url) and u != url})
            return httpx.Response(200, text=listing(names, url)) if names else httpx.Response(403)
        if url in files:
            return httpx.Response(200, content=files[url], headers={'etag': f'"{len(files[url])}"'})
        return httpx.Response(403)
    return httpx.Client(transport=httpx.MockTransport(handler))


SYNC = 'https://www.3gpp.org/ftp/Meetings_3GPP_SYNC/RAN3/'
ARCHIVE = 'https://www.3gpp.org/ftp/tsg_ran/WG3_Iu/'
ROWS = [{'Title': '3GPPRAN3#133-bis', 'TBId': 381, 'Id': 1, 'StartDate': '2026-10-12 09:00:00',
         'EndDate': '2026-10-16 17:00:00', 'StartTimeZone': '(GMT+09:00) Seoul', 'Country': 'KR'},
        {'Title': '3GPPRAN3#134', 'TBId': 381, 'Id': 2, 'StartDate': '2026-11-16 09:00:00',
         'EndDate': '2026-11-20 17:00:00', 'StartTimeZone': '(GMT-07:00) Mountain Time (US & Canada)',
         'Country': 'CA'}]


@pytest.fixture
def server_files():
    body = (FIXTURES / MID).read_bytes()
    return {SYNC + sources.CHAIRS + MID: body,
            SYNC + sources.CHAIRS + 'RAN3_133bis_agenda_20261014_1900_LX.docx': b'a vice chair copy',
            SYNC + sources.OFFLINE + 'RAN3 offline discussion_ver_02.xlsx': b'older',
            SYNC + sources.OFFLINE + 'RAN3 offline discussion_ver_03.xlsx':
                (FIXTURES / 'offline_133bis_ver_03.xlsx').read_bytes(),
            ARCHIVE + 'TSGR3_133-bis/Agenda/agenda.csv': (FIXTURES / 'agenda_133bis.csv').read_bytes(),
            ARCHIVE + 'TSGR3_133-bis/Inbox/Chairs_Notes/RAN3_133bis_agenda_20261009.docx': b'pre-meeting',
            ARCHIVE + 'TSGR3_134/Inbox/Chairs_Notes/' + PRE: (FIXTURES / PRE).read_bytes(),
            ARCHIVE + 'TSGR3_134/Agenda/agenda.csv': (FIXTURES / 'agenda_134.csv').read_bytes()}


def test_fetch_bundle_takes_the_live_version_and_its_offline_sheet(tmp_path, monkeypatch, server_files):
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')
    monkeypatch.setattr(sources.meeting_folders, 'CACHE_PATH', tmp_path / 'folders.json')
    with mock_server(server_files) as http:
        bundle = sources.fetch_bundle(sources.DEFAULTS, {}, http, rows=ROWS, today='2026-10-14',
                                      cache_dir=tmp_path / 'inputs')
    manifest = bundle.manifest
    assert (manifest['meeting_id'], manifest['version']) == ('ran3#133bis', '20261014.0.1825')
    assert manifest['offline']['name'] == 'RAN3 offline discussion_ver_03.xlsx'
    assert manifest['agenda']['url'] == ARCHIVE + 'TSGR3_133-bis/Agenda/agenda.csv'
    assert sources.source_names(manifest) == [MID, 'RAN3 offline discussion_ver_03.xlsx', 'agenda.csv']


def test_before_a_meeting_the_sync_folders_last_sheet_is_not_taken(tmp_path, monkeypatch, server_files):
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')
    monkeypatch.setattr(sources.meeting_folders, 'CACHE_PATH', tmp_path / 'folders.json')
    with mock_server(server_files) as http:
        bundle = sources.fetch_bundle(sources.DEFAULTS, {}, http, rows=ROWS, today='2026-10-30',
                                      cache_dir=tmp_path / 'inputs')
    assert (bundle.manifest['meeting_id'], bundle.manifest['schedule']['name']) == ('ran3#134', PRE)
    assert bundle.manifest['offline'] is None and bundle.offline == b''
    assert bundle.manifest['agenda']['url'] == ARCHIVE + 'TSGR3_134/Agenda/agenda.csv'


def test_a_vanished_live_version_is_kept_instead_of_stepping_back(tmp_path, monkeypatch, server_files):
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')
    monkeypatch.setattr(sources.meeting_folders, 'CACHE_PATH', tmp_path / 'folders.json')
    with mock_server(server_files) as http:
        built = sources.fetch_bundle(sources.DEFAULTS, {}, http, rows=ROWS, today='2026-10-14',
                                     cache_dir=tmp_path / 'inputs')
    del server_files[SYNC + sources.CHAIRS + MID]
    sources.remote_files.new_run()
    with mock_server(server_files) as http:
        kept = sources.fetch_bundle(sources.DEFAULTS, built.manifest, http, rows=ROWS, today='2026-10-14',
                                    cache_dir=tmp_path / 'inputs')
    assert kept.manifest['schedule']['name'] == MID and kept.schedule == built.schedule


def test_config_rejects_unknown_keys_and_bad_aliases(tmp_path, monkeypatch):
    monkeypatch.setattr(sources, 'CONFIG_PATH', tmp_path / 'config.json')
    for bad in ({'unknown': 1}, {'chair_aliases': {'L': 'Lixiang'}}, {'meetings': {'x': {}}},
                {'meeting_specific': ['https://www.3gpp.org/ftp/tsg_ran/WG3_Iu/TSGR2_133/']}):
        sources.write_json(tmp_path / 'config.json', bad)
        with pytest.raises(ValueError):
            sources.load_config()
    sources.write_json(tmp_path / 'config.json', {'meeting_specific': [ARCHIVE + 'TSGR3_133-bis/']})
    assert sources.load_config()['chair_aliases']['L'] == 'XU, Lixiang'


# ----------------------------------------------------------- pipeline and CI

@pytest.fixture
def offline_config(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = {**sources.DEFAULTS, 'local_agenda': str(FIXTURES / 'agenda_133bis.csv'),
           'local_offline': str(FIXTURES / 'offline_133bis_ver_03.xlsx'),
           'meetings': {'133bis': dict(META_133BIS)}}
    monkeypatch.setattr(pipeline, 'load_config', lambda: cfg)
    monkeypatch.setattr(sources, 'load_config', lambda: cfg)
    monkeypatch.setattr(pipeline, 'portal_rows', lambda: pytest.fail('Portal queried offline'))
    monkeypatch.setattr(sources, 'portal_rows', lambda: [])
    monkeypatch.setattr(pipeline, 'has_key', lambda: False)   # never reach Gemini from tests
    return cfg


def test_local_builds_need_no_network_and_keep_the_cb_blocks_past_the_eom(offline_config, tmp_path, monkeypatch):
    monkeypatch.setattr(httpx.Client, 'send', lambda *a, **k: pytest.fail('Offline network request'))
    out = tmp_path / 'out'
    first = pipeline.build_schedule(SimpleNamespace(local=str(FIXTURES / MID), output_dir=out))
    assert first.wg_id == 'ran3' and first.meeting_name == 'RAN3#133bis'
    assert first.source_labels == {MID: 'agenda_20261014_1825', 'offline_133bis_ver_03.xlsx': 'offline_133bis_ver_03',
                                   'agenda_133bis.csv': 'agenda.csv'}
    state = sources.read_json(out / 'ran3/.schedule_state.json')
    assert state['cbs']['8']['agenda_item'] == '10.3.1.3' and state['llm'] == {'used': False}
    closing = pipeline.build_schedule(SimpleNamespace(local=str(FIXTURES / EOM), output_dir=out))
    coffee = next(s for d in closing.days if d.day_name == 'Wednesday' for s in d.sessions if s.start_time == '10:30')
    assert coffee.agenda_item == '10.3.1.3'        # the EOM has no CB blocks; the state remembers them
    again = pipeline.build_schedule(SimpleNamespace(no_download=True, output_dir=out))
    assert again.days == closing.days


def test_a_new_meeting_without_dates_fails_clearly(offline_config, tmp_path, monkeypatch):
    offline_config['meetings'] = {}
    monkeypatch.setattr(pipeline, 'portal_rows', lambda: None)
    with pytest.raises(ValueError, match='Add them under "meetings"'):
        pipeline.build_schedule(SimpleNamespace(local=str(FIXTURES / PRE), output_dir=tmp_path / 'out'))
    assert not (tmp_path / 'out/ran3/.schedule_state.json').exists()


def test_lifecycle_check_stages_without_advancing_state(offline_config, tmp_path, monkeypatch):
    bundle = sources.local_bundle(FIXTURES / EOM, offline_config)
    monkeypatch.setattr(sources, 'fetch_bundle', lambda *a, **k: bundle)
    before = {'meeting_id': 'ran3#133'}
    sources.write_json(sources.OUTPUT / '.schedule_state.json', before)
    result = lifecycle.check_updates()
    assert result.changed and 'meeting_id changed' in result.reasons and not result.errors
    assert sources.read_json(sources.OUTPUT / '.schedule_state.json') == before
    lifecycle.prepare_build()
    assert sources.Bundle.load(sources.DOWNLOADS / 'prepared').offline == bundle.offline
    sources.write_json(sources.OUTPUT / '.schedule_state.json', {**bundle.manifest, 'metadata': META_133BIS})
    assert not lifecycle.check_updates().changed
    monkeypatch.setattr(sources, 'fetch_bundle', lambda *a, **k: (_ for _ in ()).throw(ValueError('down')))
    assert lifecycle.check_updates().errors == ['RAN3 source check failed: down']


def test_reset_keeps_the_last_successful_output(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    sources.write_json(sources.OUTPUT / '.schedule_state.json', {'meeting_id': 'ran3#133bis'})
    sources.write_json(sources.DOWNLOADS / 'inputs/manifest.json', {})
    lifecycle.reset_cache()
    assert (sources.OUTPUT / '.schedule_state.json').exists() and not sources.DOWNLOADS.exists()


def test_lifecycle_import_is_isolated():
    result = subprocess.run([sys.executable, '-c', """
import sys
from working_groups.registry import get_working_group
get_working_group('ran3')
assert not any(name.startswith(('working_groups.ran1', 'working_groups.ran2', 'working_groups.ran_plenary',
                                'google.genai', 'docx', 'openpyxl'))
               for name in sys.modules)
"""], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
