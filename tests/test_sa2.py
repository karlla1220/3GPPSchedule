"""SA2 Session Plans: real chair documents plus the rules they rely on."""
import io
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
import pytest

from shared.renderer import generate_html
from shared.schedule import load_schedule, save_schedule
from working_groups.sa2 import document as docmod, sessions

FIXTURES = Path(__file__).parent / 'fixtures/sa2'
R20 = 'S2-2608833_SA2-176-Session Plan - r20.docx'
R12 = 'S2-2509839_SA2-172-Session Plan - r12.docx'
R9 = 'S2-2601721_SA2-174-Session Plan -r9.docx'
META_176 = {'starts_on': '2026-08-24', 'ends_on': '2026-08-28', 'timezone': 'Europe/Prague',
            'starts_at': '2026-08-24T09:00:00+02:00', 'ends_at': '2026-08-28T16:30:00+02:00'}


def extract(name):
    return docmod.extract_document((FIXTURES / name).read_bytes(), name)


def agenda():
    return docmod.agenda_map((FIXTURES / 'agenda_176.csv').read_bytes())


def read(name, items=None):
    return sessions.interpret(extract(name), items)[0]


def find(found, day, start, room=None, part=None, scope='assigned'):
    return [s for s in found if s['day'] == day and s['start'] == start and s['scope'] == scope
            and (room is None or s['rooms'] == [room]) and (part is None or s['part'] == part)]


@pytest.mark.parametrize('name, expected', [
    ('S2-2610195_SA2-177-Session Plan - r7.docx', ('sa2#177', 7, '')),
    ('S2-2610195_SA2-177-Session Plan -r14 - Friday plan.docx', ('sa2#177', 14, 'Friday plan')),
    ('S2-2610195_SA2-177-Session Plan -r13 with Thursday revisions plan.docx', ('sa2#177', 13, 'with Thursday revisions plan')),
    ('S2-260nnnn_SA2-177-Session Plan.zip', ('sa2#177', 0, '')),
    ('Draft S2-2606837_SA2-176-Session Plan - r1.docx', ('sa2#176', 1, '')),
    ('S2-2603539_SA2-175-Session Plan -r12.zip', ('sa2#175', 12, '')),
    ('Draft Feb 2026 session plan -r9.zip', None),   # #173's slot list, not a Session Plan
    ('S2-2610193.zip', None),
])
def test_session_plan_filenames(name, expected):
    assert docmod.file_info(name) == expected


def test_template_days_quarters_breaks_rooms_and_conveners():
    doc = extract(R20)
    assert doc['meeting_id'] == 'sa2#176'
    assert doc['days'] == ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    assert [(q['start'], q['end']) for q in doc['quarters']] == [
        ('08:00', '08:50'), ('09:00', '10:30'), ('11:00', '12:30'), ('14:00', '15:30'), ('16:00', '17:30'),
        ('18:00', '19:30')]
    # Break rows carry no times; the quarters around them do.
    assert [(b['name'], b['start'], b['end']) for b in doc['breaks']] == [
        ('Break', '08:50', '09:00'), ('Coffee', '10:30', '11:00'), ('Lunch', '12:30', '14:00'),
        ('Coffee', '15:30', '16:00'), ('Break', '17:30', '18:00')]
    assert [(r['id'], r['name']) for r in doc['rooms']] == [
        ('stream1', 'Main'), ('stream2', 'Breakout 1'), ('stream3', 'Breakout 2')]
    # "Convener 1: Mike, Convenor 2: Wanqiang", linked by highlight to the stream labels.
    assert doc['conveners'] == {'green': 'Mike', 'cyan': 'Wanqiang'}
    assert doc['stream_colours'] == {2: 'green', 3: 'cyan'}
    assert {'day': 'Thursday', 'text': '(Revisions only unless stated)'} in doc['notes']


def test_tracked_changes_are_read_in_accepted_form():
    lines = [line for cell in extract(R20)['cells'] for line in cell['lines']]
    assert 'KI#9 (Localized Net) (20.6.9)' in lines   # inserted
    assert not any('Other 6G (tbd)' in line for line in lines)   # deleted


def test_the_heading_and_text_outside_the_table_are_kept_except_the_copied_title():
    doc = extract(R20)
    texts = [b.get('text', '') for b in doc['supplements']]
    assert any(t.startswith('When sessions run in parallel') for t in texts)
    assert any(t.startswith('Convener 1: Mike') for t in texts)
    # The first upload copies the meeting line from the previous meeting, so it is left out.
    assert not any(t.startswith('SA WG2 Meeting') or 'Prague, Czech Republic' in t for t in texts)
    assert any(b['type'] == 'table' and 'Main Room' in b['rows'][0] for b in doc['supplements'])


def test_another_layout_is_a_layout_error():
    doc = Document()
    doc.add_paragraph('SA2 #177 meeting session plan')
    table = doc.add_table(rows=3, cols=2)
    table.cell(1, 0).text, table.cell(1, 1).text = '6G slot 1', 'KI#1'
    data = io.BytesIO()
    doc.save(data)
    with pytest.raises(docmod.LayoutError):
        docmod.extract_document(data.getvalue(), 'S2-2610195_SA2-177-Session Plan - r7.docx')


# ------------------------------------------------------------------ lines

@pytest.mark.parametrize('raw, items, unread', [
    ('KI#18 (AI for Network) (20.6.18) [5+15]', ['20.6.18'], []),
    ('KI#12 – KI#16 (20.6.12 - 16) [5,0,2,0,1]', ['20.6.12', '20.6.13', '20.6.14', '20.6.15', '20.6.16'], []),
    ('KI#12 to KI#16 (20.6.14, 12, 13, 15, 16)', ['20.6.14', '20.6.12', '20.6.13', '20.6.15', '20.6.16'], []),
    ('Rel-19 (19.48/49/50) [37]', ['19.48', '19.49', '19.50'], []),
    ('UIA_ARC (19.8.2 [4]', ['19.8.2'], []),   # never closed
    ('NG_RTC_Ph2 [19.2.2] [15]', ['19.2.2'], []),   # an item in brackets, then a count
    ('20.48, 20.49, 20.50 [15]', ['20.48', '20.49', '20.50'], []),
    ('(5.x, 6.x, 7.x ) [9]', ['5.x', '6.x', '7.x'], []),
    ('Opening (1), Agenda (2), Reports (3)', ['1', '2', '3'], []),   # agenda.csv has 1.1, not 1
    ('Revisions (including 30.1)', ['30.1'], []),
    ('Postponed TEI20 item* (see 30.1) [2]', ['30.1'], []),
    ('Drafting: 6G Architecture drafting (Option 1, 2, 3)', [], []),
    ('Plenary session (1330 - 1630)', [], []),
    ('5.x [0], 6.7 [3], 7.x [0], eNA_Ph2 (8.1) [6],', ['8.1'], ['5.x', '6.7', '7.x']),
])
def test_agenda_items_are_read_from_the_original_line(raw, items, unread):
    assert sessions.line_items(raw, agenda()) == (items, unread)


@pytest.mark.parametrize('raw, label', [
    ('KI#18 (AI for Network) (20.6.18) [5+15]', 'KI#18'),
    ('UIA_ARC (19.8.2 [4]', 'UIA_ARC'),
    ('Opening (1), Agenda (2), Reports (3)', 'Opening, Agenda, Reports'),
    ('13:00 (Main room) Drafting for SoH questions based on rapporteur input',
     'Drafting for SoH questions based on rapporteur input'),
    ('Drafting:', 'Drafting'),
    ('K#1.2 (20.6.1) [3]', 'K#1.2'),   # the chair's typo stays
])
def test_titles_drop_every_bracket_and_keep_the_words(raw, label):
    assert sessions.topic(raw) == label


# ------------------------------------------------------------------ cells

def test_the_two_cells_of_a_quarter_are_pieces_and_the_same_text_twice_is_one():
    found = read(R20, agenda())
    left, right = find(found, 'Tuesday', '09:00', 'stream1', [1, 2]), find(found, 'Tuesday', '09:00', 'stream1', [2, 2])
    assert [s['name'] for s in left + right] == ['KI#18', 'Rel-20 LS’s and CAT B/C alignment']
    assert left[0]['end'] == '10:30'   # a piece keeps the block's times
    # The same text in both cells is the whole quarter.
    whole = find(found, 'Monday', '14:00', 'stream1')
    assert [(s['name'], s['part'], s['end']) for s in whole] == [('KI#18', None, '15:30')]
    # Spaces and paragraph breaks do not make the text different.
    assert [s['part'] for s in find(found, 'Wednesday', '09:00', 'stream2')] == [None]
    # Only one cell filled: that piece alone.
    assert [(s['name'], s['part']) for s in find(found, 'Tuesday', '09:00', 'stream3')] == [('KI#17', [2, 2])]


def test_q0_is_one_session_with_the_left_cell_first():
    found = read(R20, agenda())
    q0 = find(found, 'Thursday', '08:00', 'stream1')
    assert [(s['name'], s['end'], s['part']) for s in q0] == [('KI#1.1 / KI#1.2', '08:50', None)]
    assert q0[0]['lines'] == ['KI#1.1 (NAS) (20.6.1) [9]', 'KI#1.2 (20.6.1) [8]']


def test_topics_join_three_at_most_and_number_lines_add_to_the_topic_above():
    found = read(R20, agenda())
    opening = find(found, 'Monday', '09:00', scope='shared')[0]
    assert opening['name'].startswith('Opening, Agenda, Reports / Common Issues / ') and opening['name'].endswith(' …')
    assert opening['agenda_items'][:4] == ['1', '2', '3', '4.1']
    maintenance = find(found, 'Wednesday', '09:00', 'stream2')[0]
    assert (maintenance['name'], maintenance['agenda_items']) == ('Pre-Rel-19 maintenance', ['5.x', '6.x', '7.x'])
    single = find(found, 'Wednesday', '14:00', 'stream1')[0]
    assert (single['name'], single['lines']) == ('6G Single Stream', ['6G Single Stream (20.6.0)', 'Architecture'])


def test_streams_that_rest_and_announcements_are_no_sessions():
    doc = extract(R20)
    found, notes = sessions.interpret(doc, agenda())
    # "-" in Streams 2 and 3 during the 6G single stream.
    assert not find(found, 'Wednesday', '14:00', 'stream2') and not find(found, 'Wednesday', '14:00', 'stream3')
    assert notes == [{'day': 'Monday', 'text': 'Opening of meeting at 0900'},
                     {'day': 'Friday', 'text': 'Close of meeting by 1630'}]
    assert not find(found, 'Monday', '08:00', scope='shared')


def test_placeholders_stay_cells_and_merged_rows_are_shared():
    found = read(R9)
    revisions = find(found, 'Thursday', '18:00', scope='shared')
    assert [s['name'] for s in revisions] == ['Revisions']
    drafting = find(found, 'Tuesday', '08:00', 'stream2')
    assert [s['name'] for s in drafting] == ['Drafting']


def test_the_plenary_is_split_at_its_written_times():
    found = read(R20, agenda())
    plenary = [s for s in found if s['day'] == 'Friday' and s['scope'] == 'shared']
    assert [(s['start'], s['end'], s['name'], s['agenda_items']) for s in plenary] == [
        ('13:30', '14:00', 'Plenary session', []),
        ('14:00', '15:00', 'List of agreed tdocs for block approval / Revisions', ['30.1']),
        ('15:00', '16:30', 'block approval of agreed tdocs', [])]
    assert plenary[1]['lines'] == ['14:00 List of agreed tdocs for block approval', 'Revisions (including 30.1)']


def test_break_rows_hold_sessions_of_their_own():
    # "Drafting: (1300) FS_ARCH_enIMS (main room)" / "... (breakout 1)": one session per room.
    lunch = find(read(R12), 'Monday', '13:00')
    assert [(s['rooms'], s['end'], s['name']) for s in lunch] == [
        (['stream1'], '14:00', 'Drafting: FS_ARCH_enIMS'), (['stream2'], '14:00', 'Drafting: FS_AIML_CN_Ph2')]
    # No room written: across the rooms, unassigned.
    election = find(read(R9), 'Monday', '12:30', scope='unassigned')
    assert [(s['end'], s['name'], s['rooms']) for s in election] == [('14:00', 'SA2 VC Election', [])]


def test_cells_the_rules_are_unsure_about_say_why():
    run_together = find(read(R9), 'Monday', '18:00', 'stream1', [1, 2])[0]
    assert run_together['reasons'] == ["two topics run together: 'TEI20_DAMP [1]TEI20_UDCE-ARC [2]'"]
    loose = find(read(R12), 'Monday', '11:00', 'stream1', [1, 2])[0]
    assert loose['reasons'] == ["agenda numbers not in agenda.csv: ['5.x', '6.7', '7.x']"]
    assert sessions.uncertainty(['Start 10:45 in front of the Main room'], []) == [
        "time written in a form the rules do not read: 'Start 10:45 in front of the Main room'"]
    # Times in the forms the rules read, and TDoc numbers, are no doubt.
    assert sessions.uncertainty(['Plenary session (1330 - 1630)', '14:00 List of agreed tdocs',
                                 'Drafting: (1300) FS_ARCH_enIMS (main room)', 'LS’s, 9779, 10337, RAN-CN'], []) == []
    assert all(not s['reasons'] for s in read(R20, agenda()) if s['day'] != 'Tuesday' or s['start'] != '18:00')


# --------------------------------------------------------------- schedule

@pytest.fixture(scope='module')
def r20():
    return sessions.make_schedule(extract(R20), agenda(), META_176, [R20, 'agenda.csv'],
                                  {R20: 'Session Plan r20', 'agenda.csv': 'Agenda'}, '2026-08-28 12:00',
                                  chairs={'Mike': 'Michael STARSINIC', 'Wanqiang': 'Wanqiang ZHANG'})


def test_chairs_come_from_the_convener_line_only(r20):
    monday = next(d for d in r20.days if d.day_name == 'Monday')
    by_room = {}
    for s in monday.sessions:
        if s.start_time == '16:00':
            by_room.setdefault(tuple(s.room_ids), set()).add(s.chair)
    # Stream 1 has no name in the plan; the chair is not inferred.
    assert by_room[('stream1',)] == {None}
    assert by_room[('stream2',)] == {'Michael STARSINIC'} and by_room[('stream3',)] == {'Wanqiang ZHANG'}
    thursday = next(d for d in r20.days if d.day_name == 'Thursday')
    q0 = [s for s in thursday.sessions if s.start_time == '08:00' and s.room_ids == ['stream2']]
    assert q0[0].chair == 'Michael STARSINIC'   # Q0's stream label has no highlight; same stream
    assert all(s.chair is None for d in r20.days for s in d.sessions if s.room_scope != 'assigned')


def test_schedule_days_breaks_groups_and_notes(r20):
    assert (r20.meeting_name, r20.wg_id, r20.timezone) == ('SA2#176', 'sa2', 'Europe/Prague')
    friday = r20.days[-1]
    assert (friday.date, friday.timeline.start, friday.timeline.end) == ('2026-08-28', '08:00', '16:30')
    lunch = next(b for b in friday.timeline.breaks if b['name'] == 'Lunch')
    assert lunch['label_position'] == 'time-axis'   # the plenary starts at 13:30
    monday = r20.days[0]
    assert monday.timeline.start == '09:00'   # Q0 holds only the opening note
    tuesday = r20.days[1]
    piece = next(s for s in tuesday.sessions if s.room_ids == ['stream1'] and s.start_time == '09:00')
    assert (piece.part, piece.duration_minutes, piece.group_header) == ([1, 2], 90, 'Rel-20 SIDs and WIDs')
    assert piece.notes == ['KI#18 (AI for Network) (20.6.18)'] and piece.agenda_item == '20.6.18'
    note_table = r20.supplements[-1]['rows']
    assert ['Monday', 'Opening of meeting at 0900'] in note_table
    assert ['Thursday', '(Revisions only unless stated)'] in note_table


def test_render_and_snapshot_roundtrip(r20, tmp_path):
    save_schedule(r20, tmp_path / 'schedule.json')
    assert load_schedule(tmp_path / 'schedule.json') == r20
    html = BeautifulSoup(generate_html(r20), 'html.parser')
    piece = html.select('#tuesday [data-name="KI#18"]')[0]
    assert '1/2' in piece.get_text() and 'of' in piece.get_text()
    shared = html.select_one('#friday [data-name="List of agreed tdocs for block approval / Revisions"]')
    assert 'Revisions (including 30.1)' in shared['data-popup']
