"""Real v04 source regression plus deterministic remote/LLM failure scenarios."""
import io
import json
from pathlib import Path
from types import SimpleNamespace
import zipfile

from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
import httpx
import pytest

from shared.renderer import generate_html
from shared.schedule import load_schedule, save_schedule
from working_groups.ran_plenary import document as docmod, interpreter as interp, pipeline, sources, lifecycle

FIXTURES = Path(__file__).parent / 'fixtures/ran_plenary'
FILENAME = 'RAN#113 time plan v04.zip'


@pytest.fixture
def source():
    return docmod.extract_document((FIXTURES / FILENAME).read_bytes(), FILENAME)


@pytest.fixture
def parsed():
    return interp.Interpretation.model_validate_json((FIXTURES / 'interpretation.json').read_text())


@pytest.fixture
def metadata():
    return {'starts_on': '2026-09-14', 'ends_on': '2026-09-17', 'location': 'Madrid',
            'country': 'ES', 'timezone': 'Europe/Madrid'}


@pytest.fixture
def schedule(source, parsed, metadata):
    return interp.make_schedule(source, parsed, docmod.agenda_map((FIXTURES / 'agenda.csv').read_bytes()),
                                metadata, FILENAME, '2026-09-15 02:00')


def test_real_document_preserves_color_and_order(source):
    assert len(source['cells']) == 34
    assert len(source['topics']) == 22
    assert len(source['breaks']) == 3
    assert source['days'] == ['Monday', 'Tuesday', 'Wednesday', 'Thursday']
    assert source['breaks'][0]['days'] == source['days']
    colors = {r['color'] for c in source['cells'] for p in c['paragraphs'] for r in p['runs']}
    assert {'FF0000', 'EE0000', '00B050', '00B0F0', '000000'} <= colors
    header = source['headers'][1]
    assert 'until Tue AM' in header['text'] and 'from Tue PM' in header['text']
    assert any(r['text'] == '\n' for c in source['cells'] for p in c['paragraphs'] for r in p['runs'])
    assert source['topics'][0]['agenda_item'] == '3, 4'
    assert source['topics'][-1]['title'] == 'TEI'
    assert '1605/1605/1619' in source['topics'][2]['documents']


def test_inherited_theme_color_and_strike_are_preserved():
    doc = Document()
    doc.styles['Normal'].font.color.rgb = RGBColor.from_string('112233')
    p = doc.add_paragraph()
    p.add_run('inherited')
    run = p.add_run('theme\ncontinued')
    color = OxmlElement('w:color'); color.set(qn('w:themeColor'), 'accent1'); color.set(qn('w:val'), '000000')
    run._r.get_or_add_rPr().append(color)
    run.font.strike = True
    result = docmod.Formatting(doc).paragraph(p, 'p0')
    assert result['runs'][0]['color'] == '112233'
    assert result['runs'][1]['raw_color']['themeColor'] == 'accent1'
    assert result['runs'][1]['color'] != '000000'
    assert result['runs'][1]['strike']
    assert result['text'] == 'inheritedtheme\ncontinued'


def test_zip_rejects_wrong_or_multiple_members():
    for names in [('RAN#112 time plan v04.docx',), ('a/RAN#113 time plan v04.docx', 'b/RAN#113 time plan v04.docx')]:
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as z:
            for name in names:
                z.writestr(name, b'invalid')
        with pytest.raises(ValueError, match='exactly one'):
            docmod.unpack_docx(output.getvalue(), FILENAME)


def test_agenda_is_string_keyed_and_rejects_conflicts():
    assert docmod.agenda_map(b'"9.10","Ten"\n"9.1","One"')['9.10'] == 'Ten'
    with pytest.raises(ValueError, match='Conflicting'):
        docmod.agenda_map(b'"1","One"\n"1","Other"')
    with pytest.raises(ValueError):
        docmod.agenda_map(b'')
    assert docmod.describe('9.10.1, 9.1', {'9.10': 'Ten', '9.1': 'One'})[0]['matched_agenda_item'] == '9.10'


def test_real_interpretation_has_correct_physical_rooms_and_merged_cells(schedule, source, parsed):
    interp.validate(parsed, source)
    assert [len(day.rooms) for day in schedule.days] == [2, 3, 2, 2]
    tuesday = schedule.days[1]
    old = next(r.id for r in tuesday.rooms if 'Madrid 1' in r.name)
    new = next(r.id for r in tuesday.rooms if 'Neptuno' in r.name)
    assert any(old in s.room_ids and s.start_time == '08:00' for s in tuesday.sessions)
    assert any(new in s.room_ids and s.start_time == '12:30' and s.end_time == '13:30' for s in tuesday.sessions)
    assert not any(old in s.room_ids and s.end_time > '12:30' for s in tuesday.sessions)
    assert all(s.room_scope != 'assigned' for d in schedule.days for s in d.sessions if s.name == 'Early items')
    main_itu = next(s for s in tuesday.sessions if 'ITU' in s.name)
    assert main_itu.room_scope == 'assigned'
    assert any(s.name == 'TBD' and s.room_scope == 'assigned' for s in tuesday.sessions)
    assert schedule.days[1].timeline.end == '19:00'
    assert schedule.days[3].timeline.end == '17:00'
    assert schedule.days[0].date == '2026-09-14' and schedule.timezone == 'Europe/Madrid'
    assert len(schedule.topic_references) == 22
    assert sum(len(d.sessions) for d in schedule.days) == 51
    assert next(s for s in tuesday.sessions if 'modernization' in s.name).agenda_item is None
    assert any('3, 4' == t['agenda_item'] for t in schedule.topic_references)


def test_rendering_spans_rooms_without_claiming_them_and_preserves_notes(schedule):
    html = BeautifulSoup(generate_html(schedule), 'html.parser')
    tuesday = html.select_one('#tuesday')
    merged = tuesday.select_one('[data-room-scope="unassigned"]')
    assert 'grid-column:2/5' in merged['style']
    assert 'Room unassigned' in merged['data-popup']
    assert 'Room: Main' not in merged['data-popup']
    assert len(tuesday.select('.room-header')) == 4
    axis = tuesday.select('.break-axis')
    assert [(b.get_text(), b['title']) for b in axis] == [
        ('Coffee', 'MORNING COFFEE BREAK'), ('Lunch', 'LUNCH BREAK'), ('Coffee', 'AFTERNOON COFFEE BREAK')]
    assert 'Close by' in html.get_text() or 'Close by' in str(html)
    assert len(html.select('.topic-references tbody tr')) == 22
    assert html.select_one('.topic-references summary').get_text() == 'Topics'
    assert html.select('.topic-references > p')[-1].get_text() == (
        'Plan is to discuss the RAN#113 topics according to the ordering in the table below unless noted otherwise.'
    )
    assert not html.select('.demo-notice')


def test_snapshot_roundtrip_and_old_snapshot_defaults(schedule, tmp_path):
    path = tmp_path / 'schedule.json'
    save_schedule(schedule, path)
    assert load_schedule(path) == schedule
    assert schedule.source_labels == {FILENAME: 'Time plan', 'agenda.csv': 'Agenda'}
    raw = json.loads(path.read_text())
    raw.pop('topic_references')
    raw.pop('topic_preamble')
    raw.pop('source_labels')
    for day in raw['days']:
        for session in day['sessions']:
            session.pop('room_scope'); session.pop('notes')
    path.write_text(json.dumps(raw))
    old = load_schedule(path)
    assert all(s.room_scope == 'assigned' and s.notes == [] for d in old.days for s in d.sessions)
    assert old.source_labels == {}


@pytest.mark.parametrize('damage,match', [('missing', 'Unaccounted'), ('room', 'unavailable'),
                                         ('closing', 'closing'), ('overlap', 'Overlapping')])
def test_invalid_llm_result_is_rejected(source, parsed, damage, match):
    if damage == 'missing':
        parsed.sessions.pop(0)
    elif damage == 'room':
        old = next(r.id for r in parsed.rooms if 'Madrid 1' in r.name)
        block = next(s for s in parsed.sessions if s.day == 'Tuesday' and s.start == '12:30' and s.room_scope == 'assigned' and 'Rel-21' in s.name)
        block.room_ids = [old]
    elif damage == 'closing':
        next(s for s in parsed.sessions if s.day == 'Thursday' and s.start == '16:00').end = '18:00'
    else:
        parsed.sessions.append(parsed.sessions[0].model_copy(deep=True))
    with pytest.raises(ValueError, match=match):
        interp.validate(parsed, source)


def test_color_only_change_invalidates_interpretation_cache(source, parsed, tmp_path):
    calls = []
    def request(*args, **_):
        calls.append(args)
        return parsed.model_dump()
    interp.interpret(source, 'test-model', request, tmp_path)
    interp.interpret(source, 'test-model', request, tmp_path)
    assert len(calls) == 1
    source['cells'][0]['paragraphs'][0]['runs'][0]['color'] = '001122'
    interp.interpret(source, 'test-model', request, tmp_path)
    assert len(calls) == 2
    interp.interpret(source, 'other-model', request, tmp_path)
    assert len(calls) == 3


def test_one_validation_repair_then_failure(source, parsed, tmp_path):
    calls = []
    bad = parsed.model_dump(); bad['sessions'] = []
    def repair(*args, **_):
        calls.append(args)
        return bad if len(calls) == 1 else parsed.model_dump()
    interp.interpret(source, 'repair', repair, tmp_path)
    assert len(calls) == 2 and 'VALIDATION ERROR' in calls[1][0]
    with pytest.raises(ValueError, match='failed validation'):
        interp.interpret(source, 'broken', lambda *args, **_: bad, tmp_path)
    assert len(list(tmp_path.glob('parse-*.json'))) == 1


def test_listing_uses_numeric_meeting_and_version_and_refuses_rollback():
    html = ''.join(f'<a href="RAN%23{n}%20time%20plan%20v{v}.zip">file</a>' for n,v in [(112,99),(113,4),(113,10),(113,9)])
    chosen = sources.select_timeplan(html, 'https://example.org/Chair/', {})
    assert (chosen['meeting_number'], chosen['version']) == (113,10)
    with pytest.raises(ValueError, match='roll back'):
        sources.select_timeplan(html, 'https://example.org/', {'meeting_number':114, 'version':0})


def test_http_conditional_get_recovers_missing_body_and_detects_same_url_changes():
    calls = []
    def handler(req):
        calls.append(req)
        return httpx.Response(304) if req.headers.get('if-none-match') else httpx.Response(200, content=b'new', headers={'ETag':'v2'})
    with httpx.Client(transport=httpx.MockTransport(handler)) as http:
        body, info = sources.fetch_file(http, 'https://example.org/file',
                                        {'url':'https://example.org/file','etag':'v1','sha256':sources.digest(b'old')}, None)
    assert body == b'new' and info['sha256'] == sources.digest(b'new')
    assert len(calls) == 2


def test_valid_cached_body_accepts_304():
    with httpx.Client(transport=httpx.MockTransport(lambda req: httpx.Response(304))) as http:
        body, _ = sources.fetch_file(http, 'https://example.org/file',
                                    {'url':'https://example.org/file','etag':'v1','sha256':sources.digest(b'old')}, b'old')
    assert body == b'old'


def test_bundle_corruption_is_not_silently_accepted(tmp_path):
    b = sources.Bundle({'timeplan':{'sha256':sources.digest(b'x')},'agenda':{'sha256':sources.digest(b'y')}},b'x',b'y')
    b.save(tmp_path)
    (tmp_path/'agenda.csv').write_bytes(b'wrong')
    with pytest.raises(ValueError, match='Corrupt agenda'):
        sources.Bundle.load(tmp_path)


def test_portal_matches_meeting_number_not_ongoing_status(monkeypatch):
    rows = [{'Id':1,'TBId':373,'Title':'3GPPRAN#113','StartDate':'2026-09-14 09:00:00','EndDate':'2026-09-17 17:30:00','Location':'Madrid','Country':'ES'},
            {'Id':2,'TBId':373,'Title':'3GPPRAN#114','StartDate':'2026-12-07 09:00:00','EndDate':'2026-12-10 17:30:00','Location':'Boston','Country':'US'},
            {'Id':3,'TBId':380,'Title':'3GPPRAN#112','StartDate':'2026-06-01 09:00:00','EndDate':'2026-06-04 17:30:00'}]
    # The shared query (one per process, every WG) answers; RAN P sends none of its own.
    monkeypatch.setattr(sources, 'get_meetings', lambda: rows)
    assert sources.portal_meeting(114)['location'] == 'Boston'
    with pytest.raises(ValueError):
        sources.portal_meeting(112)


def test_metadata_fallback_cannot_cross_meetings(metadata):
    cfg = {**sources.DEFAULTS}
    previous = {'meeting_number':113, 'metadata':metadata}
    assert pipeline.meeting_metadata(113,cfg,previous,None,None) == metadata
    with pytest.raises(ValueError, match='Missing'):
        pipeline.meeting_metadata(114,cfg,previous,None,None)


def test_timezone_uses_location_and_reuses_the_venue_list(tmp_path, monkeypatch):
    calls = []
    def request(payload, *_, **__):
        calls.append(payload)
        return {'timezone':'Europe/Madrid'}
    data = {'location':'Madrid','country':'ES','portal_timezone':'GMT+02:00'}
    assert interp.resolve_timezone(data,'test',request) == 'Europe/Madrid'
    assert interp.resolve_timezone(data,'test',request) == 'Europe/Madrid'
    assert len(calls) == 1 and 'Madrid' in calls[0] and 'ES' in calls[0]


@pytest.fixture
def local_config(metadata, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = {**sources.DEFAULTS, 'local_agenda':str(FIXTURES/'agenda.csv'), 'meetings':{'113':metadata}}
    monkeypatch.setattr(pipeline, 'load_config', lambda:cfg)
    monkeypatch.setattr(sources, 'load_config', lambda:cfg)
    monkeypatch.setattr(sources, 'portal_rows', lambda: [])
    return cfg


def test_offline_pipeline_uses_real_inputs_and_retains_state_on_failure(local_config, parsed, tmp_path, monkeypatch):
    opts = SimpleNamespace(local=str(FIXTURES/FILENAME), no_download=True, output_dir=tmp_path/'out')
    monkeypatch.setattr(httpx.Client, 'send', lambda *a, **k: pytest.fail('Offline network request'))
    schedule = pipeline.build_schedule(opts, request=lambda *args, **_:parsed.model_dump())
    state = tmp_path/'out/ran-plenary/.schedule_state.json'
    before = state.read_bytes()
    assert not schedule.is_demo and schedule.timezone == 'Europe/Madrid'
    monkeypatch.setattr(pipeline, 'interpret', lambda *a: (_ for _ in ()).throw(ValueError('bad LLM')))
    with pytest.raises(ValueError):
        pipeline.build_schedule(opts)
    assert state.read_bytes() == before


def test_agenda_only_change_reuses_llm(local_config, parsed, tmp_path):
    agenda = tmp_path/'agenda.csv'; agenda.write_bytes((FIXTURES/'agenda.csv').read_bytes())
    local_config['local_agenda'] = str(agenda)
    opts = SimpleNamespace(local=str(FIXTURES/FILENAME), no_download=True)
    calls=[]
    def request(*a, **_):
        calls.append(a);return parsed.model_dump()
    first = pipeline.build_schedule(opts,request=request)
    agenda.write_text(agenda.read_text().replace('Opening of the meeting','Updated opening'))
    second = pipeline.build_schedule(opts,request=request)
    assert len(calls) == 1
    assert first.days == second.days  # AI 1 is not assigned to these timed blocks.
    assert sources.read_json(Path('docs/ran-plenary/.schedule_state.json'))['agenda']['sha256'] == sources.digest(agenda.read_bytes())


def test_lifecycle_check_stages_without_advancing_success_state(local_config, tmp_path, monkeypatch):
    bundle = sources.local_bundle(FIXTURES/FILENAME,local_config,{})
    monkeypatch.setattr(sources,'fetch_bundle',lambda *a, **k:bundle)
    before = {'meeting_number':113,'metadata':local_config['meetings']['113']}
    sources.write_json(sources.OUTPUT/'.schedule_state.json',before)
    result = lifecycle.check_updates()
    assert result.changed and not result.errors
    assert sources.read_json(sources.OUTPUT/'.schedule_state.json') == before
    lifecycle.prepare_build()
    assert sources.Bundle.load(sources.DOWNLOADS/'prepared').timeplan == bundle.timeplan
    state = {**bundle.manifest,'metadata':local_config['meetings']['113']}
    sources.write_json(sources.OUTPUT/'.schedule_state.json',state)
    assert not lifecycle.check_updates().changed
    monkeypatch.setattr(sources,'fetch_bundle',lambda *a, **k: (_ for _ in ()).throw(ValueError('network failed')))
    assert lifecycle.check_updates().errors
    assert sources.read_json(sources.OUTPUT/'.schedule_state.json') == state


def test_force_reset_preserves_successful_output(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    state=sources.OUTPUT/'.schedule_state.json';sources.write_json(state,{'meeting_number':113})
    sources.write_json(sources.CACHE/'x.json',{})
    sources.write_json(sources.DOWNLOADS/'inputs/manifest.json',{})
    lifecycle.reset_cache()
    assert state.exists() and not sources.CACHE.exists() and not sources.DOWNLOADS.exists()


def test_unsupported_agenda_inference_is_rejected(source, parsed):
    block = next(s for s in parsed.sessions if 'modernization' in s.name)
    block.agenda_item = '8.2.1';block.topic_refs = ['t1.r16']
    with pytest.raises(ValueError, match='specific textual evidence'):
        interp.validate(parsed,source)


def test_untimed_topics_cannot_be_given_invented_duration(source, parsed):
    block = next(s for s in parsed.sessions if s.name == 'Early items')
    block.end = '12:00' if block.start == '11:00' else '10:00'
    with pytest.raises(ValueError, match='invent'):
        interp.validate(parsed,source)


def test_room_names_notes_topics_and_filter_data_are_escaped(schedule):
    attack = '</script><script>alert(1)</script>'
    schedule.days[0].sessions[0].name = attack
    schedule.days[0].sessions[0].notes = [attack]
    schedule.days[0].rooms[0].name = attack
    schedule.topic_references[0]['title'] = attack
    html = generate_html(schedule)
    assert '<script>alert(1)</script>' not in html
    parsed_html = BeautifulSoup(html,'html.parser')
    data = json.loads(parsed_html.select_one('#filter-data').string)
    assert any(s['name'] == attack for g in data['groups'] for s in g['sessions'])


def test_tint_and_shade_follow_word_luminance_rules():
    doc = Document();p=doc.add_paragraph();run=p.add_run('tint')
    color=OxmlElement('w:color');color.set(qn('w:themeColor'),'accent1')
    color.set(qn('w:themeTint'),'99');color.set(qn('w:themeShade'),'BF')
    run._r.get_or_add_rPr().append(color)
    formatting = docmod.Formatting(doc);formatting.theme['accent1']='4F81BD'
    assert formatting.paragraph(p,'p')['runs'][0]['color'] == '95B3D7'


def test_direct_builder_restores_state_if_snapshot_write_fails(local_config, parsed, tmp_path, monkeypatch):
    import build
    import schedule_fixture
    from shared.schedule import save_schedule
    state=tmp_path/'docs/ran-plenary/.schedule_state.json'
    sources.write_json(state,{'old':'checkpoint'})
    before=state.read_bytes()
    save_schedule(schedule_fixture.build_schedule(),state.parent/'schedule.json')
    old_snapshot=(state.parent/'schedule.json').read_bytes()
    config=tmp_path/'site.json';config.write_text(json.dumps({'default_wg':'ran-plenary','working_groups':[{'id':'ran-plenary'}]}))
    monkeypatch.setattr(build,'build_schedule',lambda wg,opts: pipeline.build_schedule(opts,request=lambda *a, **_:parsed.model_dump()))
    monkeypatch.setattr(build,'save_schedule',lambda *a: (_ for _ in ()).throw(OSError('disk failed')))
    opts=SimpleNamespace(wg='ran-plenary',site_config=config,output_dir=tmp_path/'docs',render_only=False,
                         local=str(FIXTURES/FILENAME),no_download=True)
    with pytest.raises(RuntimeError):
        build.build_site(opts)
    assert state.read_bytes() == before
    assert (state.parent/'schedule.json').read_bytes() == old_snapshot


def test_prepared_bundle_matches_config_before_build(local_config, tmp_path):
    bundle=sources.local_bundle(FIXTURES/FILENAME,local_config,{})
    bundle.manifest['config_hash']='stale'
    bundle.save(sources.DOWNLOADS/'prepared')
    with pytest.raises(ValueError,match='Prepared inputs'):
        pipeline.build_schedule()
    assert not (sources.DOWNLOADS/'prepared').exists()


def test_final_direct_render_failure_restores_all_outputs(schedule, tmp_path, monkeypatch):
    import build
    import schedule_fixture
    config = tmp_path/'site.json'
    config.write_text(json.dumps({'default_wg':'ran-plenary','working_groups':[{'id':'ran-plenary'}]}))
    monkeypatch.chdir(tmp_path)
    save_schedule(schedule_fixture.build_schedule(),tmp_path/'docs/ran-plenary/schedule.json')
    state=tmp_path/'docs/ran-plenary/.schedule_state.json';sources.write_json(state,{'old':True})
    before={str(p):p.read_bytes() for p in (tmp_path/'docs').rglob('*') if p.is_file()}
    def build_one(*a):
        sources.write_json(state,{'advanced':True})
        return schedule
    monkeypatch.setattr(build,'build_schedule',build_one)
    monkeypatch.setattr(build,'render_site',lambda *a, **kw: (_ for _ in ()).throw(OSError('render failed')))
    opts=SimpleNamespace(wg='ran-plenary',site_config=config,output_dir=tmp_path/'docs',render_only=False)
    with pytest.raises(OSError):
        build.build_site(opts)
    assert {str(p):p.read_bytes() for p in (tmp_path/'docs').rglob('*') if p.is_file()} == before


def test_location_override_discards_cached_timezone(metadata,tmp_path,monkeypatch):
    monkeypatch.setattr(interp,'CACHE',tmp_path)
    cfg={**sources.DEFAULTS,'meetings':{'113':{'location':'Tokyo','country':'JP'}}}
    prior={'meeting_number':113,'metadata':metadata}
    result=pipeline.meeting_metadata(113,cfg,prior,None,lambda *a, **_:{'timezone':'Asia/Tokyo'})
    assert result['timezone'] == 'Asia/Tokyo'


def test_same_meeting_and_version_use_upload_time_before_file_extension():
    html = '''<table><tr><td><a href="RAN%23113%20time%20plan%20v04.docx">DOCX</a></td><td>2026/09/15 12:00</td></tr>
    <tr><td><a href="RAN%23113%20time%20plan%20v04.zip">ZIP</a></td><td>2026/09/14 12:00</td></tr></table>'''
    assert sources.select_timeplan(html,'https://example.org/',{})['name'].endswith('.docx')


def test_direct_docx_and_zip_extract_identically(source):
    docx=docmod.unpack_docx((FIXTURES/FILENAME).read_bytes(),FILENAME)
    assert docmod.extract_document(docx,FILENAME.replace('.zip','.docx')) == source


def test_topic_source_highlights_and_hyperlinks_survive_snapshot(schedule, tmp_path):
    path=tmp_path/'schedule.json';save_schedule(schedule,path)
    restored=load_schedule(path)
    html=BeautifulSoup(generate_html(restored),'html.parser')
    topics=html.select_one('.topic-references')
    links=topics.select('a[href]')
    assert len(links)==59
    assert any(a.get_text()=='1576' and a['href'].endswith('RP-261576.zip') for a in links)
    yellow=topics.select_one('mark[style="background-color:#FEF3C7"]')
    assert 'Meeting report:' in yellow.get_text()
    assert topics.select('mark[style="background-color:#DCF3E3"]')
    assert all(a['rel']==['noopener','noreferrer'] for a in links)


def test_topic_agenda_tree_preserves_ancestry_and_merges_common_parents():
    from shared.topic_references import agenda_tree
    data=docmod.describe('6.1.1, 6.1.2',{'6':'Other coordination matters','6.1':'ITU-R ad hoc',
                                       '6.1.1':'Status report','6.1.2':'Documents'})
    html=BeautifulSoup(agenda_tree(data),'html.parser')
    assert html.get_text().count('Other coordination matters')==1
    assert html.get_text().count('ITU-R ad hoc')==1
    leaves=html.select('ul ul ul > li')
    assert [' '.join(x.get_text().split()) for x in leaves]==['6.1.1 Status report','6.1.2 Documents']
    assert all('agenda-selected' in item.get('class',[]) for item in leaves)


def test_rich_reference_blocks_unsafe_urls_and_html():
    from shared.topic_references import rich_text
    topic={'documents':'fallback','rich_text':{'documents':[{'runs':[
        {'text':'<img onerror=alert(1)>','href':'javascript:alert(1)','highlight':'yellow'},
        {'text':'safe','href':'https://example.org/?x="bad"','highlight':'green'}]}]}}
    html=BeautifulSoup(rich_text(topic,'documents'),'html.parser')
    assert not html.select('img')
    assert len(html.select('a'))==1
    assert html.select_one('a')['href']=='https://example.org/?x="bad"'


def test_topic_presentation_changes_do_not_rerun_llm(source, parsed, tmp_path):
    calls=[]
    def request(*args, **_):
        calls.append(args);return parsed.model_dump()
    interp.interpret(source,'test',request,tmp_path)
    source['topics'][0]['rich_text']['documents'][0]['runs'][0]['highlight']='green'
    interp.interpret(source,'test',request,tmp_path)
    assert len(calls)==1


def test_portal_instants_reach_plenary_now_controls(source, parsed, monkeypatch):
    row = {'Id': 60799, 'TBId': 373, 'Title': '3GPPRAN#113',
           'StartDate': '2026-09-14 09:00:00', 'EndDate': '2026-09-17 17:30:00',
           'StartTimeZone': '(GMT+02.00) Brussels, Copenhagen, Madrid, Paris',
           'EndTimeZone': '(GMT+02.00) Brussels, Copenhagen, Madrid, Paris',
           'Country': 'ES', 'Location': 'Madrid'}
    with httpx.Client(transport=httpx.MockTransport(lambda req: httpx.Response(200, json=[row]))) as http:
        metadata = pipeline.meeting_metadata(113, sources.DEFAULTS, {}, http,
                                             lambda *a, **_: pytest.fail('Known Portal zone must not call LLM'))
    assert metadata['starts_at'] == '2026-09-14T09:00:00+02:00'
    assert metadata['ends_at'] == '2026-09-17T17:30:00+02:00'
    schedule = interp.make_schedule(source, parsed, {}, metadata, FILENAME, 'test')
    assert schedule.starts_at == metadata['starts_at']
    assert schedule.ends_at == metadata['ends_at']
    html = generate_html(schedule)
    assert 'const MEETING_START_MS = null' not in html
    assert 'const MEETING_END_MS = null' not in html


def test_date_override_discards_stale_portal_instant(metadata):
    previous = {'meeting_number': 113, 'metadata': {**metadata,
                'starts_at': '2026-09-14T09:00:00+02:00', 'ends_at': '2026-09-17T17:30:00+02:00'}}
    cfg = {**sources.DEFAULTS, 'meetings': {'113': {'starts_on': '2026-09-15'}}}
    result = pipeline.meeting_metadata(113, cfg, previous, None, None)
    assert 'starts_at' not in result
    assert result['ends_at'] == previous['metadata']['ends_at']


@pytest.mark.parametrize('paragraphs', [
    ['Changed heading', 'Discuss B before A <unless amended> & confirmed.\nSecond line.'],
    [],
])
def test_topic_preamble_comes_from_document(paragraphs, parsed, metadata, tmp_path):
    doc = Document(io.BytesIO(docmod.unpack_docx((FIXTURES / FILENAME).read_bytes(), FILENAME)))
    table = doc.tables[1]
    for sibling in list(table._tbl.itersiblings(preceding=True)):
        if sibling.tag == qn('w:tbl'):
            break
        sibling.getparent().remove(sibling)
    for text in (paragraphs or ['', '']):
        paragraph = doc.add_paragraph(text)
        table._tbl.addprevious(paragraph._p)
    data = io.BytesIO()
    doc.save(data)
    source = docmod.extract_document(data.getvalue(), 'RAN#113 time plan v04.docx')
    assert source['topic_preamble'] == paragraphs
    result = interp.make_schedule(source, parsed, docmod.agenda_map((FIXTURES / 'agenda.csv').read_bytes()),
                                  metadata, FILENAME, '2026-09-15 02:00')
    path = tmp_path / 'schedule.json'
    save_schedule(result, path)
    restored = load_schedule(path)
    assert restored.topic_preamble == paragraphs
    html = BeautifulSoup(generate_html(restored), 'html.parser')
    assert [p.get_text(separator='\n') for p in html.select('.topic-references > p')] == paragraphs
    assert len(html.select('.topic-references tbody tr')) == 22
    assert not html.select('.topic-references > p unless')


# ------------------------------------------------- sync folders come and go

CHAIR = sources.DEFAULTS['chair_url']
SYNC_AGENDA = sources.DEFAULTS['agenda_url']
MEETING_FOLDER = sources.DEFAULTS['archive_url'] + 'TSGR_113/'
MEETING_AGENDA = MEETING_FOLDER + 'Agenda/agenda.csv'
# Portal names the meeting's Docs/ subfolder on ftp.3gpp.org.
ROWS_113 = [{'Title': '3GPPRAN#113', 'TBId': 373, 'Id': 1, 'StartDate': '2026-09-14 09:00:00',
             'EndDate': '2026-09-17 17:00:00',
             'MtgDocURL': 'https://ftp.3gpp.org/tsg_ran/TSG_RAN/TSGR_113/Docs/'}]
V09 = 'RAN%23113%20time%20plan%20v09.zip'


def chair_listing(*names):
    rows = ''.join(f'<tr><td><a href="{CHAIR}{n}">{n}</a></td><td>2026/09/17 10:00</td></tr>' for n in names)
    return f'<table>{rows}</table>'


def server(files, seen=None):
    """3GPP-like: listings and files from files, 403 (IIS 'not found') for anything else."""
    def handler(request):
        url = str(request.url)
        if seen is not None:
            seen.append(url)
        if url in files:
            body = files[url]
            return httpx.Response(200, text=body) if isinstance(body, str) else httpx.Response(200, content=body)
        return httpx.Response(403)
    return httpx.Client(transport=httpx.MockTransport(handler))


@pytest.fixture
def built(tmp_path, monkeypatch):
    """State and cached bodies of a RAN#113 v09 build, as CI keeps them."""
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')
    with server({CHAIR + V09: b'PK-v09', SYNC_AGENDA: b'"1","Opening"\n'}) as http:
        _, plan = sources.remote_files.fetch_file(http, CHAIR + V09)
        _, agenda = sources.remote_files.fetch_file(http, SYNC_AGENDA)
    sources.remote_files.new_run()
    return {'meeting_number': 113, 'version': 9, 'config_hash': 'x',
            'timeplan': {**plan, 'name': 'RAN#113 time plan v09.zip'}, 'agenda': agenda}


@pytest.mark.parametrize('listing', [None, chair_listing(), chair_listing('RAN%23113%20time%20plan%20v08.zip')])
def test_a_missing_or_older_chair_listing_keeps_the_built_timeplan(built, tmp_path, listing):
    files = {} if listing is None else {CHAIR: listing}
    with server(files) as http:
        bundle = sources.fetch_bundle(sources.DEFAULTS, built, http, tmp_path / 'inputs')
    assert (bundle.manifest['version'], bundle.timeplan) == (9, b'PK-v09')
    assert bundle.manifest['timeplan'] == built['timeplan'] and bundle.agenda == b'"1","Opening"\n'


def test_without_a_build_a_missing_chair_folder_is_still_an_error(tmp_path, monkeypatch):
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')
    with server({}) as http, pytest.raises(httpx.HTTPStatusError):
        sources.fetch_bundle(sources.DEFAULTS, {}, http, tmp_path / 'inputs')


def test_agenda_comes_from_the_meeting_folder_then_sync_then_the_build(built, tmp_path):
    v10 = 'RAN%23113%20time%20plan%20v10.zip'
    base = {CHAIR: chair_listing(V09, v10), CHAIR + v10: b'PK-v10'}
    seen = []
    with server({**base, MEETING_AGENDA: b'meeting', SYNC_AGENDA: b'sync'}, seen) as http:
        bundle = sources.fetch_bundle(sources.DEFAULTS, built, http, tmp_path / 'inputs', rows=ROWS_113)
    # Portal's MtgDocURL names the folder: nothing is listed to find it.
    assert bundle.agenda == b'meeting' and SYNC_AGENDA not in seen and MEETING_FOLDER not in seen
    sources.remote_files.new_run()
    with server(base) as http:
        bundle = sources.fetch_bundle(sources.DEFAULTS, built, http, tmp_path / 'inputs')
    assert bundle.timeplan == b'PK-v10' and bundle.agenda == b'"1","Opening"\n'
    sources.remote_files.new_run()
    with server({**base, SYNC_AGENDA: b'sync'}) as http:
        assert sources.fetch_bundle(sources.DEFAULTS, built, http, tmp_path / 'inputs').agenda == b'sync'
    # A new meeting does not inherit the last meeting's agenda.
    sources.remote_files.new_run()
    v01 = 'RAN%23114%20time%20plan%20v01.zip'
    with server({CHAIR: chair_listing(v01), CHAIR + v01: b'PK-114'}) as http, pytest.raises(ValueError, match='RAN#114'):
        sources.fetch_bundle(sources.DEFAULTS, built, http, tmp_path / 'inputs')


def test_without_portal_the_meeting_folder_is_found_by_name(built, tmp_path):
    v10 = 'RAN%23113%20time%20plan%20v10.zip'
    files = {CHAIR: chair_listing(V09, v10), CHAIR + v10: b'PK-v10',
             MEETING_FOLDER: '<table></table>', MEETING_AGENDA: b'meeting', SYNC_AGENDA: b'sync'}
    seen = []
    with server(files, seen) as http:
        bundle = sources.fetch_bundle(sources.DEFAULTS, built, http, tmp_path / 'inputs')
    assert bundle.agenda == b'meeting' and sources.DEFAULTS['archive_url'] not in seen
    # An override wins over the name, and is not hashed: it does not change how the agenda is read.
    other = 'https://www.3gpp.org/ftp/tsg_ran/TSG_RAN/Moved/TSGR_113/'
    cfg = {**sources.DEFAULTS, 'meeting_specific_override': [other]}
    sources.remote_files.new_run()
    with server({**files, other + 'Agenda/agenda.csv': b'moved'}) as http:
        assert sources.fetch_bundle(cfg, built, http, tmp_path / 'inputs').agenda == b'moved'
    assert sources.config_hash(cfg) == sources.config_hash(sources.DEFAULTS)


def test_venue_list_wins_over_last_build_zone_which_is_only_a_fallback(metadata):
    import json
    from shared import timezones
    cfg = {**sources.DEFAULTS}
    live = {'starts_on': '2026-09-14', 'ends_on': '2026-09-17', 'location': 'Madrid', 'country': 'ES',
            'portal_timezone': '(GMT+02:00) Somewhere new'}
    previous = {'meeting_number': 113, 'metadata': {**live, 'timezone': 'Europe/Paris'}}
    def broken(*_, **__):
        raise ValueError('no key')
    # Nothing resolves the venue: the last build's zone for the same venue is kept.
    assert pipeline.meeting_metadata(113, cfg, previous, None, broken, live)['timezone'] == 'Europe/Paris'
    # A line in the venue list (e.g. a correction) is used instead.
    timezones.MAP_PATH.write_text(json.dumps({'label': live['portal_timezone'], 'country': 'ES',
                                              'location': 'Madrid', 'timezone': 'Europe/Madrid'}) + '\n')
    assert pipeline.meeting_metadata(113, cfg, previous, None, broken, live)['timezone'] == 'Europe/Madrid'
    # Another venue never borrows the old zone.
    with pytest.raises(ValueError):
        pipeline.meeting_metadata(113, cfg, previous, None, broken, {**live, 'location': 'Lyon', 'country': 'FR'})
