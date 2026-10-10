"""RAN2 session schedules: real chair documents plus the rules they rely on."""
import io
import json
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace

from bs4 import BeautifulSoup
from docx import Document
import httpx
import pytest

from shared.renderer import generate_html
from shared.schedule import load_schedule, save_schedule, time_to_minutes
from working_groups.ran2 import document as docmod, lifecycle, pipeline, sessions, sources

FIXTURES = Path(__file__).parent / 'fixtures/ran2'
V11 = 'R2_135_Schedule_v11.docx'
BIS = 'R2_135b_Schedule_v00.docx'
OLD = 'R2_131bis_Schedule v19.docx'
META_135 = {'starts_on': '2026-08-24', 'ends_on': '2026-08-28', 'timezone': 'Europe/Amsterdam',
            'starts_at': '2026-08-24T09:00:00+02:00', 'ends_at': '2026-08-28T17:30:00+02:00'}
META_131BIS = {'starts_on': '2025-10-13', 'ends_on': '2025-10-17', 'timezone': 'Europe/Prague'}


def extract(name):
    return docmod.extract_document((FIXTURES / name).read_bytes(), name)


def agenda(name):
    return docmod.agenda_map((FIXTURES / name).read_bytes())


@pytest.fixture(scope='module')
def v11():
    return sessions.make_schedule(extract(V11), agenda('agenda_135.csv'), META_135,
                                  [V11, 'agenda.csv'], '2026-08-28 12:00')


def find(schedule, day, room, start):
    day = next(d for d in schedule.days if d.day_name == day)
    return next(s for s in day.sessions if s.room_ids == [room] and s.start_time == start)


def synthetic(cells, slots=(('08:30', '10:30'), ('11:00', '13:00'))):
    """A document dict with Monday cells: (room, start, end, lines)."""
    rooms = [{'id': i, 'name': n, 'label': n} for i, n in
             [('main', 'Main'), ('brk1', 'Breakout 1'), ('brk2', 'Breakout 2'), ('brk3', 'Breakout 3')]]
    return {'meeting_id': 'ran2#135', 'title': 'RAN2-135 Session Schedule', 'rooms': rooms,
            'days': [{'day': 'Monday', 'slots': list(slots), 'notes': [],
                      'cells': [{'id': f'c{i}', 'rooms': [room], 'start': start, 'end': end, 'lines': lines}
                                for i, (room, start, end, lines) in enumerate(cells)]}],
            'breaks': [], 'supplements': []}


def blocks(document, agenda_items=None):
    return [(s['start'], s['end'], s['rooms'][0], s['name'], s['chair'], s['agenda_items'])
            for s in sessions.interpret(document, agenda_items)]


# ------------------------------------------------------------------ documents

@pytest.mark.parametrize('name, expected', [
    ('R2_135_Schedule_v11.docx', ('ran2#135', 11)),
    ('R2_135b_Schedule_v00.docx', ('ran2#135bis', 0)),
    ('R2_131bis_Schedule v19.docx', ('ran2#131bis', 19)),
    ('R2_132_Schedule v15.docx', ('ran2#132', 15)),
    ('R2_135_Agenda_v02.docx', None),
    ('R2-2605436.zip', None),
])
def test_schedule_filenames(name, expected):
    assert docmod.file_info(name) == expected


def supplement(document, heading):
    """The block right after a heading in the supplements."""
    blocks = document['supplements']
    index = next(i for i, b in enumerate(blocks) if b['type'] == 'heading' and b['text'] == heading)
    return blocks[index + 1]


def test_real_table_rooms_days_and_breaks():
    document = extract(V11)
    assert document['meeting_id'] == 'ran2#135'
    assert [r['name'] for r in document['rooms']] == ['Main', 'Breakout 1', 'Breakout 2', 'Breakout 3']
    assert [d['day'] for d in document['days']] == ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    thursday = document['days'][3]
    assert thursday['slots'] == [('08:30', '10:30'), ('10:50', '12:50'), ('14:15', '16:15'), ('16:40', '18:30')]
    assert [b['name'] for b in document['breaks']] == ['Morning coffee', 'Lunch', 'Afternoon coffee']


def test_text_outside_the_table_is_kept_in_document_order():
    document = extract(V11)
    kinds = [(b['type'], b.get('text')) for b in document['supplements'] if b['type'] != 'table']
    assert kinds[:3] == [('heading', 'Dates and deadlines'),
                         ('paragraph', 'NOTE that this schedule may be modified on short notice.\n'
                                       'Some Expectations: Details may be added every day. The Schedule for CBs on '
                                       'Thursday (and Friday) will be updated on Wednesday, and the schedule for CBs '
                                       'on Friday will be further updated on Thursday.'),
                         ('paragraph', '* Offline discussions should be well scoped and only 30mins in duration.')]
    assert 'RAN2-135 Session Schedule' not in [text for _, text in kinds]   # the title is the page heading
    assert supplement(document, 'Dates and deadlines')['rows'] == [['August 14 10:00 UTC', 'Tdoc Submission Deadline.']]
    assert supplement(document, 'Breaks')['rows'][1] == ['Lunch:', '13:00 to 14:30']
    offline = supplement(document, 'List of Offline Face to Face discussions')
    assert offline['header'] and offline['rows'][0] == ['Number', 'Title', 'Day/Time', 'Place', 'Coordinator']
    assert len(offline['rows']) == 18
    # A wrapped title continues on the next tab-aligned line.
    assert offline['rows'][2] == ['[004]', 'Spectrum aggregation: Extract simulation results and observations '
                                  'based on submitted documents', 'Tue 11:00-12:00', 'BO2', 'Henning Wiemann (Ericsson)']
    notes = supplement(document, 'Notes in the schedule table')['rows']
    assert notes == [['Day', 'Note'], ['Thursday', 'Colorful Polo Day'],
                     ['Thursday', 'Social event – end at 18:30 (times to be readjusted for this day)']]


def test_an_empty_offline_list_still_shows_its_header_and_unnumbered_rows_keep_their_columns():
    header = supplement(extract(BIS), 'List of Offline Face to Face discussions')
    assert header == {'type': 'table', 'header': True, 'rows': [['Number', 'Title', 'Day/Time', 'Place', 'Coordinator']]}
    rows = supplement(extract(OLD), 'List of Offline Face to Face discussions')['rows']
    assert ['', '[AIoT] MAC open issues offline', 'Tue 10:30-11:30', 'BO3', 'Rui Wang (Huawei)'] in rows


def test_tracked_changes_are_read_in_accepted_form():
    friday = extract(V11)['days'][4]
    main = next(c for c in friday['cells'] if c['rooms'] == ['main'] and c['start'] == '08:30')
    assert 'All other UP CBs' in main['lines']   # inside <w:ins>; python-docx .text drops it
    old_friday = extract(OLD)['days'][4]
    closing = next(c for c in old_friday['cells'] if c['rooms'] == ['main'] and c['start'] == '11:00')
    assert closing['lines'][1].startswith('@11:00-12:00')   # the deleted "-1" is gone


def test_a_cell_merged_across_two_slots_keeps_the_whole_span():
    monday = extract(OLD)['days'][0]
    main = next(c for c in monday['cells'] if c['rooms'] == ['main'])
    assert (main['start'], main['end']) == ('09:00', '13:00')


def docx_bytes(rows, header=('', 'Main room', 'Brk 1 room', 'Brk 2 room', 'Brk 3 room*'),
               title='RAN2-135\tSession Schedule', after=()):
    doc = Document()
    if title:
        doc.add_paragraph(title)
    table = doc.add_table(rows=1, cols=len(header))
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        if isinstance(row, str):
            merged = cells[0].merge(cells[-1])
            merged.text = row
        else:
            for cell, text in zip(cells, row):
                cell.text = text
    for text in after:
        doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_missing_weekday_label_is_estimated_and_said_so():
    data = docx_bytes(['Monday', ['09:00 – 10:30', 'Opening', '', '', ''],
                       ['14:30-16:30', 'Afternoon', '', '', ''], '',
                       ['08:30 – 10:30', 'Next morning', '', '', '']])
    document = docmod.extract_document(data, 'R2_135_Schedule_v00.docx')
    assert [(d['day'], d['slots']) for d in document['days']] == [
        ('Monday', [('09:00', '10:30'), ('14:30', '16:30')]), ('Tuesday', [('08:30', '10:30')])]
    notes = supplement(document, 'Notes in the schedule table')['rows']
    assert ['Tuesday', 'Weekday label missing in the table; rows from 08:30 are shown as Tuesday.'] in notes


def test_a_different_layout_degrades_to_notes_instead_of_failing():
    # Another author: no "Main" column, a row without a time, a typo time, no title paragraph.
    data = docx_bytes(['Monday', ['09:00-10:30', 'Opening (Diana)', 'Topic A', ''],
                       ['TBD', 'Something unplaced', '', ''],
                       ['25:00-26:00', 'Bad clock', '', ''],
                       ['11:00-12:30', '', '', 'Topic B']],
                      header=('Time', 'Plenary', 'Room A', 'Room B'), title=None,
                      after=['Contacts', 'Chair\tDiana Pani'])
    document = docmod.extract_document(data, 'R2_136_Schedule_v00.docx')
    assert document['meeting_id'] == 'ran2#136'
    assert [r['name'] for r in document['rooms']] == ['Plenary', 'Room A', 'Room B']
    found = blocks(document)
    assert [(b[0], b[1], b[2], b[3]) for b in found] == [
        ('09:00', '10:30', 'plenary', 'Opening'), ('09:00', '10:30', 'room-a', 'Topic A'),
        ('11:00', '12:30', 'room-b', 'Topic B')]
    notes = supplement(document, 'Notes in the schedule table')['rows'][1:]
    assert [n[0] for n in notes] == ['Monday', 'Monday']
    assert 'Something unplaced' in notes[0][1] and 'Bad clock' in notes[1][1]
    assert document['supplements'][:2] == [{'type': 'paragraph', 'text': 'Contacts', 'bold': False},
                                           {'type': 'table', 'header': False, 'rows': [['Chair', 'Diana Pani']]}]


def test_the_filename_decides_the_meeting_when_the_title_was_not_updated():
    data = docx_bytes(['Monday', ['09:00-10:30', 'x', '', '', '']])   # title still says RAN2-135
    assert docmod.extract_document(data, 'R2_136_Schedule_v01.docx')['meeting_id'] == 'ran2#136'
    assert docmod.extract_document(data, 'my copy.docx')['meeting_id'] == 'ran2#135'
    with pytest.raises(ValueError, match='Cannot tell the RAN2 meeting'):
        docmod.extract_document(docx_bytes(['Monday', ['09:00-10:30', 'x', '', '', '']], title=None), 'x.docx')
    with pytest.raises(ValueError, match='No table with time ranges'):
        docmod.extract_document(docx_bytes(['Monday', ['all day', 'x', '', '', '']]), 'R2_136_Schedule_v01.docx')


def test_agenda_csv_skips_rows_it_cannot_read():
    data = '"1","Opening"\n"x","junk"\n"7.4","LP-WUS"\n"7.4"\n'.encode()
    assert docmod.agenda_map(data) == {'1': 'Opening', '7.4': 'LP-WUS'}


# ---------------------------------------------------------------- cell rules

def test_markers_split_a_cell():
    document = synthetic([('brk1', '08:30', '10:30', [
        '@8:30-9:30', '[7.1] NR19 AI/ML PHY [0] (Erlin)', '@9:30-10:30', '[8.1] NR20 AI/M PHY [2] (Erlin)',
        '[8.1.1]', '[8.1.2] if time allows'])])
    assert blocks(document) == [
        ('08:30', '09:30', 'brk1', 'NR19 AI/ML PHY', 'Erlin', ['7.1']),
        ('09:30', '10:30', 'brk1', 'NR20 AI/M PHY', 'Erlin', ['8.1', '8.1.1', '8.1.2'])]


def test_a_header_written_after_its_marker_is_not_moved():
    # RAN2#133bis wrote some headers above their time. Guessing that would
    # misplace other cells, so the text stays where it is written.
    document = synthetic([('brk1', '17:30', '19:30', [
        '[7.1] NR19 AI/ML PHY [0] (Erlin)', '@17:30-18:30', '[8.1] NR20 AI/M PHY [1] (Erlin)',
        '@18:30-19:30', '[8.1.1]', '[8.1.2]'])], slots=[('17:30', '19:30')])
    assert blocks(document) == [
        ('17:30', '18:30', 'brk1', 'NR19 AI/ML PHY / NR20 AI/M PHY', 'Erlin', ['7.1', '8.1']),
        ('18:30', '19:30', 'brk1', 'AI 8.1.1, 8.1.2', None, ['8.1.1', '8.1.2'])]


def test_from_and_bare_clock_markers_but_no_prose_times():
    document = synthetic([
        ('brk2', '14:30', '16:30', ['14:30-15:30 [004] (Xiaomi)', '', 'From 15:30:',
                                    '[8.8] E-UTRA TN to NR NTN HO (Sergio)']),
        ('brk3', '11:00', '13:00', ['12:00 [9.3.2.5] CA offline (Mattias)']),
        ('brk1', '17:00', '19:00', ['[8.2] NR20 AIoT [2] (Nathan)', 'Overflow from afternoon session, end by 18:30']),
    ], slots=[('11:00', '13:00'), ('14:30', '16:30'), ('17:00', '19:00')])
    found = {(b[0], b[1], b[2], b[3]) for b in blocks(document)}
    assert found == {('14:30', '15:30', 'brk2', '[004] (Xiaomi)'), ('15:30', '16:30', 'brk2', 'E-UTRA TN to NR NTN HO'),
                     ('12:00', '13:00', 'brk3', 'CA offline'), ('17:00', '19:00', 'brk1', 'NR20 AIoT')}


def test_agenda_items_budgets_offline_numbers_and_ranges():
    line = sessions.analyze('[6.0.2.1] - [6.0.2.4], [6.0.2.14]', set())
    assert line.ais == ['6.0.2.1', '6.0.2.2', '6.0.2.3', '6.0.2.4', '6.0.2.14'] and line.lead
    line = sessions.analyze('[8.3] NR20 AI mobility [1.5] (Kyeongin)', {'Kyeongin'})
    assert (line.text, line.ais, line.chairs) == ('NR20 AI mobility', ['8.3'], ['Kyeongin'])
    assert sessions.analyze('[011] [9.4.1] Mobility offline (Jedrzej)', set()).offline == ['011']
    line = sessions.analyze('[8.2.1 Organizational', set())   # unclosed in RAN2#134
    assert (line.text, line.ais) == ('Organizational', ['8.2.1'])
    line = sessions.analyze('[605] [SONMDT] Rel19 37.320 SON MDT corrections', set())
    assert line.ais == [] and line.offline == ['605'] and '[SONMDT]' in line.text
    line = sessions.analyze('6.0.2.4, 5.1.3.2, 6.0.2', set(), {'6.0.2.4': 'a', '5.1.3.2': 'b', '6.0.2': 'c'})
    assert line.ais == ['6.0.2.4', '5.1.3.2', '6.0.2']
    # Unbracketed numbers inside prose are not agenda items.
    assert sessions.analyze('Breakout to start after completion of 7.0 and 8.0', set()).ais == []
    assert sessions.analyze('- 8.6.1 Organizational', set()).ais == []
    line = sessions.analyze('(if time allows) [6.0.2.16] R18 XR (Dawid)', {'Dawid'})
    assert (line.text, line.ais, line.lead) == ('R18 XR', ['6.0.2.16'], True)
    assert sessions.analyze('[8.12.1] (e)RedCap Less than 5 MHz', set()).text == '(e)RedCap Less than 5 MHz'


def test_only_names_in_parentheses_are_chairs():
    document = synthetic([('main', '08:30', '10:30', ['Session report from Mattias', '[9.3.3] Common CP/UP']),
                          ('brk1', '08:30', '10:30', ['[7.10] NR19 SONMDT [0] (Mattias)']),
                          ('brk2', '08:30', '10:30', ['11:00-12:00 [004] (Ericsson, Nokia)']),
                          ('brk3', '08:30', '10:30', ['CB Kyeongin', 'R19 NES comebacks'])])
    found = {b[3]: b[4] for b in blocks(document)}
    assert found == {'Session report from Mattias / Common CP/UP': None, 'NR19 SONMDT': 'Mattias',
                     '[004] (Ericsson, Nokia)': None, 'CB Kyeongin': None}
    assert sessions.learn_chairs(document) == {'Mattias'}


def test_untimed_sub_row_fills_the_time_after_a_timed_offline():
    document = synthetic([('brk2', '10:50', '12:50', ['UP offline', '10:50-11:50 [009] (InterDigital)']),
                          ('brk2', '10:50', '12:50', ['CB Mattias', 'CB NR19 SONMDT [0] (Mattias)'])],
                         slots=[('10:50', '12:50')])
    assert [(b[0], b[1], b[3], b[4]) for b in blocks(document)] == [
        ('10:50', '11:50', 'UP offline', None), ('11:50', '12:50', 'CB Mattias / CB NR19 SONMDT', 'Mattias')]


def test_a_blank_line_ends_an_explicitly_timed_part():
    document = synthetic([('brk1', '08:30', '10:30', [
        'CB Kyeongin', '@8:30-9:30', 'R20 AI Mob comebacks/ and continue [8.3]', '', 'CB Erlin',
        '[8.1] NR20 AI/M PHY [1] (Erlin)'])])
    assert [(b[0], b[1], b[3], b[4]) for b in blocks(document)] == [
        ('08:30', '09:30', 'CB Kyeongin', None), ('09:30', '10:30', 'CB Erlin / NR20 AI/M PHY', 'Erlin')]


def test_marker_like_agenda_numbers_are_content():
    document = synthetic([('main', '08:30', '10:30', ['8.10 NR20 MIMO', '@25:00 nonsense'])])
    assert [(b[0], b[1]) for b in blocks(document)] == [('08:30', '10:30')]


# ------------------------------------------------------ real schedule output

def test_real_schedule_days_breaks_and_key_blocks(v11):
    assert [(d.day_name, d.date) for d in v11.days] == [
        ('Monday', '2026-08-24'), ('Tuesday', '2026-08-25'), ('Wednesday', '2026-08-26'),
        ('Thursday', '2026-08-27'), ('Friday', '2026-08-28')]
    thursday = v11.days[3]
    assert [(b['name'], b['start'], b['end']) for b in thursday.timeline.breaks] == [
        ('Morning coffee', '10:30', '10:50'), ('Lunch', '12:50', '14:15'), ('Afternoon coffee', '16:15', '16:40')]
    assert [r.name for r in v11.days[0].rooms] == ['Main', 'Breakout 1', 'Breakout 2']

    opening = find(v11, 'Monday', 'main', '09:00')
    assert opening.chair == 'Diana' and opening.agenda_item.startswith('1, 2, 3, 6.0')
    assert opening.group_header == 'NR Rel-18'
    eutra = find(v11, 'Monday', 'brk1', '09:00')
    assert eutra.name == 'EUTRA&NR15161718' and eutra.chair == 'Mattias'
    assert {'6.0.2.1', '6.0.2.7', '6.0.2.12'} <= set(eutra.agenda_item.split(', '))
    offline = find(v11, 'Tuesday', 'brk2', '11:00')
    assert (offline.name, offline.end_time, offline.chair, offline.group_header) == (
        '[004] (Ericsson, Nokia)', '12:00', None, 'Offline')
    cb = find(v11, 'Thursday', 'brk2', '11:50')
    assert cb.name.startswith('CB Mattias / CB EUTRA&NR15161718') and cb.end_time == '12:50'
    xr = find(v11, 'Wednesday', 'brk1', '17:00')
    assert xr.agenda_item == '7.7, 8.5' and xr.chair == 'Dawid'
    late = find(v11, 'Tuesday', 'brk1', '10:00')
    assert (late.name, late.end_time) == ('Offline disc on RRC details', '11:00')
    # The offline list is shown below the grid, never merged into it.
    assert not any(s.name.startswith('[203]') for d in v11.days for s in d.sessions)


@pytest.mark.parametrize('name, agenda_name, metadata', [
    (V11, 'agenda_135.csv', META_135),
    (BIS, 'agenda_135bis.csv', {'starts_on': '2026-10-12', 'ends_on': '2026-10-16', 'timezone': 'Asia/Seoul'}),
    (OLD, None, META_131BIS),
])
def test_every_room_shows_one_block_at_a_time(name, agenda_name, metadata):
    schedule = sessions.make_schedule(extract(name), agenda(agenda_name) if agenda_name else None, metadata,
                                      [name], '2026-10-05 12:00')
    for day in schedule.days:
        start, end = time_to_minutes(day.timeline.start), time_to_minutes(day.timeline.end)
        for room in day.rooms:
            spans = sorted((time_to_minutes(s.start_time), time_to_minutes(s.end_time))
                           for s in day.sessions if room.id in s.room_ids)
            assert all(a_end <= b_start for (_, a_end), (b_start, _) in zip(spans, spans[1:])), (day.day_name, room.id)
            assert all(start <= a < b <= end and a % 5 == 0 and b % 5 == 0 for a, b in spans)
    html = BeautifulSoup(generate_html(schedule), 'html.parser')
    assert len(html.select('.session-block')) == sum(len(d.sessions) for d in schedule.days)


def test_render_roundtrip_popup_and_additional_information(v11, tmp_path):
    save_schedule(v11, tmp_path / 'schedule.json')
    assert load_schedule(tmp_path / 'schedule.json') == v11
    html = BeautifulSoup(generate_html(v11), 'html.parser')
    block = html.select_one('#monday [data-name="R17/18 NR / IoT NTN / R17 NR NTN corrections / R18 NR NTN corrections …"]')
    popup = block['data-popup']
    notes = BeautifulSoup(popup, 'html.parser').select_one('.popup-notes')
    assert '[7.8] NR19 NR NTN [0] (Sergio)' in notes.get_text('\n').split('\n')
    assert 'Note:' not in popup                    # one block, no label on every line
    assert '[7.8.1], [7.8.2]' not in popup         # repeated by the AI field
    # A two-hour block has room to show the notes in the cell too.
    cell = [n.get_text() for n in block.select('.session-notes .session-note')]
    assert cell[0] == 'R17/18 NR / IoT NTN (Sergio)'
    assert '[7.8] NR19 NR NTN [0] (Sergio)' in cell
    assert '7.8: NTN for NR Ph3' in popup.replace('<strong>', '').replace('</strong>', '')
    assert popup.index('popup-notes') < popup.index('popup-description')   # notes before agenda names
    section = html.select_one('details.supplements')
    assert section.summary.get_text() == 'Additional information'
    assert [h.get_text() for h in section.select('h3')] == [
        'Dates and deadlines', 'Breaks', 'List of Offline Face to Face discussions', 'Notes in the schedule table']
    offline = section.select('table')[2]
    assert [th.get_text() for th in offline.select('thead th')] == ['Number', 'Title', 'Day/Time', 'Place', 'Coordinator']
    assert len(offline.select('tbody tr')) == 17


def test_days_outside_the_meeting_dates_are_shown_without_a_date():
    schedule = sessions.make_schedule(extract(V11), None, {**META_135, 'ends_on': '2026-08-27'}, [V11], 'now')
    assert [(d.day_name, d.date) for d in schedule.days][-1] == ('Friday', None)


# ------------------------------------------------------------------- sources

def listing(names, base='https://www.3gpp.org/ftp/x/'):
    rows = ''.join(f'<tr><td><a href="{base}{n}">{n}</a></td><td>2026/09/30 8:32</td></tr>' for n in names)
    return f'<table>{rows}</table>'


def test_listing_parses_names_times_and_meeting_folders():
    entries = sources.listing_entries(listing(['TSGR2_135bis/', 'TSGR2_133b/', 'TSGR2_99/', 'Invitation/']),
                                      'https://www.3gpp.org/ftp/tsg_ran/WG2_RL2/')
    assert entries[0]['uploaded_at'] == '2026-09-30 08:32'
    assert sources.archive_folders(entries) == {
        'ran2#135bis': 'https://www.3gpp.org/ftp/x/TSGR2_135bis/',
        'ran2#133bis': 'https://www.3gpp.org/ftp/x/TSGR2_133b/',
        'ran2#99': 'https://www.3gpp.org/ftp/x/TSGR2_99/'}
    assert sources.rank('ran2#135') < sources.rank('ran2#135bis') < sources.rank('ran2#136')


def candidate(meeting, version, url='https://www.3gpp.org/ftp/tsg_ran/x', uploaded='2026-09-30 08:32'):
    return {'meeting_id': meeting, 'version': version, 'url': url, 'uploaded_at': uploaded, 'name': f'{meeting}-{version}'}


def test_the_current_meeting_wins_over_an_early_next_one():
    dates = {'ran2#135': ('2026-08-24', '2026-08-28'), 'ran2#135bis': ('2026-10-12', '2026-10-16'),
             'ran2#136': ('2026-11-16', '2026-11-20')}
    found = [candidate('ran2#135', 11), candidate('ran2#135bis', 0), candidate('ran2#136', 0)]
    assert sources.choose_schedule(found, dates, '2026-10-05')['meeting_id'] == 'ran2#135bis'
    assert sources.choose_schedule(found, dates, '2026-10-14')['meeting_id'] == 'ran2#135bis'
    assert sources.choose_schedule(found, dates, '2026-10-17')['meeting_id'] == 'ran2#136'
    # Nothing current yet: keep the latest meeting rather than an undated one.
    assert sources.choose_schedule(found[:2], dates, '2026-10-17')['meeting_id'] == 'ran2#135bis'
    assert sources.choose_schedule(found[:2], {}, '2026-10-17')['meeting_id'] == 'ran2#135bis'


def test_newest_version_wins_and_the_live_copy_breaks_a_tie():
    sync = 'https://www.3gpp.org/ftp/Meetings_3GPP_SYNC/RAN2/Agenda/R2_135_Schedule_v11.docx'
    found = [candidate('ran2#135', 10), candidate('ran2#135', 11, uploaded='2026-08-29 21:10'),
             candidate('ran2#135', 11, url=sync, uploaded='2026-08-27 20:49')]
    assert sources.choose_schedule(found, {}, '2026-08-27')['url'] == sync
    with pytest.raises(ValueError, match='No RAN2 session schedule'):
        sources.choose_schedule([], {}, '2026-08-27')


def test_portal_titles_with_and_without_space():
    rows = [{'Title': '3GPPRAN2#135-bis', 'TBId': 380, 'Id': 1, 'StartDate': '2026-10-12 09:00:00',
             'EndDate': '2026-10-16 17:30:00', 'StartTimeZone': '(GMT+09:00) Seoul', 'Country': 'KR',
             'Location': 'South Korea'},
            {'Title': '3GPP RAN2#132', 'TBId': 380, 'Id': 2, 'StartDate': '2025-11-17 09:00:00',
             'EndDate': '2025-11-21 17:30:00', 'StartTimeZone': '(GMT-06:00) Central Time (US & Canada)',
             'Country': 'US', 'Location': 'Dallas'}]
    bis = sources.portal_metadata(rows, 'ran2#135bis')
    assert (bis['timezone'], bis['starts_at'], bis['ends_on']) == ('Asia/Seoul', '2026-10-12T09:00:00+09:00', '2026-10-16')
    assert sources.portal_metadata(rows, 'ran2#132')['timezone'] == 'America/Chicago'
    assert sources.portal_metadata(rows, 'ran2#136') is None


def mock_server(files):
    """3GPP-like server: listings for folders, 403 for missing paths."""
    def handler(request):
        url = str(request.url)
        if url.endswith('/'):
            names = sorted({u[len(url):].split('/')[0] + ('/' if '/' in u[len(url):] else '')
                            for u in files if u.startswith(url) and u != url})
            if not names:
                return httpx.Response(403)
            return httpx.Response(200, text=listing(names, base=url))
        if url in files:
            return httpx.Response(200, content=files[url], headers={'etag': f'"{len(files[url])}"'})
        return httpx.Response(403)
    return httpx.Client(transport=httpx.MockTransport(handler))


def test_fetch_bundle_reads_the_upcoming_archive_folder_and_its_agenda(tmp_path, monkeypatch):
    monkeypatch.setattr(sources.remote_files, 'CACHE_DIR', tmp_path / 'cache')
    sync = 'https://www.3gpp.org/ftp/Meetings_3GPP_SYNC/RAN2/Agenda/'
    archive = 'https://www.3gpp.org/ftp/tsg_ran/WG2_RL2/'
    files = {sync + V11: (FIXTURES / V11).read_bytes(),
             archive + 'TSGR2_135bis/Agenda/' + BIS: (FIXTURES / BIS).read_bytes(),
             archive + 'TSGR2_135bis/Agenda/agenda.csv': (FIXTURES / 'agenda_135bis.csv').read_bytes(),
             archive + 'TSGR2_136/Invitation/x.doc': b'invitation'}
    rows = [{'Title': '3GPPRAN2#135-bis', 'TBId': 380, 'Id': 1, 'StartDate': '2026-10-12 09:00:00',
             'EndDate': '2026-10-16 17:30:00', 'StartTimeZone': '(GMT+09:00) Seoul', 'Country': 'KR'},
            {'Title': '3GPPRAN2#136', 'TBId': 380, 'Id': 2, 'StartDate': '2026-11-16 09:00:00',
             'EndDate': '2026-11-20 17:30:00', 'StartTimeZone': '(GMT-07:00) Mountain Time (US & Canada)',
             'Country': 'CA'}]
    with mock_server(files) as http:
        bundle = sources.fetch_bundle(sources.DEFAULTS, {}, http, rows=rows, today='2026-10-05',
                                      cache_dir=tmp_path / 'inputs')
    assert (bundle.manifest['meeting_id'], bundle.manifest['version']) == ('ran2#135bis', 0)
    assert bundle.schedule == files[archive + 'TSGR2_135bis/Agenda/' + BIS]
    assert bundle.manifest['agenda']['url'].endswith('TSGR2_135bis/Agenda/agenda.csv')
    assert bundle.manifest['portal_metadata']['timezone'] == 'Asia/Seoul'
    # A real access error is not mistaken for an empty future folder.
    with mock_server({}) as http, pytest.raises(httpx.HTTPStatusError):
        sources.fetch_bundle(sources.DEFAULTS, {}, http, rows=rows, today='2026-10-05', cache_dir=tmp_path / 'inputs')


def test_bundle_cache_rejects_corruption(tmp_path):
    bundle = sources.Bundle({'schedule': {'sha256': sources.digest(b'docx')}, 'agenda': None}, b'docx', b'')
    bundle.save(tmp_path)
    assert sources.Bundle.load(tmp_path).schedule == b'docx'
    (tmp_path / 'agenda.csv').write_bytes(b'tampered')
    with pytest.raises(ValueError, match='agenda'):
        sources.Bundle.load(tmp_path)


def test_config_rejects_unknown_keys_and_bad_meetings(tmp_path, monkeypatch):
    path = tmp_path / 'config.json'
    monkeypatch.setattr(sources, 'CONFIG_PATH', path)
    path.write_text(json.dumps({'chair_url': 'x'}))
    with pytest.raises(ValueError, match='Unknown'):
        sources.load_config()
    path.write_text(json.dumps({'meetings': {'135bis': {'city': 'Seoul'}}}))
    with pytest.raises(ValueError, match='override'):
        sources.load_config()
    path.write_text(json.dumps({'meetings': {'135bis': {'timezone': 'Asia/Seoul'}}}))
    assert sources.override_for(sources.load_config(), 'ran2#135bis') == {'timezone': 'Asia/Seoul'}


# ----------------------------------------------------------- pipeline and CI

@pytest.fixture
def offline_config(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = {**sources.DEFAULTS, 'local_agenda': str(FIXTURES / 'agenda_135.csv'),
           'meetings': {'135': {k: v for k, v in META_135.items()}}}
    monkeypatch.setattr(pipeline, 'load_config', lambda: cfg)
    monkeypatch.setattr(sources, 'load_config', lambda: cfg)
    monkeypatch.setattr(pipeline, 'portal_rows', lambda: pytest.fail('Portal queried offline'))
    monkeypatch.setattr(sources, 'portal_rows', lambda: [])
    monkeypatch.setattr(pipeline, 'has_key', lambda: False)   # never reach Gemini from tests
    return cfg


def test_local_build_needs_no_network_and_records_state(offline_config, tmp_path, monkeypatch):
    monkeypatch.setattr(httpx.Client, 'send', lambda *a, **k: pytest.fail('Offline network request'))
    options = SimpleNamespace(local=str(FIXTURES / V11), no_download=True, output_dir=tmp_path / 'out')
    schedule = pipeline.build_schedule(options)
    assert schedule.wg_id == 'ran2' and schedule.meeting_name == 'RAN2#135'
    assert schedule.source_files == [V11, 'agenda.csv']
    assert schedule.source_labels == {V11: 'Schedule', 'agenda.csv': 'Agenda'}
    assert schedule.starts_at == '2026-08-24T09:00:00+02:00'
    state = sources.read_json(tmp_path / 'out/ran2/.schedule_state.json')
    assert state['meeting_id'] == 'ran2#135' and state['metadata']['timezone'] == 'Europe/Amsterdam'
    assert state['llm'] == {'used': False}   # no GEMINI_API_KEY: the rules' reading as is
    # The next offline build reuses the downloaded inputs.
    again = pipeline.build_schedule(SimpleNamespace(no_download=True, output_dir=tmp_path / 'out'))
    assert again.days == schedule.days


def test_a_new_meeting_without_dates_fails_clearly(offline_config, tmp_path, monkeypatch):
    offline_config['meetings'] = {}
    monkeypatch.setattr(pipeline, 'portal_rows', lambda: None)
    with pytest.raises(ValueError, match='Add them under "meetings"'):
        pipeline.build_schedule(SimpleNamespace(local=str(FIXTURES / V11), output_dir=tmp_path / 'out'))
    assert not (tmp_path / 'out/ran2/.schedule_state.json').exists()


def test_override_dates_drop_stale_portal_instants():
    cfg = {**sources.DEFAULTS, 'meetings': {'135': {'starts_on': '2026-08-25'}}}
    live = {**META_135}
    result = pipeline.meeting_metadata('ran2#135', cfg, {}, live)
    assert result['starts_on'] == '2026-08-25' and 'starts_at' not in result and result['ends_at']


def test_lifecycle_check_stages_without_advancing_state(offline_config, tmp_path, monkeypatch):
    bundle = sources.local_bundle(FIXTURES / V11, offline_config)
    monkeypatch.setattr(sources, 'fetch_bundle', lambda *a, **k: bundle)
    before = {'meeting_id': 'ran2#134'}
    sources.write_json(sources.OUTPUT / '.schedule_state.json', before)
    result = lifecycle.check_updates()
    assert result.changed and 'meeting_id changed' in result.reasons and not result.errors
    assert sources.read_json(sources.OUTPUT / '.schedule_state.json') == before
    lifecycle.prepare_build()
    assert sources.Bundle.load(sources.DOWNLOADS / 'prepared').schedule == bundle.schedule
    sources.write_json(sources.OUTPUT / '.schedule_state.json', {**bundle.manifest, 'metadata': META_135})
    assert not lifecycle.check_updates().changed
    monkeypatch.setattr(sources, 'fetch_bundle', lambda *a, **k: (_ for _ in ()).throw(ValueError('down')))
    assert lifecycle.check_updates().errors == ['RAN2 source check failed: down']


def test_prepared_bundle_is_consumed_once_and_checked_against_config(offline_config, tmp_path):
    bundle = sources.local_bundle(FIXTURES / V11, offline_config)
    bundle.manifest['config_hash'] = 'stale'
    bundle.save(sources.DOWNLOADS / 'prepared')
    with pytest.raises(ValueError, match='Prepared inputs'):
        pipeline.build_schedule(SimpleNamespace(output_dir=tmp_path / 'out'))
    assert not (sources.DOWNLOADS / 'prepared').exists()


def test_reset_keeps_the_last_successful_output(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    sources.write_json(sources.OUTPUT / '.schedule_state.json', {'meeting_id': 'ran2#135'})
    sources.write_json(sources.DOWNLOADS / 'inputs/manifest.json', {})
    lifecycle.reset_cache()
    assert (sources.OUTPUT / '.schedule_state.json').exists() and not sources.DOWNLOADS.exists()


def test_lifecycle_import_is_isolated():
    result = subprocess.run([sys.executable, '-c', """
import sys
from working_groups.registry import get_working_group
get_working_group('ran2')
assert not any(name.startswith(('working_groups.ran1', 'working_groups.ran_plenary', 'google.genai', 'docx'))
               for name in sys.modules)
"""], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_supplements_are_escaped_and_absent_elsewhere():
    from shared.topic_references import render_supplements
    from schedule_fixture import build_schedule as plenary_schedule
    attack = '<script>alert(1)</script>'
    html = render_supplements([{'type': 'heading', 'text': attack},
                               {'type': 'paragraph', 'text': f'a\n{attack}', 'bold': True},
                               {'type': 'table', 'header': True, 'rows': [[attack], ['x\ny']]},
                               {'type': 'table', 'rows': []}])
    assert '<script>' not in html and 'a<br>&lt;script&gt;' in html and '<td>x<br>y</td>' in html
    assert render_supplements([]) == ''
    assert BeautifulSoup(generate_html(plenary_schedule()), 'html.parser').select_one('.supplements') is None


# --------------------------------------------------------------- LLM fallback

from working_groups.ran2 import llm   # noqa: E402

HEADER_ABOVE = ['[7.1] NR19 AI/ML PHY [0] (Erlin)', '@17:30-18:30', '[8.1] NR20 AI/M PHY [1] (Erlin)',
                '@18:30-19:30', '[8.1.1]', '[8.1.2]']


def doubtful(document):
    seen = []
    sessions.interpret(document, None, refine=lambda cells: seen.extend(cells) or {})
    return {item['cell']['id']: item['reasons'] for item in seen}


def test_cells_the_rules_are_unsure_about_are_flagged_with_reasons():
    document = synthetic([
        ('main', '08:30', '10:30', ['[9.3.1] User Plane', '[9.3.1.2] QoS, QoE']),
        ('brk1', '17:30', '19:30', HEADER_ABOVE),
        ('brk2', '17:00', '19:00', ['[8.2] NR20 AIoT [2] (Nathan)', 'Overflow, end by 18:30']),
        ('brk3', '08:30', '10:30', ['CB Sergio NTN (from 9:00)']),
        ('main', '11:00', '13:00', ['@ TBD', 'Something']),
        ('brk1', '11:00', '13:00', ['[8.10] NR20 MIMO', '[7.7] NR19 XR cont. (~15 minutes)']),
        ('brk2', '11:00', '13:00', ['11:00-12:00 [004] (Ericsson, Nokia)']),
    ], slots=[('08:30', '10:30'), ('11:00', '13:00'), ('17:00', '19:30')])
    reasons = doubtful(document)
    assert set(reasons) == {'c1', 'c2', 'c3', 'c4', 'c5'}   # plain cells, "[8.10]" and "[004]" are clear
    assert reasons['c1'] == ["header written right above a time marker: '[8.1] NR20 AI/M PHY [1] (Erlin)' / '@18:30-19:30'",
                             'a part has no line that names it']
    assert reasons['c2'] == ["time written inside a line: 'Overflow, end by 18:30'"]
    assert reasons['c4'] == ["unreadable time marker: '@ TBD'"]
    assert reasons['c5'] == ["time written inside a line: '[7.7] NR19 XR cont. (~15 minutes)'"]


def answer(cell_id, *parts):
    return {'id': cell_id, 'sessions': [dict(zip(('start', 'end', 'title', 'chair', 'agenda_items', 'offline', 'lines'), p))
                                        for p in parts]}


class FakeGemini:
    def __init__(self, *responses):
        self.responses, self.calls = list(responses), []

    def __call__(self, payload, schema, model):
        self.calls.append(payload)
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response


def test_checked_llm_answers_replace_the_rules_for_uncertain_cells(tmp_path):
    document = synthetic([('brk1', '17:30', '19:30', HEADER_ABOVE),
                          ('main', '17:30', '19:30', ['[9.3.2] 6GR Control Plane'])], slots=[('17:30', '19:30')])
    gemini = FakeGemini({'cells': [answer('c0', ('17:30', '18:30', 'NR19 AI/ML PHY', 'Erlin', ['7.1'], False, [0]),
                                          ('18:30', '19:30', 'NR20 AI/M PHY', 'Erlin', ['8.1', '8.1.1', '8.1.2'], False,
                                           [1, 2, 3, 4, 5]))]})
    refine = llm.CellRefiner('test-model', gemini, tmp_path)
    found = [(s['start'], s['end'], s['rooms'][0], s['name'], s['chair'], s['agenda_items'])
             for s in sessions.interpret(document, None, refine)]
    assert found == [('17:30', '18:30', 'brk1', 'NR19 AI/ML PHY', 'Erlin', ['7.1']),
                     ('17:30', '19:30', 'main', '6GR Control Plane', None, ['9.3.2']),
                     ('18:30', '19:30', 'brk1', 'NR20 AI/M PHY', 'Erlin', ['8.1', '8.1.1', '8.1.2'])]
    assert refine.summary == {'model': 'test-model', 'uncertain': 1, 'applied': 1, 'kept_rules': 0}
    assert len(gemini.calls) == 1 and '"c0"' in gemini.calls[0] and '"c1"' not in gemini.calls[0]
    # The same input is answered from the cache.
    again = llm.CellRefiner('test-model', FakeGemini(), tmp_path)
    assert len(sessions.interpret(document, None, again)) == 3 and again.summary['applied'] == 1


@pytest.mark.parametrize('bad, problem', [
    (('17:30', '18:00', 'NR19 AI/ML PHY', 'Erlin', ['7.1'], False, [0, 1, 2, 3, 4, 5]), 'not written in the source'),
    (('17:30', '19:30', 'AI/ML for physical layer', 'Erlin', ['7.1'], False, [0, 1, 2, 3, 4, 5]), 'title words'),
    (('17:30', '19:30', 'NR19 AI/ML PHY', 'Ericsson', ['7.1'], False, [0, 1, 2, 3, 4, 5]), 'chair'),
    (('17:30', '19:30', 'NR19 AI/ML PHY', 'Erlin', ['7.2'], False, [0, 1, 2, 3, 4, 5]), 'agenda item'),
    (('17:30', '19:30', 'NR19 AI/ML PHY', 'Erlin', ['7.1'], False, [0, 1, 2]), 'not cited'),
])
def test_unchecked_answers_are_retried_once_then_the_rules_stand(tmp_path, bad, problem, capsys):
    document = synthetic([('brk1', '17:30', '19:30', HEADER_ABOVE)], slots=[('17:30', '19:30')])
    rules = blocks(document)
    gemini = FakeGemini({'cells': [answer('c0', bad)]}, {'cells': [answer('c0', bad)]})
    refine = llm.CellRefiner('test-model', gemini, tmp_path)
    found = [(s['start'], s['end'], s['rooms'][0], s['name'], s['chair'], s['agenda_items'])
             for s in sessions.interpret(document, None, refine)]
    assert found == rules
    assert len(gemini.calls) == 2 and problem in gemini.calls[1]
    assert refine.summary['kept_rules'] == 1 and problem in capsys.readouterr().out


def test_an_llm_outage_never_breaks_the_build(tmp_path):
    document = synthetic([('brk1', '17:30', '19:30', HEADER_ABOVE)], slots=[('17:30', '19:30')])
    refine = llm.CellRefiner('test-model', FakeGemini(RuntimeError('quota')), tmp_path)
    assert blocks(document) == [(s['start'], s['end'], s['rooms'][0], s['name'], s['chair'], s['agenda_items'])
                                for s in sessions.interpret(document, None, refine)]
    assert refine.summary['error'] == 'quota' and not list(tmp_path.glob('*.json'))


def test_real_cell_with_a_written_end_is_read_by_the_llm(tmp_path):
    """RAN2#135 v11, Wednesday Brk 2: "Overflow from afternoon session, end by 18:30"."""
    document = extract(V11)
    wednesday = next(d for d in document['days'] if d['day'] == 'Wednesday')
    cell = next(c for c in wednesday['cells'] if c['rooms'] == ['brk2'] and c['start'] == '17:00')
    assert cell['lines'] == ['[8.2] NR20 AIoT [2] (Nathan)', 'Overflow from afternoon session, end by 18:30']
    gemini = FakeGemini({'cells': [answer(cell['id'], ('17:00', '18:30', 'NR20 AIoT', 'Nathan', ['8.2'], False, [0, 1]))]},
                        {'cells': []})
    refine = llm.CellRefiner('test-model', gemini, tmp_path)
    schedule = sessions.make_schedule(document, agenda('agenda_135.csv'), META_135, [V11], 'now', refine=refine)
    aiot = find(schedule, 'Wednesday', 'brk2', '17:00')
    assert (aiot.end_time, aiot.name, aiot.chair, aiot.group_header) == ('18:30', 'NR20 AIoT', 'Nathan', 'NR Rel-20')
    # The other uncertain cells got no answer, so they keep the rules' reading.
    assert refine.summary['applied'] == 1 and refine.summary['kept_rules'] == refine.summary['uncertain'] - 1


def plain_document():
    """A schedule written as paragraphs: no table for the rules to read."""
    return docx_bytes([], title='RAN2-136 Session Schedule', after=[
        'Monday', 'Main room 09:00-10:30: [1], [2] Opening (Diana)',
        'Breakout 1 room 11:00-13:00: [8.3] NR20 AI mobility (Kyeongin)'])


def plain_answer(blocks_by_text):
    ref = {text: key for key, text in blocks_by_text.items()}
    return {'rooms': [{'id': 'main', 'name': 'Main room'}, {'id': 'brk1', 'name': 'Breakout 1 room'}],
            'days': [{'day': 'Monday', 'slots': [{'start': '09:00', 'end': '10:30'}, {'start': '11:00', 'end': '13:00'}]}],
            'sessions': [
                {'day': 'Monday', 'start': '09:00', 'end': '10:30', 'room_ids': ['main'], 'title': 'Opening',
                 'chair': 'Diana', 'agenda_items': ['1', '2'], 'offline': False,
                 'refs': [ref['Main room 09:00-10:30: [1], [2] Opening (Diana)']]},
                {'day': 'Monday', 'start': '11:00', 'end': '13:00', 'room_ids': ['brk1'], 'title': 'NR20 AI mobility',
                 'chair': 'Kyeongin', 'agenda_items': ['8.3'], 'offline': False,
                 'refs': [ref['Breakout 1 room 11:00-13:00: [8.3] NR20 AI mobility (Kyeongin)']]}]}


def test_an_unrecognised_layout_is_read_by_the_llm(offline_config, tmp_path, monkeypatch):
    data = plain_document()
    with pytest.raises(docmod.LayoutError):
        docmod.extract_document(data, 'R2_136_Schedule_v00.docx')
    texts = {b['id']: b['text'] for b in docmod.outline(Document(io.BytesIO(data))) if 'text' in b}
    gemini = FakeGemini(plain_answer(texts))
    local = tmp_path / 'R2_136_Schedule_v00.docx'
    local.write_bytes(data)
    offline_config['local_agenda'] = None
    offline_config['meetings'] = {'136': {'starts_on': '2026-11-16', 'ends_on': '2026-11-20', 'timezone': 'America/Edmonton'}}
    monkeypatch.setattr(pipeline, 'CACHE', tmp_path / 'cache')
    monkeypatch.setattr(llm, 'CACHE', tmp_path / 'cache')
    schedule = pipeline.build_schedule(SimpleNamespace(local=str(local), output_dir=tmp_path / 'out'), request=gemini)
    monday = schedule.days[0]
    assert [r.name for r in monday.rooms] == ['Main room', 'Breakout 1 room']
    assert [(s.start_time, s.name, s.chair, s.agenda_item) for s in monday.sessions] == [
        ('09:00', 'Opening', 'Diana', '1, 2'), ('11:00', 'NR20 AI mobility', 'Kyeongin', '8.3')]
    assert schedule.supplements[0]['text'].startswith('The schedule table layout was not recognised')
    state = sources.read_json(tmp_path / 'out/ran2/.schedule_state.json')
    assert state['llm'] == {'model': offline_config['model'], 'document': True}


def test_an_unchecked_document_answer_fails_the_build(tmp_path):
    data = plain_document()
    texts = {b['id']: b['text'] for b in docmod.outline(Document(io.BytesIO(data))) if 'text' in b}
    invented = plain_answer(texts)
    invented['sessions'][0]['start'] = '09:15'   # not written anywhere
    gemini = FakeGemini(invented, invented)
    with pytest.raises(docmod.LayoutError, match='could not read the schedule either'):
        llm.read_document(data, 'R2_136_Schedule_v00.docx', 'test-model', gemini, tmp_path)
    assert len(gemini.calls) == 2


def test_without_a_key_an_unrecognised_layout_still_fails(offline_config, tmp_path):
    local = tmp_path / 'R2_136_Schedule_v00.docx'
    local.write_bytes(plain_document())
    with pytest.raises(docmod.LayoutError):
        pipeline.build_schedule(SimpleNamespace(local=str(local), output_dir=tmp_path / 'out'))
