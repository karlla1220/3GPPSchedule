"""Office documents as blocks for the whole-document LLM fallback.

When a WG's rules cannot read a schedule's layout, the whole file goes to the
LLM (``llm_cells.read_document``). The layout is what broke, so these outlines
keep everything a reader could need in any arrangement, each piece with an id
the answer cites: text in order, table cells with their spans and merges, and
the marks chairs use to mean something (struck text, coloured text, cell fill,
highlight). Choices and the tools compared: ref_docs/office-document-readers.md.

- DOCX: the docx-cli AST (text boxes in reading order, tracked changes,
  merges and symbols in one tree). Without the binary (local runs, tests) the
  same blocks come from python-docx. Ids are docx-cli locators: ``p12``,
  ``t1``, ``t1:r3c2``.
- PPTX: every slide's shapes, groups and tables in drawing order, with their
  place on the slide and run marks inside table cells (RAN4 strikes single runs).
  Ids: ``s2``, ``s2:sh5``, ``s2:sh5:r3c2``.
- XLSX: every sheet's non-empty cells with fill, strike and colour, and its
  merged ranges. Ids: ``Sheet1!B4``.

Accepted text only: tracked insertions kept, deletions dropped. Struck text is
kept apart in ``struck`` so the LLM sees what the chair crossed out.
"""
from __future__ import annotations

import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile

from defusedxml import ElementTree

DOCX_CLI_VERSION = '0.28.0'
INK = {None, '', 'auto', '000000'}


def texts(blocks: list[dict]) -> dict[str, str]:
    """Every id's text, for checking what an answer cites."""
    found = {}

    def walk(node):
        if isinstance(node, dict):
            if node.get('id') and isinstance(node.get('text'), str):
                found[node['id']] = node['text']
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for value in node:
                walk(value)

    walk(blocks)
    return found


def _marks(entry: dict, struck: list[str], colours: set, highlights: set) -> dict:
    if struck:
        entry['struck'] = ' '.join(''.join(struck).split())
    if colours:
        entry['colours'] = sorted(colours)
    if highlights:
        entry['highlight'] = sorted(highlights)
    return entry


# ------------------------------------------------------------------- DOCX

def docx_cli() -> str | None:
    """The docx-cli binary: $DOCX_CLI, else ``docx`` on PATH. CI installs a pinned one."""
    configured = os.environ.get('DOCX_CLI')
    if configured is not None:
        return configured or None
    return shutil.which('docx')


def docx_blocks(data: bytes) -> list[dict]:
    binary = docx_cli()
    if binary:
        with tempfile.NamedTemporaryFile(suffix='.docx') as handle:
            handle.write(data)
            handle.flush()
            result = subprocess.run([binary, 'read', handle.name, '--ast'], capture_output=True, timeout=120)
        if result.returncode == 0:
            return blocks_from_ast(json.loads(result.stdout))
        print(f'docx-cli failed ({result.stderr.decode(errors="replace")[:200]}); reading with python-docx')
    return blocks_from_python_docx(data)


def _ast_runs(runs: list) -> tuple[str, list[str], set, set]:
    text, struck, colours, highlights = [], [], set(), set()
    for run in runs or []:
        kind = run.get('type')
        if (run.get('trackedChange') or {}).get('kind') in ('del', 'moveFrom'):
            continue
        piece = {'text': run.get('text', ''), 'tab': '\t', 'break': '\n'}.get(kind, run.get('text', ''))
        if not piece:
            continue
        if run.get('strike') or run.get('dstrike'):
            struck.append(piece)
            continue
        text.append(piece)
        if run.get('color') not in INK and piece.strip():
            colours.add(run['color'])
        if run.get('highlight') and piece.strip():
            highlights.add(run['highlight'])
    return ''.join(text), struck, colours, highlights


def _ast_paragraph(node: dict) -> dict:
    text, struck, colours, highlights = _ast_runs(node.get('runs'))
    return _marks({'id': node['id'], 'text': text}, struck, colours, highlights)


def _ast_cell_text(blocks: list) -> tuple[str, list[str], set, set]:
    lines, struck, colours, highlights = [], [], set(), set()
    for block in blocks or []:
        if block.get('type') == 'paragraph':
            text, s, c, h = _ast_runs(block.get('runs'))
            lines.append(text)
        elif block.get('type') == 'table':   # a table inside a cell: its text in order
            for row in block.get('rows', []):
                for cell in row.get('cells', []):
                    text, s, c, h = _ast_cell_text(cell.get('blocks'))
                    lines.append(text)
                    struck += s
                    colours |= c
                    highlights |= h
            continue
        else:
            continue
        struck += s
        colours |= c
        highlights |= h
    return '\n'.join(lines).strip('\n'), struck, colours, highlights


def blocks_from_ast(ast: dict) -> list[dict]:
    """The docx-cli AST (``docx read --ast``) as blocks."""
    out = []

    def add(block):
        kind = block.get('type')
        if kind == 'paragraph':
            entry = _ast_paragraph(block)
            if entry['text'].strip() or entry.get('struck'):
                out.append(entry)
            for run in block.get('runs') or []:   # text boxes anchored in the paragraph
                for inner in run.get('blocks') or []:
                    add(inner)
        elif kind == 'table':
            rows = []
            for r, row in enumerate(block.get('rows', [])):
                cells = []
                for c, cell in enumerate(row.get('cells', [])):
                    text, struck, colours, highlights = _ast_cell_text(cell.get('blocks'))
                    entry = {'id': f'{block["id"]}:r{r}c{c}', 'text': text}
                    if cell.get('gridSpan', 1) and cell.get('gridSpan', 1) > 1:
                        entry['span'] = cell['gridSpan']
                    if cell.get('vMerge'):
                        entry['merge'] = cell['vMerge']   # restart / continue
                    if cell.get('shading') not in INK | {'FFFFFF'}:
                        entry['fill'] = cell['shading']
                    cells.append(_marks(entry, struck, colours, highlights))
                rows.append(cells)
            out.append({'id': block['id'], 'rows': rows})
        elif kind in ('textBox', 'textbox'):
            for inner in block.get('blocks') or []:
                add(inner)

    for block in ast.get('blocks', []):
        add(block)
    return out


W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _docx_runs(element) -> tuple[str, list[str], set, set]:
    from shared.docx_text import struck as is_struck, tracked_deletion
    text, struck, colours, highlights = [], [], set(), set()
    for run in element.iter(W + 'r'):
        if tracked_deletion(run):
            continue
        piece = ''.join(node.text or '' if node.tag == W + 't' else '\t' if node.tag == W + 'tab'
                        else '\n' if node.tag in (W + 'br', W + 'cr') else '' for node in run)
        if not piece:
            continue
        if is_struck(run):
            struck.append(piece)
            continue
        text.append(piece)
        props = run.find(W + 'rPr')
        if props is not None and piece.strip():
            color = props.find(W + 'color')
            if color is not None and color.get(W + 'val') not in INK:
                colours.add(color.get(W + 'val'))
            mark = props.find(W + 'highlight')
            if mark is not None:
                highlights.add(mark.get(W + 'val'))
    return ''.join(text), struck, colours, highlights


def blocks_from_python_docx(data: bytes) -> list[dict]:
    """The same blocks from python-docx, for runs without docx-cli."""
    from docx import Document
    body = Document(io.BytesIO(data)).element.body
    out, p, t = [], 0, 0
    for child in body:
        if child.tag == W + 'p':
            text, struck, colours, highlights = _docx_runs(child)
            if text.strip() or struck:
                out.append(_marks({'id': f'p{p}', 'text': text}, struck, colours, highlights))
            p += 1
        elif child.tag == W + 'tbl':
            rows = []
            for r, row in enumerate(child.findall(W + 'tr')):
                cells = []
                for c, cell in enumerate(row.findall(W + 'tc')):
                    lines, struck, colours, highlights = [], [], set(), set()
                    for para in cell.iter(W + 'p'):
                        text, s, co, h = _docx_runs(para)
                        lines.append(text)
                        struck += s
                        colours |= co
                        highlights |= h
                    entry = {'id': f't{t}:r{r}c{c}', 'text': '\n'.join(lines).strip('\n')}
                    props = cell.find(W + 'tcPr')
                    if props is not None:
                        span = props.find(W + 'gridSpan')
                        if span is not None and int(span.get(W + 'val', '1')) > 1:
                            entry['span'] = int(span.get(W + 'val'))
                        merge = props.find(W + 'vMerge')
                        if merge is not None:
                            entry['merge'] = merge.get(W + 'val', 'continue')
                        fill = props.find(W + 'shd')
                        if fill is not None and fill.get(W + 'fill') not in INK | {'FFFFFF'}:
                            entry['fill'] = fill.get(W + 'fill')
                    cells.append(_marks(entry, struck, colours, highlights))
                rows.append(cells)
            out.append({'id': f't{t}', 'rows': rows})
            t += 1
    return out


# ------------------------------------------------------------------- PPTX

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def _pptx_paragraphs(body) -> tuple[str, list[str], set]:
    lines, struck, colours = [], [], set()
    for para in body.iter(A + 'p'):
        kept = []
        for node in para:
            if node.tag not in (A + 'r', A + 'fld'):
                if node.tag == A + 'br':
                    kept.append('\n')
                continue
            piece = ''.join(t.text or '' for t in node.iter(A + 't'))
            props = node.find(A + 'rPr')
            if props is not None and props.get('strike') not in (None, 'noStrike'):
                struck.append(piece)
                continue
            kept.append(piece)
            colour = _fill(props) if props is not None else None
            if colour and piece.strip():
                colours.add(colour)
        lines.append(''.join(kept))
    return '\n'.join(lines).strip('\n'), struck, colours


def _fill(props) -> str | None:
    """A solid fill's colour: "FF0000", or a theme colour by name ("accent1")."""
    if props is None:
        return None
    solid = props.find(f'{A}solidFill')
    if solid is None:
        return None
    rgb = solid.find(f'{A}srgbClr')
    if rgb is not None:
        return None if rgb.get('val') in INK else rgb.get('val')
    scheme = solid.find(f'{A}schemeClr')
    return scheme.get('val') if scheme is not None and scheme.get('val') not in ('tx1', 'dk1') else None


def _box(shape, size) -> list[int] | None:
    """Where a shape sits, in percent of the slide: [left, top, width, height]."""
    xfrm = shape.find(f'{P}spPr/{A}xfrm')
    if xfrm is None:
        xfrm = shape.find(f'{P}xfrm')   # a graphic frame (table)
    off = xfrm.find(f'{A}off') if xfrm is not None else None
    ext = xfrm.find(f'{A}ext') if xfrm is not None else None
    if off is None or ext is None or not size[0] or not size[1]:
        return None
    return [round(100 * int(off.get('x')) / size[0]), round(100 * int(off.get('y')) / size[1]),
            round(100 * int(ext.get('cx')) / size[0]), round(100 * int(ext.get('cy')) / size[1])]


def pptx_blocks(data: bytes) -> list[dict]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = set(archive.namelist())
        presentation = ElementTree.fromstring(archive.read('ppt/presentation.xml'))
        size_node = presentation.find(f'{P}sldSz')
        size = (int(size_node.get('cx')), int(size_node.get('cy'))) if size_node is not None else (0, 0)
        rels = ElementTree.fromstring(archive.read('ppt/_rels/presentation.xml.rels'))
        targets = {rel.get('Id'): rel.get('Target') for rel in rels}
        order = [targets[s.get(f'{R}id')] for s in presentation.iter(f'{P}sldId')]
        out = []
        for n, target in enumerate(order, 1):
            path = 'ppt/' + target.lstrip('/').removeprefix('ppt/')
            if path not in names:
                continue
            tree = ElementTree.fromstring(archive.read(path)).find(f'{P}cSld/{P}spTree')
            shapes = []
            counter = [0]

            def walk(parent):
                for shape in parent:
                    tag = shape.tag.rsplit('}', 1)[-1]
                    if tag == 'grpSp':
                        walk(shape)
                        continue
                    if tag not in ('sp', 'graphicFrame', 'cxnSp', 'pic'):
                        continue
                    counter[0] += 1
                    sid = f's{n}:sh{counter[0]}'
                    name_node = shape.find(f'.//{P}cNvPr')
                    entry = {'id': sid}
                    if name_node is not None and name_node.get('name'):
                        entry['name'] = name_node.get('name')
                    box = _box(shape, size)
                    if box:
                        entry['box'] = box
                    table = shape.find(f'.//{A}tbl')
                    if table is not None:
                        rows = []
                        for r, row in enumerate(table.findall(f'{A}tr')):
                            cells = []
                            for c, cell in enumerate(row.findall(f'{A}tc')):
                                body = cell.find(f'{A}txBody')
                                text, struck, colours = _pptx_paragraphs(body) if body is not None else ('', [], set())
                                item = {'id': f'{sid}:r{r}c{c}', 'text': text}
                                for attr, key in (('gridSpan', 'span'), ('rowSpan', 'rowspan')):
                                    if cell.get(attr):
                                        item[key] = int(cell.get(attr))
                                if cell.get('hMerge') or cell.get('vMerge'):
                                    item['merged'] = True
                                fill = _fill(cell.find(f'{A}tcPr'))
                                if fill and fill != 'FFFFFF':
                                    item['fill'] = fill
                                cells.append(_marks(item, struck, colours, set()))
                            rows.append(cells)
                        entry['rows'] = rows
                        shapes.append(entry)
                        continue
                    body = shape.find(f'{P}txBody')
                    if body is None:
                        continue
                    text, struck, colours = _pptx_paragraphs(body)
                    if text.strip() or struck:
                        entry['text'] = text
                        shapes.append(_marks(entry, struck, colours, set()))

            if tree is not None:
                walk(tree)
            out.append({'id': f's{n}', 'shapes': shapes})
        return out


# ------------------------------------------------------------------- XLSX

def xlsx_blocks(data: bytes) -> list[dict]:
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for sheet in book.worksheets:
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None or str(cell.value).strip() == '':
                    continue
                entry = {'id': f'{sheet.title}!{cell.coordinate}', 'text': str(cell.value)}
                if cell.font is not None and cell.font.strike:
                    entry['struck'] = entry.pop('text')
                    entry['text'] = ''
                colour = getattr(cell.font.color, 'rgb', None) if cell.font is not None and cell.font.color else None
                if isinstance(colour, str) and colour[-6:] not in INK:
                    entry['colours'] = [colour[-6:]]
                fill = cell.fill.fgColor.rgb if cell.fill is not None and cell.fill.fill_type == 'solid' else None
                if isinstance(fill, str) and fill[-6:] not in INK | {'FFFFFF'}:
                    entry['fill'] = fill[-6:]
                cells.append(entry)
        out.append({'id': sheet.title, 'cells': cells,
                    'merged': [str(rng) for rng in sheet.merged_cells.ranges]})
    return out


def blocks(data: bytes, name: str) -> list[dict]:
    """Blocks for any Office file by its extension."""
    suffix = Path(name).suffix.lower()
    if suffix in ('.docx', '.docm'):
        return docx_blocks(data)
    if suffix == '.pptx':
        return pptx_blocks(data)
    if suffix in ('.xlsx', '.xlsm'):
        return xlsx_blocks(data)
    raise ValueError(f'no outline for {name}')
