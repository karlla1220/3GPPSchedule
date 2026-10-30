"""Read an SA1 chair file (agenda, schedule and TDoc table in one DOCX).

The schedule is one table: weekdays across, fixed rows ``Q0``..``Q5`` down
(times in the second column), with Coffee and Lunch rows between them. Rooms
are not columns; a cell splits its rooms itself with ``Breakout 1:`` /
``=====`` / ``Breakout 2:``. The rooms are named in the ``MEETING ROOMS:``
paragraphs above the table, and the breakout chairs in the NOTE below it
(paragraphs, or a one-row table in #112 to #114).

The agenda is the heading tree of the same file's TDoc table: Word numbers
the headings, so the numbers are counted here. A heading with its own list
numbering is one level below its style (#112 "Editor's notes solving"), and a
hand-typed number (numbering removed) is taken as written.
"""
from __future__ import annotations

import io
import re

from docx import Document

from shared.docx_text import paragraph_text, W
from shared.portal_meetings import meeting_key

DAYS = ('Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday')
CLOCK = r'(\d{1,2})[:.](\d{2})'
SLOT = re.compile(CLOCK + r'\s*[-–—~/|]?\s*' + CLOCK)
# "3GPP TSG-SA WG1 Meeting #115", "3GPP TSG SA WG1 Meeting #107-bis"
TITLE = re.compile(r'TSG[- ]?SA\s*WG\s*1\s+Meeting\s*#\s*(\d+)(?:[- ]?(bis|b))?', re.I)
NAME = re.compile(r'SA1\s*#\s*(\d+)(?:[- ]?(bis|b))?', re.I)
# "AI, Computing – chaired by Vasil Aleksiev", "ISAC - chaired by Lola Awoniyi-Oteri"
CHAIRED = re.compile(r'^(.*?)\s*[–—-]\s*chaired by\s+(.+?)\s*$', re.I)
ROOM_ROLE = re.compile(r'(Breakout|Drafting)\s*(\d)', re.I)
HEADING = re.compile(r'(?:berschrift|Heading)\s*(\d)$', re.I)
HAND_NUMBER = re.compile(r'^(\d+(?:\.\d+)*)\.?\s+(.*)$')


class LayoutError(ValueError):
    """The document has no schedule table these rules recognise."""


class UnreadTable(LayoutError):
    """A weekday table whose slots the rules cannot read (time column, cells).

    ``partial`` is the document without its days: meeting, rooms, NOTE chairs,
    agenda and supplements, which the whole-document fallback builds on.
    """

    def __init__(self, message: str, partial: dict):
        super().__init__(message)
        self.partial = partial


def clock(hour, minute) -> str:
    hour, minute = int(hour), int(minute)
    if not (0 <= hour < 24 and 0 <= minute < 60):
        raise ValueError(f'Invalid time {hour}:{minute}')
    return f'{hour:02d}:{minute:02d}'


def squash(text: str) -> str:
    return ' '.join(text.replace(' ', ' ').split())


def lines_of(element) -> list[str]:
    """Non-empty text lines of a cell or paragraph, in order."""
    lines = []
    for paragraph in element.iter(W + 'p'):
        lines.extend(line for line in map(squash, paragraph_text(paragraph).split('\n')) if line)
    return lines


def meeting_id(number, suffix) -> str:
    return meeting_key(f'SA1#{number}{suffix or ""}')


def name_meeting(name: str) -> str | None:
    match = NAME.search(name)
    return meeting_id(match[1], match[2]) if match else None


def identify(doc, name: str) -> tuple[str | None, str]:
    """(meeting id, title paragraph); the title wins over the file name.

    The timetable-only file has neither ("Breakout sessions schedule_Prague_WED_1730.docx"):
    its meeting comes from where it was found.
    """
    for paragraph in doc.paragraphs[:12]:
        text = squash(paragraph_text(paragraph._p))
        match = TITLE.search(text)
        if match:
            return meeting_id(match[1], match[2]), text.split(' S1-')[0].strip()
    return name_meeting(name), ''


def _is_schedule(table) -> bool:
    first = [squash(paragraph_text(cell._tc)) for cell in table.rows[0].cells] if table.rows else []
    return sum(text in DAYS for text in first) >= 3


def has_schedule(data: bytes) -> bool:
    return any(_is_schedule(t) for t in Document(io.BytesIO(data)).tables)


# ------------------------------------------------------------------ the table

def room_list(lines: list[str]) -> list[dict]:
    """Rooms from "Plenary/ Breakout 1: D3/D4, Floor 3" lines.

    ``aliases`` are the cell heads that put a cell in the room: "plenary",
    "breakout 1", "drafting 2". The name is the part before the colon.
    """
    rooms = []
    for line in lines:
        role, _, place = line.partition(':')
        aliases = {f'{m[1].lower()} {m[2]}' for m in ROOM_ROLE.finditer(role)}
        if re.search(r'plenary', role, re.I):
            aliases.add('plenary')
        if not aliases or not place.strip():
            continue
        first = 'plenary' if 'plenary' in aliases else sorted(aliases)[0]
        rooms.append({'id': first.replace(' ', '-'), 'name': squash(role), 'place': squash(place),
                      'aliases': sorted(aliases)})
    return rooms


DEFAULT_ROOMS = [
    {'id': 'plenary', 'name': 'Plenary/ Breakout 1', 'place': '', 'aliases': ['breakout 1', 'drafting 1', 'plenary']},
    {'id': 'breakout-2', 'name': 'Breakout 2', 'place': '', 'aliases': ['breakout 2', 'drafting 2']},
]


def _times(text: str):
    match = SLOT.search(text)
    if not match:
        return None
    try:
        begin, end = clock(match[1], match[2]), clock(match[3], match[4])
    except ValueError:
        return None
    return (begin, end) if begin < end else None


def read_table(table) -> list[dict]:
    """Days with their slots, breaks and non-empty cells.

    A row is a slot when its first two cells give a time range ("Q1", "09:00 /
    10:30"); a row naming a break ("Coffee", "Lunch") lies between the slots
    around it. Cells in break rows keep the row and get its times later.
    """
    header = table.rows[0].cells
    columns = {}
    for index, cell in enumerate(header):
        text = squash(paragraph_text(cell._tc))
        if text in DAYS and text not in columns.values():
            columns[index] = text
    rows = []
    for ri, row in enumerate(table.rows[1:], 1):
        cells = row.cells
        firsts = [' '.join(lines_of(c._tc)) for c in cells[:2]]
        times = _times(' '.join(firsts))
        label = next((text for text in reversed(firsts) if text), '')
        rows.append({'index': ri, 'times': times, 'slot': bool(times), 'label': '' if times else label,
                     'cells': {day: lines_of(cells[ci]._tc) for ci, day in columns.items() if ci < len(cells)}})
    slots = [r['times'] for r in rows if r['times']]
    breaks = []
    for i, row in enumerate(rows):
        if row['times']:
            continue
        before = next((r['times'] for r in reversed(rows[:i]) if r['times']), None)
        after = next((r['times'] for r in rows[i + 1:] if r['times']), None)
        if before and after and before[1] < after[0]:
            row['times'] = (before[1], after[0])
            if row['label']:
                breaks.append({'name': row['label'], 'start': before[1], 'end': after[0]})
    days = []
    for day in columns.values():
        cells = []
        for row in rows:
            lines = row['cells'].get(day) or []
            if lines and row['times']:
                cells.append({'id': f'{day[:3].lower()}-r{row["index"]}', 'start': row['times'][0],
                              'end': row['times'][1], 'kind': 'slot' if row['slot'] else 'break',
                              'row': row['label'], 'lines': lines})
        days.append({'day': day, 'slots': list(slots), 'breaks': breaks, 'cells': cells})
    return days


# ------------------------------------------------------------- around the table

def _body(doc):
    """Body blocks in order: ('p', element) and ('tbl', table)."""
    tables = {t._tbl: t for t in doc.tables}
    for child in doc.element.body.iterchildren():
        if child.tag == W + 'p':
            yield 'p', child
        elif child.tag == W + 'tbl':
            yield 'tbl', tables[child]


def around(doc, table) -> tuple[list[str], list[str], list[str]]:
    """(room lines, NOTE paragraphs, chair lines) next to the schedule table."""
    blocks = list(_body(doc))
    at = next(i for i, (kind, value) in enumerate(blocks) if kind == 'tbl' and value._tbl is table._tbl)
    rooms = []
    for kind, value in reversed(blocks[:at]):
        if kind != 'p':
            break
        text = squash(paragraph_text(value))
        if re.match(r'MEETING ROOMS', text, re.I):
            break
        if text:
            rooms.insert(0, text)
    notes, chairs = [], []
    for kind, value in blocks[at + 1:]:
        if kind == 'tbl':
            lines = [line for row in value.rows for cell in row.cells for line in lines_of(cell._tc)]
            if len(value.rows) > 2 or not any(CHAIRED.match(line) for line in lines):
                break
            chairs.extend(line for line in dict.fromkeys(lines) if CHAIRED.match(line))
            continue
        text = squash(paragraph_text(value))
        if not text:
            continue
        if CHAIRED.match(text):
            chairs.append(text)
        elif not chairs and re.match(r'NOTE\b', text):
            notes.append(text)
        else:
            break
    return rooms, notes, chairs


def note_chairs(lines: list[str]) -> list[dict]:
    """[{'topic': 'AI, Computing', 'chair': 'Vasil Aleksiev'}] from the NOTE lines."""
    result = []
    for line in lines:
        match = CHAIRED.match(line)
        if match and match[1] and match[2]:
            result.append({'topic': match[1].strip(), 'chair': match[2].strip()})
    return result


# ------------------------------------------------------------------ the agenda

def _style_level(paragraph) -> int | None:
    props = paragraph.find(W + 'pPr')
    style = props.find(W + 'pStyle') if props is not None else None
    match = HEADING.search(style.get(W + 'val', '')) if style is not None else None
    return int(match[1]) if match else None


def _direct_numbering(paragraph) -> str | None:
    """The numId written on the paragraph itself, if any ("0" removes numbering)."""
    props = paragraph.find(W + 'pPr')
    numbering = props.find(W + 'numPr') if props is not None else None
    node = numbering.find(W + 'numId') if numbering is not None else None
    return node.get(W + 'val') if node is not None else None


def headings(body) -> list[tuple[object, str, str]]:
    """(paragraph element, agenda number, title) for each heading of the file's TDoc table."""
    counters, result = [], []
    for paragraph in body.iter(W + 'p'):
        level = _style_level(paragraph)
        if not level:
            continue
        text = squash(paragraph_text(paragraph))
        if not text:
            continue
        numbering = _direct_numbering(paragraph)
        if numbering == '0':
            match = HAND_NUMBER.match(text)
            if match:
                number, text = match[1], match[2]
                counters = [int(p) for p in number.split('.')]
                result.append((paragraph, number, text))
            continue
        if numbering is not None:
            level += 1   # a list of its own under the style's level (#112)
        counters = (counters + [0] * level)[:level]
        counters[-1] += 1
        result.append((paragraph, '.'.join(map(str, counters)), text))
    return result


def agenda_items(doc) -> dict[str, str]:
    """Agenda number -> title from the heading tree of the file's TDoc table."""
    result = {}
    for _, number, text in headings(doc.element.body):
        result.setdefault(number, text)
    return result


# ------------------------------------------------------------------- document

def extract_document(data: bytes, name: str) -> dict:
    doc = Document(io.BytesIO(data))
    meeting, title = identify(doc, name)
    table = next((t for t in doc.tables if _is_schedule(t)), None)
    if table is None:
        raise LayoutError('No schedule table with weekday columns')
    days = read_table(table)
    room_lines, notes, chair_lines = around(doc, table)
    rooms = room_list(room_lines)
    supplements = []
    if room_lines:
        supplements.append({'type': 'paragraph', 'text': '\n'.join(['MEETING ROOMS:', *room_lines]), 'bold': False})
    supplements += [{'type': 'paragraph', 'text': text, 'bold': False} for text in notes]
    if chair_lines:
        supplements.append({'type': 'paragraph', 'text': '\n'.join(chair_lines), 'bold': False})
    document = {'meeting_id': meeting, 'title': title, 'rooms': rooms or [dict(r) for r in DEFAULT_ROOMS],
                'days': days, 'chairs': note_chairs(chair_lines), 'agenda': agenda_items(doc),
                'supplements': supplements}
    if not any(day['cells'] for day in days):
        raise UnreadTable('The SA1 schedule table has no sessions', {**document, 'days': []})
    return document


def extract_agenda(data: bytes) -> dict[str, str]:
    return agenda_items(Document(io.BytesIO(data)))
