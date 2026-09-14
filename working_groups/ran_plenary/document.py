"""Lossless-enough OOXML input for the Plenary interpreter; no room inference here."""
from __future__ import annotations

import colorsys
import csv
import io
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
DAYS = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
TIME_RANGE = re.compile(r'(\d{1,2}):(\d{2})\s*[~–—-]\s*(\d{1,2}):(\d{2})')
MEETING = re.compile(r'(?:3GPP)?RAN\s*#?\s*(\d+)(?!\d)', re.I)
FILE = re.compile(r'RAN\s*#?\s*(\d+)\s+time\s*plan\s+v(\d+)\.(zip|docx)$', re.I)


def time_range(text):
    match = TIME_RANGE.search(text)
    if not match:
        raise ValueError(f'No explicit time range: {text!r}')
    h1, m1, h2, m2 = map(int, match.groups())
    if not (0 <= h1 < 24 and 0 <= h2 < 24 and 0 <= m1 < 60 and 0 <= m2 < 60):
        raise ValueError('Invalid source time')
    if h1 * 60 + m1 >= h2 * 60 + m2:
        raise ValueError('Non-positive source time range')
    return f'{h1:02d}:{m1:02d}', f'{h2:02d}:{m2:02d}'


def unpack_docx(data: bytes, name: str) -> bytes:
    if name.lower().endswith('.docx'):
        return data
    expected = FILE.fullmatch(name)
    if not expected:
        raise ValueError(f'Unsupported timeplan filename: {name}')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        candidates = []
        for info in archive.infolist():
            match = FILE.fullmatch(Path(info.filename).name)
            if match and match.group(3).lower() == 'docx' and match.groups()[:2] == expected.groups()[:2]:
                candidates.append(info)
        if len(candidates) != 1:
            raise ValueError('ZIP must contain exactly one matching timeplan DOCX')
        if candidates[0].file_size > 30_000_000:
            raise ValueError('Timeplan DOCX is too large')
        # Read the member directly; never extract archive paths onto the filesystem.
        return archive.read(candidates[0])


class Formatting:
    def __init__(self, doc):
        self.doc = doc
        self.styles = {s.style_id: s for s in doc.styles}
        self.theme = {}
        for rel in doc.part.rels.values():
            if rel.reltype.endswith('/theme'):
                root = etree.fromstring(rel.target_part.blob)
                for node in root.findall(f'.//{{{A}}}clrScheme/*'):
                    child = next(iter(node), None)
                    if child is not None:
                        self.theme[etree.QName(node).localname] = child.get('lastClr') or child.get('val')
        self.theme.update({k: self.theme.get(v) for k, v in {
            'text1': 'dk1', 'text2': 'dk2', 'background1': 'lt1', 'background2': 'lt2'}.items()})

    def properties(self, run, paragraph):
        yield run.find(f'{{{W}}}rPr')
        style_node = run.find(f'{{{W}}}rPr/{{{W}}}rStyle')
        style = self.styles.get(style_node.get(f'{{{W}}}val')) if style_node is not None else None
        for initial in (style, paragraph.style):
            seen = set()
            current = initial
            while current is not None and current.style_id not in seen:
                seen.add(current.style_id)
                yield current.element.find(f'{{{W}}}rPr')
                current = current.base_style
        yield self.doc.styles.element.find(f'{{{W}}}docDefaults/{{{W}}}rPrDefault/{{{W}}}rPr')

    def run(self, node, paragraph, source_id):
        props = list(self.properties(node, paragraph))
        color = next((p.find(f'{{{W}}}color') for p in props
                      if p is not None and p.find(f'{{{W}}}color') is not None), None)
        raw = {etree.QName(k).localname: v for k, v in color.attrib.items()} if color is not None else {}
        rgb = self.theme.get(raw.get('themeColor')) or raw.get('val', '000000')
        if rgb == 'auto':
            rgb = '000000'
        if re.fullmatch('[0-9a-fA-F]{6}', rgb or ''):
            # Word applies tint/shade to HSL luminance; tint wins when both exist.
            # https://learn.microsoft.com/en-us/openspecs/office_standards/ms-oi29500/8229a077-7fc8-4fba-96cc-c77b6a4fc768
            channels = [int(rgb[i:i+2], 16) / 255 for i in (0, 2, 4)]
            hue, lightness, saturation = colorsys.rgb_to_hls(*channels)
            if raw.get('themeTint'):
                tint = int(raw['themeTint'], 16) / 255
                lightness = lightness * tint + 1 - tint
            elif raw.get('themeShade'):
                lightness *= int(raw['themeShade'], 16) / 255
            rgb = ''.join(f'{round(c * 255):02X}' for c in colorsys.hls_to_rgb(hue, lightness, saturation))

        strike = False
        for tag in ('strike', 'dstrike'):
            element = next((p.find(f'{{{W}}}{tag}') for p in props
                            if p is not None and p.find(f'{{{W}}}{tag}') is not None), None)
            strike |= element is not None and element.get(f'{{{W}}}val', 'true') not in ('0', 'false', 'off')
        text = ''.join(e.text or '' if e.tag == f'{{{W}}}t' else '\n' if e.tag in (f'{{{W}}}br', f'{{{W}}}cr') else '\t'
                       for e in node.iter() if e.tag in (f'{{{W}}}t', f'{{{W}}}br', f'{{{W}}}cr', f'{{{W}}}tab'))
        return {'id': source_id, 'text': text, 'raw_color': raw, 'color': rgb, 'strike': strike}

    def paragraph(self, paragraph, source_id):
        runs = [self.run(r, paragraph, f'{source_id}.r{i}')
                for i, r in enumerate(paragraph._p.iter(f'{{{W}}}r'))]
        return {'id': source_id, 'text': ''.join(r['text'] for r in runs), 'runs': runs}


    def rich_paragraph(self, paragraph, source_id):
        """Preserve presentation separately from the schedule interpreter's input."""
        runs = []
        for i, node in enumerate(paragraph._p.iter(f'{{{W}}}r')):
            extracted = self.run(node, paragraph, f'{source_id}.r{i}')
            run = {key: extracted[key] for key in ('text', 'color', 'strike')}
            props = list(self.properties(node, paragraph))
            for tag, key in [('highlight', 'highlight'), ('b', 'bold'), ('i', 'italic'), ('u', 'underline')]:
                element = next((p.find(f'{{{W}}}{tag}') for p in props
                                if p is not None and p.find(f'{{{W}}}{tag}') is not None), None)
                value = element.get(f'{{{W}}}val', 'true') if element is not None else None
                run[key] = value if tag == 'highlight' else value not in (None, '0', 'false', 'off', 'none')
            shading = next((p.find(f'{{{W}}}shd') for p in props
                            if p is not None and p.find(f'{{{W}}}shd') is not None), None)
            run['shading'] = shading.get(f'{{{W}}}fill') if shading is not None else None
            for ancestor in node.iterancestors():
                if ancestor.tag == f'{{{W}}}hyperlink':
                    relationship = ancestor.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    rel = paragraph.part.rels.get(relationship)
                    if rel is not None and rel.is_external:
                        run['href'] = rel.target_ref
                    break
            runs.append(run)
        return {'runs': runs}


def extract_document(data: bytes, name: str) -> dict:
    doc = Document(io.BytesIO(unpack_docx(data, name)))
    formatting = Formatting(doc)
    headers = [formatting.paragraph(p, f'p{i}') for i, p in enumerate(doc.paragraphs) if p.text.strip()]
    title = next((p['text'] for p in headers if MEETING.search(p['text'])), '')
    meeting = MEETING.search(title)
    if not meeting:
        raise ValueError('Missing RAN meeting identity in DOCX')
    filename_match = FILE.fullmatch(name)
    if filename_match and int(filename_match.group(1)) != int(meeting.group(1)):
        raise ValueError('Filename and document meeting identities disagree')
    cells, breaks, topics = [], [], []
    days = []
    time_tables = 0
    for ti, table in enumerate(doc.tables):
        first = [c.text.strip() for c in table.rows[0].cells]
        day_columns = {i: value for i, value in enumerate(first) if value in DAYS}
        if day_columns:
            time_tables += 1
            days.extend(d for d in day_columns.values() if d not in days)
            for ri, row in enumerate(table.rows[1:], 1):
                start, end = time_range(row.cells[0].text)
                seen = set()
                for ci, day in day_columns.items():
                    cell = row.cells[ci]
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    covered_days = [d for i, d in day_columns.items() if row.cells[i]._tc is cell._tc]
                    paragraphs = [formatting.paragraph(p, f't{ti}.r{ri}.c{ci}.p{pi}')
                                  for pi, p in enumerate(cell.paragraphs) if p.text.strip()]
                    text = '\n'.join(p['text'] for p in paragraphs)
                    if not text.strip():
                        continue
                    entry = {'id': f't{ti}.r{ri}.c{ci}', 'days': covered_days,
                             'start': start, 'end': end, 'paragraphs': paragraphs}
                    if 'BREAK' in text.upper() and not re.search(r'\b(Main|Breakout)\b', text, re.I):
                        breaks.append({**entry, 'name': text})
                    else:
                        cells.append(entry)
        elif [s.lower() for s in first] == ['topics', 'documents', 'agenda item']:
            for ri, row in enumerate(table.rows[1:], 1):
                rich_text = {
                    key: [formatting.rich_paragraph(p, f't{ti}.r{ri}.c{ci}.p{pi}')
                          for pi, p in enumerate(row.cells[ci].paragraphs)]
                    for ci, key in enumerate(('title', 'documents', 'agenda_item'))}
                topics.append({'id': f't{ti}.r{ri}', 'title': row.cells[0].text.strip(),
                               'documents': row.cells[1].text.strip(),
                               'agenda_item': row.cells[2].text.strip(), 'rich_text': rich_text})
    if time_tables != 1 or not cells:
        raise ValueError('Expected one non-empty weekday timeplan table')
    return {'meeting_number': int(meeting.group(1)), 'title': title, 'days': days,
            'headers': headers, 'cells': cells, 'breaks': breaks, 'topics': topics}


def agenda_map(data: bytes) -> dict[str, str]:
    text = data.decode('utf-8-sig')
    result = {}
    for row in csv.reader(io.StringIO(text)):
        if not row or not any(v.strip() for v in row):
            continue
        if len(row) != 2:
            raise ValueError('Agenda must have two CSV columns')
        item, description = (v.strip() for v in row)
        if item.lower() in ('agenda item', 'item'):
            continue
        item = item.rstrip('.')
        if not re.fullmatch(r'\d+(?:\.\d+)*', item) or not description:
            raise ValueError('Invalid agenda row')
        if item in result and result[item] != description:
            raise ValueError(f'Conflicting agenda descriptions: {item}')
        result[item] = description
    if not result:
        raise ValueError('Empty agenda CSV')
    return result


def describe(items: str, descriptions: dict) -> list[dict]:
    result = []
    for item in dict.fromkeys(re.findall(r'\b\d+(?:\.\d+)*\b', items)):
        match = item
        while match not in descriptions and '.' in match:
            match = match.rsplit('.', 1)[0]
        if match not in descriptions:
            continue
        path = ['.'.join(match.split('.')[:i]) for i in range(1, len(match.split('.')) + 1)]
        result.append({'agenda_item': item, 'matched_agenda_item': match,
                       'description': descriptions[match],
                       'hierarchy': [{'agenda_item': p, 'description': descriptions.get(p)} for p in path]})
    return result
