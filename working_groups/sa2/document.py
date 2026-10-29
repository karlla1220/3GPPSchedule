"""Read an SA2 Session Plan DOCX.

The chair keeps one template (#172 onwards): a room table, a "Convener" line,
and one schedule table whose rows are quarters (Q0 to Q5) times streams (one
row per room), with Break/Coffee/Lunch rows between quarters, and whose columns
are five weekdays of two cells each. A cell may be merged across both cells of
a day (gridSpan) or down several rows (vMerge), so cells are identified by
their XML element. Tracked changes are read in accepted form.

The table is checked against that template; another layout raises LayoutError
rather than being guessed at. Paragraphs outside the table are kept as
supplements, which the page shows below the grid.
"""
from __future__ import annotations

import csv
import io
import re

from docx import Document
from docx.text.paragraph import Paragraph

from shared.portal_meetings import meeting_key

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
DAYS = ('Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday')
# S2-2610195_SA2-177-Session Plan - r7.docx, S2-2610195_SA2-177-Session Plan -r14 - Friday plan.docx,
# S2-260nnnn_SA2-177-Session Plan.zip (the first upload has no r), Draft S2-2606837_SA2-176-Session Plan - r1.docx
FILE = re.compile(r'.*?(?<![A-Za-z])SA2-(\d+)-Session Plan(?:\s*-\s*r(\d+))?(.*)\.(docx|zip)', re.I)
QUARTER = re.compile(r'Q(\d)')
SPAN = re.compile(r'(\d{2})(\d{2})\s*-\s*(\d{2})(\d{2})')
STREAM = re.compile(r'Stream\s*(\d)', re.I)
CONVENER = re.compile(r'Conven[eo]r\s*(\d+)\s*:\s*([^,]+)', re.I)


class LayoutError(ValueError):
    """The schedule table is not in the template these rules know."""


def meeting_id(number) -> str:
    return meeting_key(f'SA2#{number}')


def file_info(name: str):
    """(meeting_id, revision, rest of the name) for a Session Plan filename, or None.

    The first upload has no revision ("r0"); a later upload can repeat a
    revision with a suffix ("-r14 - Friday plan"), which the upload time orders.
    """
    match = FILE.fullmatch(name.strip())
    if not match:
        return None
    return meeting_id(match[1]), int(match[2] or 0), match[3].strip(' -')


def _deleted(run) -> bool:
    if any(a.tag in (W + 'del', W + 'moveFrom') for a in run.iterancestors()):
        return True
    props = run.find(W + 'rPr')
    if props is None:
        return False
    for tag in ('strike', 'dstrike'):
        node = props.find(W + tag)
        if node is not None and node.get(W + 'val', 'true') not in ('0', 'false', 'off'):
            return True
    return False


def paragraph_text(element) -> str:
    """Accepted text: insertions kept, deletions and struck runs dropped.

    python-docx's .text skips runs inside <w:ins>; r7, r11, r13 and r14 of
    SA2#177 carry the chair's tracked insertions (r11's Wednesday lunch row).
    """
    parts = []
    for run in element.iter(W + 'r'):
        if _deleted(run):
            continue
        for node in run:
            if node.tag == W + 't':
                parts.append(node.text or '')
            elif node.tag in (W + 'br', W + 'cr'):
                parts.append('\n')
            elif node.tag == W + 'tab':
                parts.append('\t')
    return ''.join(parts)


def squash(text: str) -> str:
    return ' '.join(text.replace(' ', ' ').split())


def cell_lines(tc) -> list[str]:
    """Non-blank text lines. SA2 marks no parts with blank lines (the table cells do)."""
    lines = []
    for paragraph in tc.iter(W + 'p'):
        lines.extend(squash(line) for line in paragraph_text(paragraph).split('\n'))
    return [line for line in lines if line]


def highlights(element) -> list[tuple[str, str]]:
    """(highlight colour or '', text) per run that is not deleted."""
    result = []
    for run in element.iter(W + 'r'):
        if _deleted(run):
            continue
        mark = run.find(f'{W}rPr/{W}highlight')
        result.append((mark.get(W + 'val') if mark is not None else '', paragraph_text(run)))
    return result


def clock(hour, minute) -> str:
    hour, minute = int(hour), int(minute)
    if not (0 <= hour < 24 and 0 <= minute < 60):
        raise ValueError(f'Invalid time {hour}:{minute}')
    return f'{hour:02d}:{minute:02d}'


# ------------------------------------------------------------------ the grid

def _grid(table) -> list[list[dict]]:
    """Each row as grid columns: {'tc': the cell that owns the text, 'own': True where it starts}.

    gridSpan repeats a cell across columns; vMerge="continue" points at the
    cell above in the same column.
    """
    rows = []
    for tr in table._tbl.iter(W + 'tr'):
        row = []
        for tc in tr.iter(W + 'tc'):
            props = tc.find(W + 'tcPr')
            span, merge = 1, None
            if props is not None:
                node = props.find(W + 'gridSpan')
                span = int(node.get(W + 'val')) if node is not None else 1
                node = props.find(W + 'vMerge')
                merge = None if node is None else (node.get(W + 'val') or 'continue')
            for offset in range(span):
                column = len(row)
                owner = tc
                if merge == 'continue' and rows and column < len(rows[-1]):
                    owner = rows[-1][column]['tc']
                row.append({'tc': owner, 'own': owner is tc and offset == 0, 'tr': tr})
        rows.append(row)
    return rows


def _text(tc) -> str:
    return squash(' '.join(cell_lines(tc)))


def _schedule_table(doc):
    """The last table whose first column holds Q0..Q5; LayoutError without one."""
    for table in reversed(doc.tables):
        firsts = [_text(row[0]['tc']) for row in _grid(table) if row]
        if sum(bool(QUARTER.fullmatch(f)) for f in set(firsts)) >= 3:
            return table
    raise LayoutError('No SA2 schedule table (quarters Q0..Q5 in the first column)')


def _read_table(table) -> dict:
    """Days, quarters, breaks and blocks of the schedule table, checked against the template."""
    grid = _grid(table)
    header = grid[0]
    days = []   # {'day', 'columns'} with the day's grid columns, left to right
    notes = []
    for column, cell in enumerate(header):
        text = _text(cell['tc'])
        word = text.split()[0].rstrip(':,') if text else ''
        day = next((d for d in DAYS if d.lower() == word.lower()), None)
        if day is None:
            continue
        if days and days[-1]['day'] == day:
            days[-1]['columns'].append(column)
            continue
        days.append({'day': day, 'columns': [column]})
        rest = squash(text[len(word):])
        if rest:
            notes.append({'day': day, 'text': rest})   # "(Revisions only unless stated)"
    if len(days) < 5 or any(len(d['columns']) != 2 for d in days):
        raise LayoutError('The SA2 schedule header is not five weekdays of two cells each')
    first_day = days[0]['columns'][0]

    quarters, breaks, rows = [], [], []
    for index, row in enumerate(grid[1:], 1):
        label = _text(row[first_day - 1]['tc']) if first_day >= 1 else ''
        quarter = QUARTER.fullmatch(_text(row[0]['tc']))
        stream = STREAM.fullmatch(label)
        if stream:
            # The quarter label is written once, on the quarter's first stream row.
            if not quarter and not (rows and rows[-1]['kind'] == 'stream'):
                raise LayoutError(f'Stream row {index} has no quarter')
            number = int(quarter[1]) if quarter else quarters[-1]['q']
            if not quarters or quarters[-1]['q'] != number:
                span = next((SPAN.search(_text(row[c]['tc'])) for c in range(1, first_day - 1)
                             if SPAN.search(_text(row[c]['tc']))), None)
                if span is None:
                    raise LayoutError(f'Q{number} has no time range')
                quarters.append({'q': number, 'start': clock(span[1], span[2]), 'end': clock(span[3], span[4]),
                                 'streams': {}})
            rows.append({'index': index, 'kind': 'stream', 'quarter': len(quarters) - 1,
                         'stream': int(stream[1]), 'label': row[first_day - 1]['tc']})
        elif label:
            breaks.append({'name': label, 'after': len(quarters) - 1})
            rows.append({'index': index, 'kind': 'break', 'break': len(breaks) - 1})
    if [q['q'] for q in quarters] != list(range(len(quarters))) or len(quarters) < 4:
        raise LayoutError('The SA2 schedule quarters are not Q0, Q1, ... in order')
    streams = sorted({r['stream'] for r in rows if r['kind'] == 'stream'})
    if streams != list(range(1, len(streams) + 1)) or len(streams) < 2:
        raise LayoutError('The SA2 schedule streams are not Stream 1, 2, ...')
    for item in breaks:
        # Break rows carry no times: the quarters around them do.
        if not 0 <= item['after'] < len(quarters) - 1:
            raise LayoutError(f'{item["name"]} row is not between two quarters')
        item['start'] = quarters[item['after']]['end']
        item['end'] = quarters[item['after'] + 1]['start']

    blocks = {}
    for info in rows:
        row = grid[info['index']]
        for day in days:
            for half, column in enumerate(day['columns'], 1):
                if column >= len(row):
                    continue
                tc = row[column]['tc']
                block = blocks.get(id(tc))
                if block is None:
                    block = blocks[id(tc)] = {'id': f'r{info["index"]}c{column}', 'day': day['day'],
                                              'lines': cell_lines(tc), 'rows': [], 'halves': set()}
                elif block['day'] != day['day']:
                    raise LayoutError(f'A cell spans {block["day"]} and {day["day"]}')
                if info not in block['rows']:
                    block['rows'].append(info)
                block['halves'].add(half)
    labels = {}   # stream -> highlight colour of its row label (Q0's label has none)
    for info in rows:
        if info['kind'] == 'stream':
            colour = next((c for c, text in highlights(info['label']) if c and text.strip()), '')
            if colour:
                labels.setdefault(info['stream'], colour)
        info.pop('label', None)
    cells = []
    for block in blocks.values():
        if not block['lines']:
            continue
        kinds = [r for r in block['rows'] if r['kind'] == 'stream']
        block['streams'] = sorted({r['stream'] for r in kinds})
        block['quarters'] = sorted({r['quarter'] for r in kinds})
        block['breaks'] = sorted({r['break'] for r in block['rows'] if r['kind'] == 'break'})
        block['halves'] = sorted(block['halves'])
        del block['rows']
        cells.append(block)
    return {'days': [d['day'] for d in days], 'quarters': quarters, 'breaks': breaks,
            'streams': streams, 'stream_colours': labels, 'cells': cells, 'notes': notes}


# ------------------------------------------------------------- outside the table

def room_names(doc, schedule_table, count: int) -> list[str]:
    """Room names by stream, from the room table above the schedule.

    "Main Room | B2 (Floor 2) [300] | Breakout 1 | B3 (Floor 2) [100] | ..."
    lists the rooms in stream order: the cell shading of each stream row
    matches its room, and the chair notes call SA2_Brk1 "Parallel Stream 2".
    """
    for table in doc.tables:
        if table._tbl is schedule_table._tbl:
            break
        texts = [_text(c['tc']) for row in _grid(table) for c in row if c['own']]
        names = [t for t in texts if re.fullmatch(r'Main(?:\s+Room)?|Break[- ]?out\s*\d+', t, re.I)]
        if len(names) >= count:
            return ['Main' if t.lower().startswith('main') else f'Breakout {re.search(r"\d+", t)[0]}'
                    for t in names[:count]]
    return [f'Stream {n}' for n in range(1, count + 1)]


def conveners(doc) -> dict[str, str]:
    """Convener names by highlight colour: "Convener 1: Mike, Convenor 2: Wanqiang".

    Each name carries the colour of the stream rows it convenes (green on
    Stream 2's labels, cyan on Stream 3's). A name without a colour is linked
    to nothing.
    """
    found = {}
    for paragraph in doc.paragraphs:
        runs = highlights(paragraph._p)
        text = ''.join(t for _, t in runs)
        if not CONVENER.search(text):
            continue
        position = 0
        spans = []
        for colour, part in runs:
            spans.append((position, position + len(part), colour))
            position += len(part)
        for match in CONVENER.finditer(text):
            colours = {c for a, b, c in spans if c and a < match.end(2) and b > match.start(2)}
            if len(colours) == 1:
                found.setdefault(colours.pop(), squash(match[2]))
    return found


def _bold(paragraph: Paragraph) -> bool:
    style_bold = False
    style = paragraph.style
    while style is not None:
        if style.font.bold is not None:
            style_bold = style.font.bold
            break
        style = style.base_style
    runs = [r for r in paragraph._p.iter(W + 'r') if not _deleted(r) and squash(paragraph_text(r))]
    if not runs:
        return False
    for run in runs:
        node = run.find(f'{W}rPr/{W}b')
        bold = style_bold if node is None else node.get(W + 'val', 'true') not in ('0', 'false', 'off')
        if not bold:
            return False
    return True


def _other_table(table) -> dict:
    rows = []
    for row in _grid(table):
        cells = ['\n'.join(cell_lines(c['tc'])) for c in row if c['own']]
        if any(cells):
            rows.append(cells)
    width = max((len(r) for r in rows), default=0)
    return {'type': 'table', 'header': False, 'rows': [r + [''] * (width - len(r)) for r in rows]}


def supplements(doc, schedule_table) -> list[dict]:
    """Everything outside the schedule table, in document order.

    The heading paragraphs up to the "session plan" title are left out: the
    first upload copies them from the previous meeting (#177's first plan
    named #176's dates).
    """
    blocks, titled = [], False
    for child in doc.element.body.iterchildren():
        if child.tag == W + 'tbl':
            if titled and child is not schedule_table._tbl:
                block = _other_table(next(t for t in doc.tables if t._tbl is child))
                if block['rows']:
                    blocks.append(block)
            continue
        if child.tag != W + 'p':
            continue
        text = '\n'.join(squash(line) for line in paragraph_text(child).split('\n') if squash(line))
        if not text:
            continue
        if not titled:
            titled = bool(re.search(r'session plan', text, re.I))
            continue
        bold = _bold(Paragraph(child, doc._body))
        if bold and len(text.split()) <= 8 and '\n' not in text:
            blocks.append({'type': 'heading', 'text': text})
        else:
            blocks.append({'type': 'paragraph', 'text': text, 'bold': bold})
    return blocks


def identify(doc, name: str) -> str:
    """The meeting id. The filename decides: the first upload's text names the previous meeting."""
    info = file_info(name)
    if info:
        return info[0]
    for paragraph in doc.paragraphs:
        match = re.search(r'SA\s*WG2\s+Meeting\s*#\s*(\d+)', paragraph_text(paragraph._p), re.I)
        if match:
            return meeting_id(match[1])
    raise ValueError('Cannot tell the SA2 meeting from the filename or the document')


def extract_document(data: bytes, name: str) -> dict:
    doc = Document(io.BytesIO(data))
    meeting = identify(doc, name)
    table = _schedule_table(doc)
    result = _read_table(table)
    if not result['cells']:
        raise LayoutError('The SA2 schedule table has no sessions')
    names = room_names(doc, table, len(result['streams']))
    result['rooms'] = [{'id': f'stream{n}', 'name': names[n - 1], 'stream': n} for n in result['streams']]
    result['conveners'] = conveners(doc)
    result['meeting_id'] = meeting
    # Table notes are added with the rules' own (sessions.py) below the supplements.
    result['supplements'] = supplements(doc, table)
    return result


def agenda_map(data: bytes) -> dict[str, str]:
    """agenda.csv in each meeting's Agenda folder: "item","description" (RAN's format).

    Malformed rows are skipped; descriptions are only an aid.
    """
    result = {}
    for row in csv.reader(io.StringIO(data.decode('utf-8-sig', errors='replace'))):
        if len(row) < 2:
            continue
        item, description = row[0].strip().rstrip('.'), squash(row[1])
        if re.fullmatch(r'\d+(?:\.\d+)*', item) and description:
            result.setdefault(item, description)
    return result
