"""Extract the RAN2 session schedule table without interpreting cell text.

The chair's DOCX holds one table: a time column, then one column per room
(Main, Brk 1, Brk 2, Brk 3). Full-width rows name a weekday or carry a day
note. A room cell may be merged vertically across rows, sometimes across two
time slots, so cells are identified by their XML element rather than by row.
Tracked changes are read in their accepted form.
"""
from __future__ import annotations

import csv
import io
import re

from docx import Document

from shared.portal_meetings import meeting_key

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
DAYS = ('Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday')
CLOCK = r'(\d{1,2})[:.](\d{2})'
SLOT = re.compile(CLOCK + r'\s*[-–—~]\s*' + CLOCK)
# R2_135_Schedule_v11.docx, R2_133b_Schedule_v14.docx, R2_131bis_Schedule v19.docx
FILE = re.compile(r'R2[_-]?(\d+)(bis|b)?[_ -]+Schedule[_ -]*v(\d+)\.docx', re.I)
TITLE = re.compile(r'RAN2\s*[-#]?\s*(\d+)\s*(bis|b)?(?![a-z0-9])', re.I)
# Morning coffee:   10:30 to 11:00
BREAK = re.compile(r'^\s*([A-Za-z][A-Za-z ]*?)\s*:\s*' + CLOCK + r'\s*(?:to|[-–—])\s*' + CLOCK)
# Mon 17:00-18:30, Tue 08:30-10:00
OFFLINE_TIME = re.compile(r'(?:^|\t)\s*(Mo|Tu|We|Th|Fr|Sa|Su)[a-z]*\.?\s+' + CLOCK + r'\s*[-–—]\s*' + CLOCK, re.I)
# [004], [xxx], [POST133bis][009], [AT135][110]; topic tags such as [NES] stay in titles
OFFLINE_ID = re.compile(r'\[\s*(\d{3}|x{3}|POST\s*\d+\w*|AT\s*\d+\w*)\s*\]', re.I)


def meeting_id(number, suffix) -> str:
    return meeting_key(f'RAN2#{number}{suffix or ""}')


def file_info(name: str):
    """Return (meeting_id, version) for a schedule filename, or None."""
    match = FILE.fullmatch(name.strip())
    if not match:
        return None
    return meeting_id(match[1], match[2]), int(match[3])


def clock(hour, minute) -> str:
    hour, minute = int(hour), int(minute)
    if hour < 7:
        hour += 12  # "6:45pm" style afternoon times written without a meridiem
    if not (0 <= hour < 24 and 0 <= minute < 60):
        raise ValueError(f'Invalid time {hour}:{minute}')
    return f'{hour:02d}:{minute:02d}'


def _deleted(run) -> bool:
    if any(a.tag in (W + 'del', W + 'moveFrom') for a in run.iterancestors()):
        return True
    props = run.find(W + 'rPr')
    if props is None:
        return False
    for tag in ('strike', 'dstrike'):
        node = props.find(W + tag)
        if node is not None and node.get(W + 'val', 'true') not in ('0', 'false', 'off'):
            return True
    return False


def paragraph_text(element) -> str:
    """Accepted text: insertions kept, deletions and struck runs dropped."""
    parts = []
    for run in element.iter(W + 'r'):
        if _deleted(run):
            continue
        for node in run:
            if node.tag == W + 't':
                parts.append(node.text or '')
            elif node.tag in (W + 'br', W + 'cr'):
                parts.append('\n')
            elif node.tag == W + 'tab':
                parts.append('\t')
    return ''.join(parts)


def cell_lines(cell) -> list[str]:
    """Text lines; one '' marks a blank line, which chairs use as a separator."""
    lines = []
    for paragraph in cell._tc.iter(W + 'p'):
        for line in paragraph_text(paragraph).split('\n'):
            line = ' '.join(line.replace(' ', ' ').split())
            if line or (lines and lines[-1]):
                lines.append(line)
    while lines and not lines[-1]:
        lines.pop()
    return lines


def room_info(text: str, index: int) -> dict:
    label = ' '.join(text.replace('*', ' ').split())
    match = re.fullmatch(r'(?:Brk|Break[- ]?out)\s*(\d+)(?:\s+room)?', label, re.I)
    if match:
        return {'id': f'brk{match[1]}', 'name': f'Breakout {match[1]}', 'label': label}
    if re.fullmatch(r'Main(?:\s+room)?', label, re.I):
        return {'id': 'main', 'name': 'Main', 'label': label}
    slug = re.sub(r'[^a-z0-9]+', '-', label.lower()).strip('-') or f'room{index}'
    return {'id': slug, 'name': label or f'Room {index}', 'label': label}


def _schedule_table(doc):
    for table in doc.tables:
        header = [' '.join(paragraph_text(c._tc).split()) for c in table.rows[0].cells]
        if any(re.match(r'Main\b', h, re.I) for h in header[1:]):
            return table, header
    raise ValueError('No RAN2 session table with a Main room column')


def _day_of(text: str):
    word = text.split()[0].rstrip(':,') if text.split() else ''
    return next((d for d in DAYS if d.lower() == word.lower()), None)


def extract_offline_list(paragraphs: list[str]) -> list[dict]:
    """Parse the tab-separated 'List of Offline Face to Face discussions'."""
    try:
        start = next(i for i, p in enumerate(paragraphs) if re.match(r'\s*List of Offline', p, re.I))
    except StopIteration:
        return []
    records, pending = [], []
    for text in paragraphs[start + 1:]:
        if not text.strip() or re.match(r'\s*Number\s*\t', text):
            continue
        pending.append(text)
        match = OFFLINE_TIME.search(text)
        if not match:
            continue
        head = ' '.join([*pending[:-1], text[:match.start()]])
        head = ' '.join(head.split())
        numbers = []
        while (lead := OFFLINE_ID.match(head)):
            numbers.append(lead[1])
            head = head[lead.end():].strip()
        after = [f.strip() for f in text[match.end():].split('\t') if f.strip()]
        day = next(d for d in DAYS if d.lower().startswith(match[1].lower()))
        records.append({
            'numbers': numbers,
            'title': ' '.join(OFFLINE_ID.sub(' ', head).split()),
            'day': day,
            'start': clock(match[2], match[3]), 'end': clock(match[4], match[5]),
            'place': after[0] if after else '',
            'coordinator': ' '.join(after[1:]) if len(after) > 1 else '',
        })
        pending = []
    return records


def extract_breaks(paragraphs: list[str]) -> list[dict]:
    breaks = []
    for text in paragraphs:
        match = BREAK.match(text)
        if match and re.search(r'coffee|lunch|break|tea', match[1], re.I):
            breaks.append({'name': ' '.join(match[1].split()),
                           'start': clock(match[2], match[3]), 'end': clock(match[4], match[5])})
    return breaks


def extract_document(data: bytes, name: str) -> dict:
    doc = Document(io.BytesIO(data))
    paragraphs = [paragraph_text(p._p) for p in doc.paragraphs]
    title = next((' '.join(p.split()) for p in paragraphs
                  if TITLE.search(p) and re.search(r'schedule', p, re.I)), '')
    match = TITLE.search(title)
    if not match:
        raise ValueError('Missing RAN2 meeting identity in the schedule title')
    meeting = meeting_id(match[1], match[2])
    info = file_info(name)
    if info and info[0] != meeting:
        raise ValueError(f'Filename meeting {info[0]} disagrees with document {meeting}')

    table, header = _schedule_table(doc)
    rooms = [room_info(text, i) for i, text in enumerate(header[1:], 1)]
    if len({r['id'] for r in rooms}) != len(rooms):
        raise ValueError('Duplicate room columns')
    days = []
    current = None
    for ri, row in enumerate(table.rows[1:], 1):
        cells = row.cells
        if len(cells) != len(rooms) + 1:
            raise ValueError(f'Row {ri} does not match the room columns')
        first = ' '.join(paragraph_text(cells[0]._tc).split())
        if len({id(c._tc) for c in cells}) == 1:
            day = _day_of(first)
            if day:
                current = {'day': day, 'slots': [], 'cells': [], 'notes': []}
                days.append(current)
            elif current is not None and first:
                current['notes'].append(first)
            continue
        slot = SLOT.search(first)
        if current is None:
            raise ValueError(f'Row {ri} precedes the first weekday')
        if not slot:
            if any(cell_lines(c) for c in cells[1:]):
                raise ValueError(f'Row {ri} has sessions but no time range: {first!r}')
            continue
        start, end = clock(slot[1], slot[2]), clock(slot[3], slot[4])
        if start >= end:
            raise ValueError(f'Non-positive slot {first!r}')
        if current['slots'] and start < current['slots'][-1][0]:
            # The clock went backwards under a blank full-width row: the
            # chair left the next weekday's label empty (seen in RAN2#135 v00).
            index = DAYS.index(current['day']) + 1
            if index >= len(DAYS) or any(d['day'] == DAYS[index] for d in days):
                raise ValueError(f'Row {ri} goes back in time without a weekday')
            current = {'day': DAYS[index], 'slots': [], 'cells': [], 'notes': []}
            days.append(current)
        if (start, end) not in current['slots']:
            current['slots'].append((start, end))  # row order, sorted after the table
        known = {id(c['_tc']): c for c in current['cells']}
        for ci, room in enumerate(rooms, 1):
            tc = cells[ci]._tc
            if any(c._tc is tc for c in cells[:ci]):
                continue  # horizontal merge: the leftmost column owns the text
            span = [rooms[i - 1]['id'] for i in range(1, len(cells)) if cells[i]._tc is tc]
            entry = known.get(id(tc))
            if entry is None:
                lines = cell_lines(cells[ci])
                entry = {'_tc': tc, 'id': f'r{ri}c{ci}', 'rooms': span, 'start': start, 'end': end,
                         'lines': lines}
                current['cells'].append(entry)
                known[id(tc)] = entry
            else:
                entry['start'], entry['end'] = min(entry['start'], start), max(entry['end'], end)
    for day in days:
        for cell in day['cells']:
            del cell['_tc']
        day['cells'] = [c for c in day['cells'] if c['lines']]
        day['slots'].sort()
    if not any(day['cells'] for day in days):
        raise ValueError('The RAN2 schedule table has no sessions')
    return {'meeting_id': meeting, 'title': title, 'rooms': rooms, 'days': days,
            'breaks': extract_breaks(paragraphs), 'offline': extract_offline_list(paragraphs),
            'paragraphs': [p for p in paragraphs if p.strip()]}


def agenda_map(data: bytes) -> dict[str, str]:
    """agenda.csv published in each meeting's Agenda folder: "item","description"."""
    result = {}
    for row in csv.reader(io.StringIO(data.decode('utf-8-sig'))):
        if not row or not any(v.strip() for v in row):
            continue
        if len(row) < 2:
            raise ValueError('Agenda rows need an item and a description')
        item, description = row[0].strip().rstrip('.'), ' '.join(row[1].split())
        if item.lower() in ('agenda item', 'item'):
            continue
        if not re.fullmatch(r'\d+(?:\.\d+)*', item) or not description:
            raise ValueError(f'Invalid agenda row: {row!r}')
        result[item] = description
    if not result:
        raise ValueError('Empty agenda CSV')
    return result
